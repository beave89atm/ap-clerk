"""Unique AP-inbox vendor catalog for receiving-owner mapping.

GET-only Graph metadata (From / subject / domain). Does not download PDFs,
send mail, or write KIMCO. Kyle fills Receiving owner on the sheet.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ap_clerk.graph import ALLOWED_MAILBOX
from ap_clerk.inbox import SKIP_CLASSES, sender_address, sender_name
from ap_clerk.pdf_invoice import (
    DOMAIN_VENDORS,
    SUBJECT_VENDORS,
    _COMPANY_WORD_RE,
    _looks_like_person_name,
    _title_company,
    company_from_subject_or_text,
    vendor_from_context,
)
from ap_clerk.rules import (
    classify_mail,
    distinctive_vendor_tokens,
    normalize_name,
    subject_has_invoice_bill_hint,
)

COUNT_WINDOW_START = date(2026, 8, 1)
SCAN_FLOOR = date(2026, 7, 1)

VENDOR_COLUMNS = [
    "Vendor",
    "Aliases seen",
    "Approx invoice count",
    "Sample from / subject",
    "Receiving owner",
    "Notes",
]

SKIP_COLUMNS = [
    "Bucket",
    "Count",
    "Sample from / subject",
    "Notes",
]

# Platform / mailbox noise — not a vendor. Subject may still name one.
PLATFORM_DOMAINS = frozenset(
    {
        "gmail.com",
        "googlemail.com",
        "outlook.com",
        "hotmail.com",
        "live.com",
        "yahoo.com",
        "aol.com",
        "icloud.com",
        "me.com",
        "msn.com",
        "microsoft.com",
        "microsoftonline.com",
        "office365.com",
        "protection.outlook.com",
        "bill.com",
        "melio.com",
        "invoiced.com",
        "stripe.com",
        "paypal.com",
        "intuit.com",
        "quickbooks.com",
        "mailchimp.com",
        "sendgrid.net",
        "amazonses.com",
        "sbcglobal.net",
        "att.net",
        "verizon.net",
    }
)

INTERNAL_DOMAINS = frozenset({"kannonmfg.com"})

SYSTEM_FROM_RE = re.compile(
    r"\b(postmaster|mailer-?daemon|microsoft\s*exchange|undeliverable|delivery\s+status|"
    r"accepted:\s*change\s+password|security\s+alert)\b",
    flags=re.I,
)
NOISE_FROM_RE = re.compile(
    r"^(no-?reply|do\s*not\s*reply|donotreply|invoices?|billing|accounts?\s+"
    r"(payable|receivable)|ar(\s+mailer)?|ar\s+department|ap\s+department|"
    r"customer\s+service|accounting|credit(\s+department)?|parts|pocketbook|"
    r"auto-receipt|the\s+efax\s+team)$",
    flags=re.I,
)
VIA_PLATFORM_RE = re.compile(r"\s+via\s+\S+$", flags=re.I)
SUBJECT_AS_VENDOR_RE = re.compile(
    r"^(re:|fw:|fwd:|automatic reply|your\s|please\s|new\s|notice\s|reminder|"
    r"ap\s+(dry\s+)?run|invoice\s+\d|invoices?\s+from|payment\s+confirm|"
    r"\*{2,}|\d{4}\s|open invoices|message from|transaction receipt|"
    r"have you |how |keep your |see all |set |share |simplify |stay |"
    r"try |unlock |upgrade |save |ready when |available:|connect |"
    r"changes?\s|end of |find the |losing |meet our |more ice|prime |"
    r"price drops|certificate)",
    flags=re.I,
)
FROM_COMPANY_RE = re.compile(
    r"\b(?:invoice|invoices|nvoice|payment request|ebill|receipt|reminder)"
    r".{0,60}?\bfrom\s+(.+?)\s*(?:[-–—|:.]|$|\binvoice\b|\bfor\b|\(#)",
    flags=re.I,
)
PAYMENT_TO_RE = re.compile(
    r"\b(?:your\s+payment\s+to|has sent you an invoice:?)\s+(.+?)\s*"
    r"(?:is\b|has\b|[-–—|:.]|$)",
    flags=re.I,
)
KANNON_CUSTOMER_RE = re.compile(r"\bkannon(\s+mfg|\s+manufacturing|\s+menufacturing)?\b", flags=re.I)
MARKETING_SUBJECT_RE = re.compile(
    r"(fort worth business|autopay|business prime|prime big deal|"
    r"upgrade worth|artisant|water cooler|auction|newsletter|"
    r"free sample|save up to|email exclusive|webinar|employment laws)",
    flags=re.I,
)
SKIP_MARKETING_DOMAINS = frozenset(
    {
        "resellcnc.com",
        "workwisecompliance.com",
        "secturafab.com",
        "highradius.com",
    }
)

# Live-inbox domains not yet in DOMAIN_VENDORS. Catalog-only; does not change bill parse.
EXTRA_DOMAIN_VENDORS = {
    "3pindustries.com": "3P",
    "3p.com": "3P",
    "precisionfabsvs.com": "Precision Fabrication Services",
    "precisionfabrication.com": "Precision Fabrication Services",
    "beshertsteel.com": "Beshert Steel Processing",
    "easternmetalsupply.com": "Eastern Metal Supply of Texas",
    "easternmetal.com": "Eastern Metal Supply of Texas",
    "shoppas.com": "Shoppa's Material Handling",
    "technitoolinc.com": "Techni-Tool",
    "priority1.com": "Priority 1",
    "priority1inc.com": "Priority 1",
    "recur360.com": "PCT Support",
    "readyrefresh.com": "Primo Brands",
    "toyota.com": "Toyota Commercial Finance",
    "ticf.com": "Toyota Commercial Finance",
    "tpcdm.com": "NTTA",
    "rivercitysteelco.com": "River City Steel",
    "venturisupply.com": "Venturi Supply",
    "exalloys.com": "Exotic Metals",
    "higginbotham.com": "IPFS",
    "meau.com": "MEAU",
    "tpitexas.com": "Telecom Products Inc.",
    "sss-steel.com": "Beshert Steel Processing",
    "capitalmachine.com": "Capital Machine Technologies, Inc",
    "fabcorp.com": "Fabcorp",
    "hagensfasteners.com": "Hagens Fasteners",
    "ktgalvanizing.com": "K-T Galvanizing",
    "engrcomp.com": "Engineered Components",
    "thyssenkrupp-materials.com": "Online Metals",
    "onlinemetals.com": "Online Metals",
    "kimcoerp.com": "KIMCO",
    "houstonplating.com": "Houston Plating",
    "amcastle.com": "A.M. Castle & Co.",
    "weckbrodt.de": "Weckbrodt",
    "arrowpersonnel.com": "Arrow Personnel",
    "arrowplating.com": "Arrow Plating",
    "coloniallife.com": "Colonial Life",
    "quenchusa.com": "Culligan Quench",
    "quench.com": "Culligan Quench",
    "culliganquench.com": "Culligan Quench",
    "phoenixmetals.com": "Phoenix Metals",
    "rolledalloys.com": "Rolled Alloys Inc",
    "tricormetals.com": "Tricor Metals",
    "ntta.org": "NTTA",
    "ipfs.com": "IPFS",
    "alarm-billing.com": "Alarm Billing",
    "amazon.com": "Amazon",
    "amazonbusiness.com": "Amazon",
    "graybar.com": "Graybar",
    "primobrands.com": "Primo Brands",
    "efax.com": "eFax",
    "vistaprint.com": "VistaPrint",
    "safety-kleen.com": "Safety-Kleen",
    "safetykleen.com": "Safety-Kleen",
    "ally.com": "Ally Auto",
    "spectrum.com": "Spectrum Business",
    "spectrumbusiness.com": "Spectrum Business",
    "dropbox.com": "Dropbox",
    "globelife.com": "Globe Life",
    "freepoint.com": "Freepoint Energy Solutions",
    "freepointenergy.com": "Freepoint Energy Solutions",
    "avexinstallations.com": "AVEX Installations LLC",
    "abybenefits.com": "ABY Benefits LLC",
    "tracemetalindustries.com": "Trace Metal Industries, Inc",
    "ldindustrialsolutions.com": "LD Industrial Solutions, LLC",
    "steelinspectors.com": "Steel Inspectors of Texas, Inc",
    "guerreroplating.com": "Guerrero Plating Technology, LLC",
    "greentreepackaging.com": "Greentree Packaging & Lumber",
    "metro-sprocket.com": "Metro Sprocket And Gear Inc",
    "metrosprocket.com": "Metro Sprocket And Gear Inc",
    "nationalbolt.com": "National Bolt and Ind Supply Co Inc.",
    "premiumalloys.com": "Premium Alloys",
    "productionmetals.com": "Production Metals",
    "stellasource.com": "Stella Source, Inc",
    "pittsburgsteel.com": "Pittsburgh Steel",
    "centralexpandedmetal.com": "Central Expanded Metal",
    "cofw.org": "City of Fort Worth",
    "fortworthtexas.gov": "City of Fort Worth",
    "thermofluids.com": "Thermo Fluids",
    "tolomatic.com": "Tolomatic",
    "syspro.com": "Syspro",
}

# Extra spellings that vendor_from_context / names_match would otherwise split.
EXPLICIT_ALIASES = {
    "mcmaster": "McMaster-Carr Supply Company",
    "mcmaster carr": "McMaster-Carr Supply Company",
    "mcmastercarr": "McMaster-Carr Supply Company",
    "gas supply": "Gas and Supply North Texas, LLC",
    "gas and supply": "Gas and Supply North Texas, LLC",
    "o neal": "O'Neal Steel - Dallas (GP)",
    "oneal": "O'Neal Steel - Dallas (GP)",
    "oneal steel": "O'Neal Steel - Dallas (GP)",
    "emj": "Earle M. Jorgensen Co",
    "earle m jorgensen": "Earle M. Jorgensen Co",
    "aqpc": "American Quality Powder Coating",
    "american quality powdercoating": "American Quality Powder Coating",
    "unifirst first aid": "UniFirst First Aid & Safety",
    "unifirst firstaid": "UniFirst First Aid & Safety",
    "amazon.com": "Amazon",
    "amazon business": "Amazon",
    "tpi": "Telecom Products Inc.",
    "telecom products": "Telecom Products Inc.",
    "capitalmachine": "Capital Machine Technologies, Inc",
    "capital machine": "Capital Machine Technologies, Inc",
    "phoenixmetals": "Phoenix Metals",
    "phoenix metals credit memos": "Phoenix Metals",
    "wasteconnections": "Waste Connections Lone Star, Inc",
    "coloniallife": "Colonial Life",
    "culligan quench": "Culligan Quench",
    "culliganquench": "Culligan Quench",
    "quench usa": "Culligan Quench",
    "arrowpersonnel": "Arrow Personnel",
    "arrowplating": "Arrow Plating",
    "houstonplating": "Houston Plating",
    "ktgalvanizing": "K-T Galvanizing",
    "k t galvanizing": "K-T Galvanizing",
    "cofw": "City of Fort Worth",
    "ipfs": "IPFS",
    "ntta": "NTTA",
    "tricormetals": "Tricor Metals",
    "pittsburgsteel": "Pittsburgh Steel",
    "centralexpandedmetal": "Central Expanded Metal",
    "vista print": "VistaPrint",
    "vistaprint email exclusive": "VistaPrint",
    "the efax team": "eFax",
    "efax team": "eFax",
    "primo brands customer experience": "Primo Brands",
    "freepoint solution customer relations": "Freepoint Energy Solutions",
    "freepointsolutions": "Freepoint Energy Solutions",
    "freepoint solutions": "Freepoint Energy Solutions",
    "sss steel": "Beshert Steel Processing",
    "triple s steel": "Beshert Steel Processing",
    "am castle": "A.M. Castle & Co.",
    "a m castle": "A.M. Castle & Co.",
    "online metals": "Online Metals",
    "thyssenkrupp": "Online Metals",
    "spectrum business": "Spectrum Business",
    "globe life": "Globe Life",
    "globe": "Globe Life",
    "kimcoerp": "KIMCO",
    "kimco accounting": "KIMCO",
    "kimco": "KIMCO",
    "3pindustries": "3P",
    "3p industries": "3P",
    "beshertsteel": "Beshert Steel Processing",
    "easternmetal": "Eastern Metal Supply of Texas",
    "shoppas": "Shoppa's Material Handling",
    "technitoolinc": "Techni-Tool",
    "priority1": "Priority 1",
    "readyrefresh": "Primo Brands",
    "ticf": "Toyota Commercial Finance",
    "tpcdm": "NTTA",
    "spectrumemails": "Spectrum Business",
    "precisionfabsvs": "Precision Fabrication Services",
    "rivercitysteelco": "River City Steel",
    "venturisupply": "Venturi Supply",
    "recur360": "PCT Support",
    "melody channell": "Precision Fabrication Services",
}


def known_canonical_vendors() -> list[str]:
    """Stable unique names from DOMAIN_VENDORS + SUBJECT_VENDORS + inbox extras."""
    seen: set[str] = set()
    out: list[str] = []
    extra_names = list(EXTRA_DOMAIN_VENDORS.values()) + list(EXPLICIT_ALIASES.values())
    for name in list(DOMAIN_VENDORS.values()) + [vendor for _, vendor in SUBJECT_VENDORS] + extra_names:
        if name not in seen:
            seen.add(name)
            out.append(name)
    return out


def _compact(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value or "")


def _alias_index() -> dict[str, str]:
    """normalize_name / unique-compact → canonical. Colliding compact keys dropped."""
    index: dict[str, str] = dict(EXPLICIT_ALIASES)
    compact_hits: dict[str, set[str]] = defaultdict(set)
    for vendor in known_canonical_vendors():
        norm = normalize_name(vendor)
        if norm and norm not in index:
            index[norm] = vendor
        compact_hits[_compact(norm)].add(vendor)
    for compact, vendors in compact_hits.items():
        if compact and len(vendors) == 1:
            index.setdefault(compact, next(iter(vendors)))
    return index


_ALIAS_INDEX = _alias_index()
_KNOWN = known_canonical_vendors()


def email_domain(address: str) -> str:
    addr = (address or "").strip().lower()
    if "@" not in addr:
        return ""
    return addr.split("@", 1)[1]


def parent_domains(domain: str) -> list[str]:
    parts = [p for p in (domain or "").split(".") if p]
    out = []
    for index in range(len(parts) - 1):
        out.append(".".join(parts[index:]))
    return out


def domain_looks_internal(address: str) -> bool:
    domain = email_domain(address)
    return any(part in INTERNAL_DOMAINS for part in parent_domains(domain)) or domain in INTERNAL_DOMAINS


def domain_is_platform(address: str) -> bool:
    domain = email_domain(address)
    return any(part in PLATFORM_DOMAINS for part in parent_domains(domain)) or domain in PLATFORM_DOMAINS


def domain_vendor(address: str) -> str:
    addr = (address or "").lower()
    if "firstaid" in addr:
        return "UniFirst First Aid & Safety"
    domain = email_domain(address)
    for part in [domain, *parent_domains(domain)]:
        if part in EXTRA_DOMAIN_VENDORS:
            return EXTRA_DOMAIN_VENDORS[part]
        if part in DOMAIN_VENDORS:
            return DOMAIN_VENDORS[part]
        if part in SKIP_MARKETING_DOMAINS:
            return ""
    return ""


def clean_from_display(name: str) -> str:
    cleaned = VIA_PLATFORM_RE.sub("", name or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" -–—|:")
    return cleaned[:80]


def looks_like_subject_line(name: str) -> bool:
    text = (name or "").strip()
    if not text:
        return False
    if SUBJECT_AS_VENDOR_RE.search(text):
        return True
    if re.search(r"[?!]|[\U0001F300-\U0001FAFF]", text):
        return True
    if len(text) > 48 and text not in _KNOWN:
        return True
    words = text.split()
    if len(words) >= 8:
        return True
    if text[:1].isdigit() and not distinctive_vendor_tokens(text):
        return True
    return False


def looks_like_person_name(name: str) -> bool:
    """Person From names, including lowercase 'abel jasso'."""
    cleaned = clean_from_display(name)
    if _looks_like_person_name(cleaned):
        return True
    parts = [p for p in re.split(r"\s+", cleaned) if p]
    if 2 <= len(parts) <= 4 and all(p.isalpha() for p in parts) and not _COMPANY_WORD_RE.search(cleaned):
        return True
    return False


def looks_like_role_mailbox(name: str) -> bool:
    return bool(NOISE_FROM_RE.match(clean_from_display(name)))


def extract_company_from_subject(subject: str) -> str:
    legal = company_from_subject_or_text(subject=subject or "", text="")
    if legal:
        return legal[:80]
    for rx in (FROM_COMPANY_RE, PAYMENT_TO_RE):
        match = rx.search(subject or "")
        if not match:
            continue
        raw = match.group(1).strip(" -–—|:.,")
        raw = re.sub(r"\s+\((?:#?\d+).*$", "", raw).strip()
        if raw.lower().startswith("from "):
            raw = raw[5:].strip()
        if raw and not looks_like_subject_line(raw) and not KANNON_CUSTOMER_RE.search(raw):
            return raw[:80]
    headed = re.match(
        r"^\s*([A-Za-z][A-Za-z0-9&.'/+\s]{2,50}?)\s*[-–—|:]\s*(?:Invoice|Inv\.?|Sales Invoice|eBill)\b",
        subject or "",
        flags=re.I,
    )
    if headed:
        raw = headed.group(1).strip()
        if raw and not _looks_like_person_name(raw) and not looks_like_role_mailbox(raw):
            return raw[:80]
    return ""


def guess_vendor(*, subject: str = "", from_name: str = "", from_address: str = "") -> str:
    """Cheap vendor guess: domain / subject company / From. No PDF download."""
    mapped = domain_vendor(from_address)
    if mapped:
        return mapped
    extracted = extract_company_from_subject(subject)
    if extracted:
        return extracted
    if re.search(r"\bipfs\b", subject or "", flags=re.I):
        return "IPFS"
    display = clean_from_display(from_name)
    guessed = vendor_from_context(
        subject=subject or "",
        from_name=display,
        from_address=from_address or "",
        text="",
    )
    guessed = (guessed or "").strip()
    if guessed and not looks_like_subject_line(guessed) and not KANNON_CUSTOMER_RE.search(guessed):
        if (
            guessed in _KNOWN
            or not looks_like_person_name(guessed)
            or guessed.lower() in {"melody channell", "rachel bailey"}
        ):
            if guessed.lower() == "rachel bailey":
                return "3P"
            if guessed.lower() == "melody channell":
                return "Precision Fabrication Services"
            return guessed
    if looks_like_role_mailbox(display) or looks_like_person_name(display):
        return ""
    if display and not looks_like_subject_line(display) and not KANNON_CUSTOMER_RE.search(display):
        return display
    return ""


def canonicalize_vendor(name: str) -> str:
    """Collapse McMaster-Carr / McMaster Carr / mcmaster.com to one row."""
    raw = clean_from_display(name)
    if not raw:
        return ""
    norm = normalize_name(raw)
    if norm in _ALIAS_INDEX:
        return _ALIAS_INDEX[norm]
    compact = _compact(norm)
    if compact in _ALIAS_INDEX:
        return _ALIAS_INDEX[compact]
    best = _best_known_match(raw)
    if best:
        return best
    if raw.isupper() and len(raw) > 3:
        return _title_company(raw)
    return raw[:80]


def _best_known_match(name: str) -> str:
    """Map an unknown spelling onto a listed canonical without merging UniFirst pair / MSC↔RMP."""
    candidates: list[tuple[int, int, str]] = []
    name_dist = distinctive_vendor_tokens(name)
    name_norm = normalize_name(name)
    if not name_norm:
        return ""
    for canon in _KNOWN:
        canon_norm = normalize_name(canon)
        if name_norm == canon_norm or _compact(name_norm) == _compact(canon_norm):
            return canon
        canon_dist = distinctive_vendor_tokens(canon)
        if not name_dist or not canon_dist:
            continue
        overlap = set(name_dist) & set(canon_dist)
        shorter = name_dist if len(name_dist) <= len(canon_dist) else canon_dist
        contained = bool(shorter) and set(shorter) <= overlap
        if contained or len(overlap) >= 2:
            candidates.append((len(overlap), len(canon_dist), canon))
    if not candidates:
        return ""
    candidates.sort(key=lambda item: (item[0], item[1], item[2]), reverse=True)
    return candidates[0][2]


def vendors_same_row(left: str, right: str) -> bool:
    """True when two guessed names are the same company for the catalog."""
    a = canonicalize_vendor(left)
    b = canonicalize_vendor(right)
    if not a or not b:
        return False
    if a == b:
        return True
    return normalize_name(a) == normalize_name(b) or _compact(normalize_name(a)) == _compact(normalize_name(b))


def message_received(message: dict[str, Any]) -> datetime | None:
    raw = str(message.get("receivedDateTime") or "")
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return None


def _is_weak_vendor(name: str, subject: str) -> bool:
    cleaned = (name or "").strip()
    if not cleaned:
        return True
    if looks_like_person_name(cleaned):
        return True
    if NOISE_FROM_RE.match(cleaned):
        return True
    if looks_like_subject_line(cleaned):
        return True
    if KANNON_CUSTOMER_RE.search(cleaned):
        return True
    if normalize_name(cleaned) == normalize_name(subject or "") and not distinctive_vendor_tokens(cleaned):
        return True
    if normalize_name(cleaned) in {"invoice", "invoices", "billing", "statement", "past due"}:
        return True
    return False


def classify_inbox_item(
    message: dict[str, Any],
) -> tuple[str, str, str]:
    """Return (kind, vendor_or_bucket, reason).

    kind is ``vendor`` or ``skip``.
    """
    subject = str(message.get("subject") or "")
    preview = str(message.get("bodyPreview") or "")
    from_name = sender_name(message)
    from_addr = sender_address(message)
    klass = classify_mail(subject=subject, preview=preview, from_name=from_name)

    if from_addr == ALLOWED_MAILBOX:
        return "skip", "sent-from-ap", "outbound / sent-item copy"
    if SYSTEM_FROM_RE.search(f"{from_name} {subject}"):
        return "skip", "system", "system / bounce / exchange"
    if klass == "internal":
        return "skip", "internal", "internal"

    guessed = canonicalize_vendor(guess_vendor(subject=subject, from_name=from_name, from_address=from_addr))
    known = bool(guessed and guessed in _KNOWN)
    marketing_domain = any(
        part in SKIP_MARKETING_DOMAINS for part in parent_domains(email_domain(from_addr))
    ) or email_domain(from_addr) in SKIP_MARKETING_DOMAINS

    if KANNON_CUSTOMER_RE.search(guessed) and not known:
        return "skip", "internal", "Kannon is the customer, not a vendor"
    if marketing_domain and not subject_has_invoice_bill_hint(subject):
        return "skip", "marketing", "auction / newsletter / collections platform"
    if re.search(
        r"resell\s*cnc|workwise|secturafab|dropbox|syspro|tolomatic|vistaprint|efax",
        f"{from_name} {from_addr} {subject}",
        flags=re.I,
    ) and not subject_has_invoice_bill_hint(subject):
        return "skip", "marketing", "auction / newsletter / promo"
    if re.search(r"certificate\(s\) of insurance|verification code|please update our address|share your feedback", subject, flags=re.I):
        return "skip", "not-a-bill", "insurance cert / verification / address change / survey"
    if guessed and re.fullmatch(r"[A-Za-z]{2,4}", guessed) and guessed not in _KNOWN:
        return "skip", "weak-vendor", "short unmapped From domain"
    if MARKETING_SUBJECT_RE.search(subject) and not known and not subject_has_invoice_bill_hint(subject):
        return "skip", "marketing", "promo / newsletter"
    if looks_like_subject_line(guessed):
        return "skip", "unmapped-sender", "subject used as vendor name"

    if domain_looks_internal(from_addr) and not known and not subject_has_invoice_bill_hint(subject):
        return "skip", "internal", "kannonmfg.com without a mapped vendor"

    if klass in SKIP_CLASSES and klass != "internal":
        # Invoice/INV subjects still catalog even when classify_mail says statement
        # (Greentree-style false statement). Plain account statements stay skip.
        if klass == "statement" and subject_has_invoice_bill_hint(subject):
            pass
        else:
            return "skip", klass, klass

    if klass == "auto-pay":
        return "vendor", guessed or "Toyota Commercial Finance", "auto-pay"

    if _is_weak_vendor(guessed, subject):
        if known:
            return "vendor", guessed, klass or "invoice"
        if domain_is_platform(from_addr) or domain_looks_internal(from_addr):
            return "skip", "unmapped-sender", "weak From / platform / internal"
        domain = email_domain(from_addr)
        token = ""
        if domain and not domain_is_platform(from_addr):
            parts = domain.split(".")
            token = parts[-2] if len(parts) >= 2 else domain
            if token and token not in {"com", "net", "org", "edu", "mail", "email", "invoices", "billing"} and not re.fullmatch(r"[A-Za-z]{2,4}", token):
                return "vendor", canonicalize_vendor(token.replace("-", " ").title()), "domain-token"
        if subject_has_invoice_bill_hint(subject):
            return "skip", "unmapped-invoice", "invoice-looking but no vendor name"
        return "skip", "weak-vendor", "person name / empty vendor"

    return "vendor", guessed, klass or "invoice"


@dataclass
class VendorAgg:
    vendor: str
    aliases: set[str] = field(default_factory=set)
    count_window: int = 0
    count_older: int = 0
    sample_from: str = ""
    sample_subject: str = ""
    first_seen: str = ""
    last_seen: str = ""

    def add(self, *, alias: str, received: datetime | None, from_name: str, from_addr: str, subject: str, in_window: bool) -> None:
        if alias and alias != self.vendor:
            self.aliases.add(alias)
        if in_window:
            self.count_window += 1
        else:
            self.count_older += 1
        stamp = received.isoformat() if received else ""
        if stamp and (not self.first_seen or stamp < self.first_seen):
            self.first_seen = stamp
        if stamp and stamp > self.last_seen:
            self.last_seen = stamp
        if not self.sample_subject:
            who = from_name or from_addr
            self.sample_from = who
            self.sample_subject = subject


@dataclass
class SkipAgg:
    bucket: str
    count: int = 0
    sample: str = ""
    reason: str = ""

    def add(self, *, from_name: str, from_addr: str, subject: str, reason: str) -> None:
        self.count += 1
        if not self.sample:
            who = from_name or from_addr
            self.sample = f"{who} | {subject}"[:200]
        if reason and not self.reason:
            self.reason = reason


@dataclass
class VendorCatalog:
    vendors: dict[str, VendorAgg] = field(default_factory=dict)
    skips: dict[str, SkipAgg] = field(default_factory=dict)
    messages_scanned: int = 0
    invoice_emails: int = 0
    skip_emails: int = 0
    scanned_from: str = ""
    scanned_to: str = ""
    count_window_start: str = COUNT_WINDOW_START.isoformat()
    mailbox: str = ALLOWED_MAILBOX

    def unique_vendor_names(self) -> list[str]:
        return sorted(self.vendors, key=str.casefold)


def catalog_messages(
    messages: list[dict[str, Any]],
    *,
    window_start: date = COUNT_WINDOW_START,
) -> VendorCatalog:
    """Deduplicate invoice-looking mail into canonical vendor rows + skip buckets."""
    catalog = VendorCatalog(count_window_start=window_start.isoformat())
    by_key: dict[str, str] = {}

    def _row_key(name: str) -> str:
        canon = canonicalize_vendor(name)
        norm = normalize_name(canon)
        return _compact(norm) or norm or canon.casefold()

    for message in messages:
        catalog.messages_scanned += 1
        received = message_received(message)
        if received:
            stamp = received.isoformat()
            if not catalog.scanned_from or stamp < catalog.scanned_from:
                catalog.scanned_from = stamp
            if stamp > catalog.scanned_to:
                catalog.scanned_to = stamp
        in_window = True
        if received is not None:
            in_window = received.astimezone(timezone.utc).date() >= window_start

        subject = str(message.get("subject") or "")
        from_name = sender_name(message)
        from_addr = sender_address(message)
        raw_guess = guess_vendor(subject=subject, from_name=from_name, from_address=from_addr)
        kind, label, reason = classify_inbox_item(message)

        if kind == "skip":
            catalog.skip_emails += 1
            bucket = catalog.skips.get(label)
            if bucket is None:
                bucket = SkipAgg(bucket=label, reason=reason)
                catalog.skips[label] = bucket
            bucket.add(from_name=from_name, from_addr=from_addr, subject=subject, reason=reason)
            continue

        catalog.invoice_emails += 1
        canon = canonicalize_vendor(label or raw_guess)
        if not canon:
            catalog.skip_emails += 1
            bucket = catalog.skips.setdefault("unmapped-invoice", SkipAgg(bucket="unmapped-invoice", reason="empty vendor"))
            bucket.add(from_name=from_name, from_addr=from_addr, subject=subject, reason="empty vendor")
            continue
        key = _row_key(canon)
        existing_name = by_key.get(key)
        if existing_name is None:
            # Merge into an existing row when compact/normalize already matches.
            for other_key, other_name in list(by_key.items()):
                if vendors_same_row(canon, other_name):
                    existing_name = other_name
                    by_key[key] = other_name
                    break
        if existing_name is None:
            catalog.vendors[canon] = VendorAgg(vendor=canon)
            by_key[key] = canon
            existing_name = canon
        row = catalog.vendors[existing_name]
        for alias in {raw_guess, from_name, label}:
            cleaned = clean_from_display(alias)
            if (
                cleaned
                and cleaned != row.vendor
                and not looks_like_person_name(cleaned)
                and not NOISE_FROM_RE.match(cleaned)
            ):
                row.aliases.add(cleaned)
        row.add(
            alias="",
            received=received,
            from_name=from_name,
            from_addr=from_addr,
            subject=subject,
            in_window=in_window,
        )
    return catalog


def sample_from_subject(row: VendorAgg) -> str:
    who = row.sample_from or ""
    subject = row.sample_subject or ""
    if who and subject:
        return f"{who} | {subject}"[:240]
    return (who or subject)[:240]


def count_label(row: VendorAgg, *, window_start: date = COUNT_WINDOW_START) -> str:
    if row.count_window and not row.count_older:
        return str(row.count_window)
    if row.count_window and row.count_older:
        return f"{row.count_window} (plus {row.count_older} before {window_start.isoformat()})"
    if row.count_older:
        return f"0 ({row.count_older} before {window_start.isoformat()} only)"
    return "0"


def vendor_sheet_rows(catalog: VendorCatalog) -> list[dict[str, Any]]:
    rows = []
    for name in catalog.unique_vendor_names():
        row = catalog.vendors[name]
        aliases = sorted(
            {a for a in row.aliases if a and normalize_name(a) != normalize_name(row.vendor)},
            key=str.casefold,
        )
        rows.append(
            {
                "Vendor": row.vendor,
                "Aliases seen": "; ".join(aliases),
                "Approx invoice count": count_label(row),
                "Sample from / subject": sample_from_subject(row),
                "Receiving owner": "",
                "Notes": "",
            }
        )
    return rows


def skip_sheet_rows(catalog: VendorCatalog) -> list[dict[str, Any]]:
    rows = []
    for key in sorted(catalog.skips, key=str.casefold):
        item = catalog.skips[key]
        rows.append(
            {
                "Bucket": item.bucket,
                "Count": item.count,
                "Sample from / subject": item.sample,
                "Notes": item.reason,
            }
        )
    return rows


def catalog_to_json(catalog: VendorCatalog) -> dict[str, Any]:
    return {
        "mailbox": catalog.mailbox,
        "scanned_from": catalog.scanned_from,
        "scanned_to": catalog.scanned_to,
        "count_window_start": catalog.count_window_start,
        "messages_scanned": catalog.messages_scanned,
        "invoice_emails": catalog.invoice_emails,
        "skip_emails": catalog.skip_emails,
        "unique_vendors": len(catalog.vendors),
        "unique_vendor_names": catalog.unique_vendor_names(),
        "vendors": [
            {
                "vendor": catalog.vendors[name].vendor,
                "aliases": sorted(catalog.vendors[name].aliases, key=str.casefold),
                "count_aug1": catalog.vendors[name].count_window,
                "count_older": catalog.vendors[name].count_older,
                "sample_from": catalog.vendors[name].sample_from,
                "sample_subject": catalog.vendors[name].sample_subject,
                "first_seen": catalog.vendors[name].first_seen,
                "last_seen": catalog.vendors[name].last_seen,
            }
            for name in catalog.unique_vendor_names()
        ],
        "skip": skip_sheet_rows(catalog),
        "caveat": (
            "Vendor names come from From display, From domain, and subject "
            "(DOMAIN_VENDORS / SUBJECT_VENDORS / vendor_from_context). "
            "PDF-printed vendor was not downloaded for this catalog — a "
            "forwarded invoice can show the forwarder or a billing platform "
            "instead of the remitting vendor."
        ),
    }


def write_vendor_workbook(path: Path, catalog: VendorCatalog) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    vendors = workbook.active
    vendors.title = "Vendors"
    for col, name in enumerate(VENDOR_COLUMNS, start=1):
        cell = vendors.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)
    for row_idx, row in enumerate(vendor_sheet_rows(catalog), start=2):
        for col, key in enumerate(VENDOR_COLUMNS, start=1):
            cell = vendors.cell(row_idx, col, row.get(key, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate([36, 40, 28, 70, 22, 28], start=1):
        vendors.column_dimensions[get_column_letter(idx)].width = width
    vendors.auto_filter.ref = f"A1:{get_column_letter(len(VENDOR_COLUMNS))}{max(1, len(catalog.vendors) + 1)}"
    vendors.freeze_panes = "A2"
    vendors.row_dimensions[1].height = 22

    names = workbook.create_sheet("Unique names")
    names.cell(1, 1, "Vendor").font = header_font
    names.cell(1, 1).fill = header_fill
    for row_idx, name in enumerate(catalog.unique_vendor_names(), start=2):
        names.cell(row_idx, 1, name)
    names.column_dimensions["A"].width = 44
    names.freeze_panes = "A2"

    skip = workbook.create_sheet("Skip")
    for col, name in enumerate(SKIP_COLUMNS, start=1):
        cell = skip.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for row_idx, row in enumerate(skip_sheet_rows(catalog), start=2):
        for col, key in enumerate(SKIP_COLUMNS, start=1):
            cell = skip.cell(row_idx, col, row.get(key, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate([22, 10, 70, 36], start=1):
        skip.column_dimensions[get_column_letter(idx)].width = width
    skip.freeze_panes = "A2"

    meta = workbook.create_sheet("Scan")
    meta.cell(1, 1, "Field").font = header_font
    meta.cell(1, 1).fill = header_fill
    meta.cell(1, 2, "Value").font = header_font
    meta.cell(1, 2).fill = header_fill
    payload = catalog_to_json(catalog)
    facts = [
        ("Mailbox", payload["mailbox"]),
        ("Scanned from", payload["scanned_from"]),
        ("Scanned to", payload["scanned_to"]),
        ("Count window start", payload["count_window_start"]),
        ("Messages scanned", payload["messages_scanned"]),
        ("Invoice-looking emails", payload["invoice_emails"]),
        ("Skip / noise emails", payload["skip_emails"]),
        ("Unique vendors", payload["unique_vendors"]),
        ("Caveat", payload["caveat"]),
    ]
    for row_idx, (key, value) in enumerate(facts, start=2):
        meta.cell(row_idx, 1, key)
        meta.cell(row_idx, 2, value)
        meta.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
    meta.column_dimensions["A"].width = 28
    meta.column_dimensions["B"].width = 100

    workbook.save(path)
    return path


def default_sheet_path(day: date | None = None) -> Path:
    stamp = (day or date.today()).isoformat()
    return Path("runs") / f"AP-vendors-from-inbox-{stamp}.xlsx"
