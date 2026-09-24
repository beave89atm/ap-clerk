"""Build the 3-column Gas HOLD merch-lines sheet for Shawn.

Read-only. invent=false. No KIMCO writes. No Mail.Send.
Each row is one PDF merch line for that invoice # only.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ap_clerk.pdf_invoice import parse_invoice_pdf
from gas_supply_0917 import exact_invoice_number, merchandise_lines_from_parsed

HOLD_INVOICES: tuple[tuple[str, str], ...] = (
    ("0040430010", "59081"),
    ("0040424839", "59081"),
    ("0040423658", "58948"),
    ("0040438057", "59081"),
    ("0040438056", "59081"),
    ("0040438055", "59081"),
    ("0040438053", "59081"),
    ("0040417672", "59081"),
    ("0040414962", "59006"),
    ("0040414821", "58948"),
)

# Prefer the unsplit pack so multi-page invoices (0040417672) stay together.
# parse_invoice_pdf splits siblings; we keep THIS invoice # only.
PREFERRED_PDF: dict[str, str] = {
    "0040430010": "gas_2026-09-12_billing01_A3050_c.pdf",
    "0040424839": "gas_2026-09-10_billing01_A3050_c.pdf",
    "0040423658": "gas_2026-09-10_billing01_A3050_c.pdf",
    "0040438057": "gas_2026-09-18_billing01_A3050_c.pdf",
    "0040438056": "gas_2026-09-18_billing01_A3050_c.pdf",
    "0040438055": "gas_2026-09-18_billing01_A3050_c.pdf",
    "0040438053": "gas_2026-09-18_billing01_A3050_c.pdf",
    "0040417672": "gas_2026-09-04_billing01_A3050_c.pdf",
    "0040414962": "gas_2026-09-03_billing01_A3050_c.pdf",
    "0040414821": "gas_2026-09-03_billing01_A3050_c.pdf",
}

PDF_DIR = ROOT / "runs" / "inbox-pdfs"
OUT_XLSX = ROOT / "runs" / "Gas-Supply-HOLD-lines-for-Shawn-2026-09-18.xlsx"
COLUMNS = ("INV#", "line item", "PO#")


def _money(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return ""
    if abs(number - round(number, 2)) < 1e-9:
        return f"{number:.2f}"
    text = f"{number:.4f}".rstrip("0").rstrip(".")
    return text


def format_hold_line_item(line: dict[str, Any]) -> str:
    """Clear readable merch string: part, description, qty, unit cost."""
    part = str(line.get("part") or "").strip()
    desc = str(line.get("description") or line.get("label") or "").strip()
    if part and desc and part.lower() not in desc.lower():
        name = f"{part} {desc}"
    else:
        name = desc or part
    name = " ".join(name.split())
    qty = line.get("qty")
    unit = line.get("unit_price") if line.get("unit_price") is not None else line.get("unit")
    bits = [name] if name else []
    if qty not in (None, ""):
        bits.append(f"qty {_money(qty).rstrip('0').rstrip('.') if float(qty) == int(float(qty)) else _money(qty)}")
    if unit not in (None, ""):
        bits.append(f"@ {_money(unit)}")
    return " ".join(bits) if bits else "(no merch description on PDF)"


def _bills_from_pdf(pdf_path: Path) -> list[dict[str, Any]]:
    parsed = parse_invoice_pdf(
        pdf_path,
        subject="Gas&Supply Invoice/Statement",
        from_name="Gas and Supply",
        from_address="billing@gasandsupply.com",
    )
    extras = list(parsed.pop("siblings", []) or [])
    return [parsed, *extras]


def parse_hold_invoice(invoice_number: str, po: str) -> dict[str, Any]:
    preferred = PDF_DIR / PREFERRED_PDF[invoice_number]
    search = [preferred]
    search.extend(sorted(PDF_DIR.glob(f"*{invoice_number}*")))
    seen: set[str] = set()
    last_err = "no PDF candidate"
    for pdf in search:
        if not pdf.exists() or str(pdf) in seen:
            continue
        seen.add(str(pdf))
        for bill in _bills_from_pdf(pdf):
            if exact_invoice_number(bill.get("invoice_number")) != invoice_number:
                continue
            merch = merchandise_lines_from_parsed(bill)
            return {
                "invoice_number": invoice_number,
                "po": po,
                "pdf_path": str(pdf),
                "pdf_amount": bill.get("amount"),
                "pdf_po": bill.get("po"),
                "merch": merch,
                "fees": list(bill.get("fees") or []),
                "blocker": None,
            }
        last_err = f"PDF {pdf.name} had no section for {invoice_number} only"
    return {
        "invoice_number": invoice_number,
        "po": po,
        "pdf_path": None,
        "pdf_amount": None,
        "pdf_po": None,
        "merch": [],
        "fees": [],
        "blocker": last_err,
    }


def rows_from_parsed(parsed: dict[str, Any]) -> list[dict[str, str]]:
    inv = parsed["invoice_number"]
    po = parsed["po"]
    merch = list(parsed.get("merch") or [])
    if merch:
        return [
            {"INV#": inv, "line item": format_hold_line_item(line), "PO#": po}
            for line in merch
        ]
    fees = parsed.get("fees") or []
    if parsed.get("blocker"):
        note = f"(PDF merch lines could not be extracted: {parsed['blocker']})"
    elif fees:
        fee_bits = ", ".join(
            f"{f.get('name') or 'fee'} {f.get('amount')}" for f in fees
        )
        note = f"(no merch lines on this invoice PDF; fees only: {fee_bits})"
    else:
        note = "(no merch lines on this invoice PDF)"
    return [{"INV#": inv, "line item": note, "PO#": po}]


def write_xlsx(path: Path, rows: list[dict[str, str]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    sheet = wb.active
    sheet.title = "HOLD lines for Shawn"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for col, name in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for idx, row in enumerate(rows, start=2):
        for col, name in enumerate(COLUMNS, start=1):
            cell = sheet.cell(idx, col, row.get(name, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 78
    sheet.column_dimensions["C"].width = 12
    sheet.auto_filter.ref = f"A1:C{max(1, len(rows) + 1)}"
    sheet.freeze_panes = "A2"
    wb.save(path)
    return path


def main() -> int:
    parsed_holds = [parse_hold_invoice(inv, po) for inv, po in HOLD_INVOICES]
    rows: list[dict[str, str]] = []
    for parsed in parsed_holds:
        rows.extend(rows_from_parsed(parsed))
    write_xlsx(OUT_XLSX, rows)
    sidecar = {
        "invent": False,
        "mail_send": False,
        "kimco_writes": False,
        "columns": list(COLUMNS),
        "row_count": len(rows),
        "invoices": [inv for inv, _po in HOLD_INVOICES],
        "parsed": [
            {
                "invoice_number": p["invoice_number"],
                "po": p["po"],
                "pdf_path": p["pdf_path"],
                "pdf_amount": p["pdf_amount"],
                "merch_count": len(p.get("merch") or []),
                "blocker": p.get("blocker"),
                "line_items": [format_hold_line_item(ln) for ln in p.get("merch") or []],
            }
            for p in parsed_holds
        ],
        "rows": rows,
        "xlsx": str(OUT_XLSX),
    }
    json_path = OUT_XLSX.with_suffix(".json")
    json_path.write_text(json.dumps(sidecar, indent=2, default=str) + "\n")
    print(json.dumps({"xlsx": str(OUT_XLSX), "json": str(json_path), "row_count": len(rows)}, indent=2))
    blockers = [p for p in parsed_holds if p.get("blocker") or not p.get("merch")]
    for p in parsed_holds:
        print(
            f"{p['invoice_number']} po={p['po']} merch={len(p.get('merch') or [])} "
            f"blocker={p.get('blocker')}",
            flush=True,
        )
    return 2 if any(p.get("blocker") for p in parsed_holds) else 0


if __name__ == "__main__":
    raise SystemExit(main())
