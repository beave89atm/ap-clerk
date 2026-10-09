"""Enter the 2026-10-09 live AP run. Cap 25. Does not post, close, or email."""

from __future__ import annotations

import csv
import json
import logging
import re
import sys
from datetime import date
from pathlib import Path
from typing import Any

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from ap_clerk.cli import _find_or_create_batch
from ap_clerk.graph import (
    ALLOWED_MAILBOX,
    ENTERED_IN_AI_CATEGORY,
    LEGACY_AI_SKIPPED_CATEGORY,
    GraphClient,
    load_graph_credentials,
)
from ap_clerk.kimco import KimcoClient, KimcoError, added_comment_payload
from ap_clerk.rules import SHAWN_MENTION_HTML, ap_clerk_edit_note, lookup_id, lookup_text
from scripts.ap_run_2026_10_08_enter import (
    add_charges,
    add_misc,
    attach,
    cents,
    create_header as _create_header_unused,
    load_open_receipts,
    match_lines,
    move_batch,
    po_absence_reason,
    sample_for,
)
from scripts.sept25_30_email_move_execute import archive_folder, folder_path, sender_of, stamp
from scripts import sept25_30_finish_2026_10_06 as finish
from scripts.sept_archive_check_2026_10_07 import plain
from scripts.sept_missed_entry_2026_10_07 import put, totals

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
LOGGER = logging.getLogger("ap-run-1009")

OUT = ROOT / "runs" / "ap-run-2026-10-09"
QC = OUT / "qc"
SRC = Path("/tmp/ap-run-1009/splits")
BASE = Path("/tmp/ap-run-1009")
LISTING = OUT / "inbox-listing.json"
BATCH_NAME = "API Agent - 10/9/26"
TRANSFER = 375
CAP = 25
MIN_NEW_ID = 10566
SELECT = "id,subject,from,receivedDateTime,parentFolderId,categories,lastModifiedDateTime"

# Silence the unused import while keeping the 10/8 header helper available for comparison.
_ = _create_header_unused


def job(
    number: str,
    vendor: str,
    vendor_id: int,
    day: date,
    amount: float,
    received: str,
    pdf: str,
    *,
    kind: str = "po",
    po: str = "",
    po_id: int | None = None,
    merch: float | None = None,
    lines_match: list | None = None,
    lines: list | None = None,
    freight: float | None = None,
    fee: float | None = None,
    fee_name: str = "",
    owner: str = "treyce",
    hold: str = "",
    extra: str = "",
    pages: list[int] | None = None,
    source: str = "",
    piece_inch: bool = False,
) -> dict[str, Any]:
    row: dict[str, Any] = {
        "number": number,
        "vendor": vendor,
        "vendor_id": vendor_id,
        "day": day,
        "amount": amount,
        "received": received,
        "pdf": pdf,
        "kind": kind,
        "po": po,
        "merch": merch if merch is not None else amount,
        "owner": owner,
        "extra": extra,
        "source": source,
        "pages": pages,
    }
    if po_id:
        row["po_id"] = po_id
    if lines_match:
        row["lines_match"] = lines_match
    if lines:
        row["lines"] = lines
    if freight:
        row["freight"] = freight
    if fee:
        row["fee"] = fee
        row["fee_name"] = fee_name or "fee"
    if hold:
        row["hold_without_lines"] = hold
    if piece_inch:
        row["piece_inch"] = True
    return row


D = date
JOBS: list[dict[str, Any]] = [
    job("PS-INV104087", "Legacy Wire Products", 292, D(2026, 10, 5), 1523.45, "2026-10-05T14:26:53Z", "legacy-104087.pdf", po="59320", po_id=7322, merch=1271.00, freight=252.45, source="20261005T142653Z-1.pdf", hold="No receipt recorded yet on PO 59320 for 8 each and 23 each of 75-10-201007. The only open receipt is 20 each on a different line. Nothing was selected."),
    job("130153", "Morgan Steel", 304, D(2026, 10, 1), 6040.00, "2026-10-05T14:32:08Z", "morgan-130153.pdf", po="59013", po_id=7015, merch=6040.00, lines_match=[{"qty": 40, "amount": 6040.00, "uom": "EA"}], source="20261005T143208Z-1.pdf", owner="shawn"),
    job("130154", "Morgan Steel", 304, D(2026, 10, 1), 6040.00, "2026-10-05T14:32:12Z", "morgan-130154.pdf", po="59199", po_id=7201, merch=6040.00, lines_match=[{"qty": 40, "amount": 6040.00, "uom": "EA"}], source="20261005T143212Z-1.pdf", owner="shawn", piece_inch=True, extra="The invoice bills 40 pieces. The open receipt is 40 inches for the same $6,040.00, so that receipt was selected."),
    job("130152", "Morgan Steel", 304, D(2026, 10, 1), 5628.00, "2026-10-05T14:32:21Z", "morgan-130152.pdf", po="59012", po_id=7014, merch=5628.00, lines_match=[{"qty": 42, "amount": 5628.00, "uom": "EA"}], source="20261005T143221Z-1.pdf"),
    job("130155", "Morgan Steel", 304, D(2026, 10, 1), 3640.00, "2026-10-05T14:32:22Z", "morgan-130155.pdf", po="59149", po_id=7151, merch=3640.00, lines_match=[{"qty": 26, "amount": 3640.00, "uom": "EA"}], source="20261005T143222Z-1.pdf"),
    job("130156", "Morgan Steel", 304, D(2026, 10, 1), 14288.00, "2026-10-05T14:32:33Z", "morgan-130156.pdf", po="59134", po_id=7136, merch=14288.00, lines_match=[{"qty": 68, "amount": 10268.00, "uom": "EA", "token": "A-06810"}, {"qty": 30, "amount": 4020.00, "uom": "EA", "token": "A-02393"}], source="20261005T143233Z-1.pdf", owner="shawn"),
    job("PS-INV104088", "Legacy Wire Products", 292, D(2026, 10, 5), 1080.00, "2026-10-05T14:49:56Z", "legacy-104088.pdf", po="59320", po_id=7322, merch=880.00, freight=200.00, lines_match=[{"qty": 20, "amount": 880.00, "uom": "EA"}], source="20261005T144956Z-1.pdf"),
    job("PS-INV104089", "Legacy Wire Products", 292, D(2026, 10, 5), 734.25, "2026-10-05T18:06:54Z", "legacy-104089.pdf", po="59283", po_id=7285, merch=660.00, freight=74.25, lines_match=[{"qty": 15, "amount": 660.00, "uom": "EA"}], source="20261005T180654Z-1.pdf"),
    job("PS-INV104090", "Legacy Wire Products", 292, D(2026, 10, 5), 783.20, "2026-10-05T19:00:16Z", "legacy-104090.pdf", po="59303", po_id=7305, merch=704.00, freight=79.20, lines_match=[{"qty": 16, "amount": 704.00, "uom": "EA"}], source="20261005T190016Z-1.pdf"),
    job("PS-INV104091", "Legacy Wire Products", 292, D(2026, 10, 5), 487.60, "2026-10-05T19:14:46Z", "legacy-104091.pdf", po="59275", po_id=7277, merch=93.60, freight=394.00, lines_match=[{"qty": 100, "amount": 78.00, "uom": "EA"}, {"qty": 20, "amount": 15.60, "uom": "EA"}], source="20261005T191446Z-1.pdf", owner="shawn"),
    job("0214825", "Freepoint Energy", 448, D(2026, 10, 5), 8392.46, "2026-10-05T20:52:55Z", "freepoint-0214825.pdf", kind="misc", lines=[{"item_id": 29, "description": "Electric bill 08/31/2026 through 10/01/2026", "qty": 1, "unit_price": 8392.46, "gl_id": 79}], source="20261005T205255Z-1.pdf", pages=[0, 1], extra="Printed taxes are included in the utilities line, the same way as Freepoint bill 0204730."),
    job("52649", "AFT Industries", 385, D(2026, 10, 5), 280.00, "2026-10-05T21:15:51Z", "aft-52649.pdf", po="59255", po_id=7257, merch=280.00, source="20261005T211551Z-1.pdf", owner="shawn", hold="The invoice bills 1, 7, 2, and 2 each on PO 59255. The open receipts are 1, 7, 2, and 1 each. The last line bills 2 each and the receipt is 1 each. Nothing was selected."),
    job("32684025", "Purvis Industries", 333, D(2026, 10, 5), 8078.43, "2026-10-06T01:12:44Z", "purvis-32684025.pdf", po="59065", po_id=7067, merch=7970.00, freight=108.43, lines_match=[{"qty": 40, "amount": 7970.00, "uom": "EA", "token": "22209"}], source="20261006T011244Z-1.pdf", pages=[0, 1]),
    job("32684026", "Purvis Industries", 333, D(2026, 10, 5), 7219.99, "2026-10-06T01:12:44Z", "purvis-32684026.pdf", po="59111", po_id=7113, merch=7173.00, freight=46.99, lines_match=[{"qty": 36, "amount": 7173.00, "uom": "EA", "token": "22209"}], source="20261006T011244Z-2.pdf", pages=[0, 1]),
    job("Z251088432", "EMJ", 208, D(2026, 10, 5), 1445.57, "2026-10-06T03:09:51Z", "emj-Z251088432.pdf", po="59340", po_id=7342, merch=1445.57, lines_match=[
        {"qty": 432, "amount": 395.20, "uom": "IN", "token": "Square Bar"},
        {"qty": 288, "amount": 166.25, "uom": "IN", "token": "RB-1.25-6061"},
        {"qty": 576, "amount": 416.00, "uom": "IN", "token": "RB-1.50"},
        {"qty": 288, "amount": 273.24, "uom": "IN", "token": "7075"},
        {"qty": 288, "amount": 194.88, "uom": "IN", "token": "1.375"},
    ], source="20261006T030951Z-1.pdf", pages=[0, 1], extra="The invoice bills pounds. The receipts are the same 12-foot bars in inches."),
    job("885055", "O'Neal Steel", 137, D(2026, 10, 5), 9490.50, "2026-10-06T04:30:39Z", "oneal-885055.pdf", po="59341", po_id=7343, merch=9490.50, lines_match=[{"qty": 24, "amount": 9490.50, "uom": "EA"}], source="20261006T043039Z-1.pdf", owner="anthony"),
    job("0040488455", "Gas and Supply", 71, D(2026, 10, 5), 596.96, "2026-10-06T04:38:48Z", "gas-0040488455.pdf", kind="misc", lines=[{"item_id": 31, "description": "Shop supplies", "qty": 1, "unit_price": 596.96}], source="20261006T043848Z-1.pdf", owner="none"),
    job("9142053", "Phoenix Metals", 291, D(2026, 10, 5), 39761.69, "2026-10-06T05:06:44Z", "phoenix-9142053.pdf", po="59241", po_id=7243, merch=39690.00, fee=71.69, fee_name="fuel surcharge", lines_match=[{"qty": 100, "amount": 39690.00, "uom": "EA"}], source="20261006T050644Z-1.pdf", pages=[0], owner="anthony"),
    job("151291", "Pittsburg Steel", 0, D(2026, 10, 5), 4594.50, "2026-10-06T13:01:46Z", "pittsburg-151291.pdf", po="59174", po_id=7176, merch=4594.50, lines_match=[{"qty": 9, "amount": 4594.50, "uom": "EA"}], source="20261006T130146Z-1.pdf", owner="anthony"),
    job("PS-INV104098", "Legacy Wire Products", 292, D(2026, 10, 6), 3249.20, "2026-10-06T14:33:31Z", "legacy-104098.pdf", po="59282", po_id=7284, merch=2592.00, freight=657.20, lines_match=[{"qty": 42, "amount": 1512.00, "uom": "EA"}, {"qty": 20, "amount": 1080.00, "uom": "EA"}], source="20261006T143331Z-1.pdf"),
    job("PS-INV104099", "Legacy Wire Products", 292, D(2026, 10, 6), 1292.00, "2026-10-06T14:46:00Z", "legacy-104099.pdf", po="59271", po_id=7273, merch=1080.00, freight=212.00, lines_match=[{"qty": 20, "amount": 1080.00, "uom": "EA", "token": "59271-03"}], source="20261006T144600Z-1.pdf", extra="Purchase order 59271 also has an open receipt for 17 each that this invoice does not bill."),
    job("PS-INV104100", "Legacy Wire Products", 292, D(2026, 10, 6), 326.20, "2026-10-06T15:15:32Z", "legacy-104100.pdf", po="59303", po_id=7305, merch=252.00, freight=74.20, lines_match=[{"qty": 7, "amount": 252.00, "uom": "EA"}], source="20261006T151532Z-1.pdf", owner="shawn"),
    job("9162719", "Phoenix Metals", 291, D(2026, 10, 6), 440.41, "2026-10-06T17:47:14Z", "phoenix-9162719.pdf", po="59356", po_id=7358, merch=438.57, fee=1.84, fee_name="fuel surcharge", lines_match=[{"qty": 2160, "amount": 438.57, "uom": "IN", "token": "ANG-1.50"}], source="20261006T174714Z-1.pdf", pages=[0], owner="anthony", extra="The invoice bills 9 pieces of 20-foot angle. The receipt is 2,160 inches."),
    job("1474805", "RMP Industrial Supply", 322, D(2026, 10, 6), 277.18, "2026-10-06T23:02:04Z", "rmp-1474805.pdf", kind="misc", lines=[{"item_id": 53, "description": "Inserts", "qty": 1, "unit_price": 249.51}], freight=27.67, source="20261006T230204Z-1.pdf", pages=[0, 1], extra="The printed reference VERBAL : MANUEL is not a KIMCO purchase order."),
]


def prepare_pdfs() -> None:
    SRC.mkdir(parents=True, exist_ok=True)
    for item in JOBS:
        src = BASE / item["source"]
        dest = SRC / item["pdf"]
        reader = PdfReader(str(src))
        indexes = item["pages"] if item["pages"] is not None else list(range(len(reader.pages)))
        from pypdf import PdfWriter

        writer = PdfWriter()
        for index in indexes:
            writer.add_page(reader.pages[index])
        with dest.open("wb") as handle:
            writer.write(handle)
        document = pymupdf.open(dest)
        for page in document:
            if item["number"] not in page.get_text("text"):
                raise SystemExit(f"{dest.name} page is missing invoice {item['number']}")
        LOGGER.info("PDF %s pages %s", dest.name, document.page_count)


def remember_created(number: str, bill_id: int, vendor_id: int) -> None:
    path = OUT / "created-ids.json"
    rows = json.loads(path.read_text()) if path.exists() else []
    if any(int(row.get("bill_id") or 0) == int(bill_id) for row in rows):
        return
    rows.append({"invoice": number, "bill_id": int(bill_id), "vendor_id": int(vendor_id)})
    path.write_text(json.dumps(rows, indent=2))


def scan_live_bills(client: KimcoClient) -> dict[tuple[int, str], int]:
    found: dict[tuple[int, str], int] = {}
    misses = 0
    for invoice_id in range(MIN_NEW_ID + 1, MIN_NEW_ID + 40):
        try:
            record = client.get_item("ap_invoices", invoice_id)
        except KimcoError:
            misses += 1
            if misses >= 4:
                break
            continue
        misses = 0
        values = record.get("values") or {}
        vendor_id = lookup_id(values.get("Vendor"))
        number = str(values.get("Invoice_Number") or "")
        if vendor_id and number:
            found[(int(vendor_id), number)] = int(invoice_id)
            LOGGER.info("Existing new bill %s vendor %s invoice %s", invoice_id, vendor_id, number)
    return found


def existing_id(client: KimcoClient, number: str, vendor_id: int, live: dict[tuple[int, str], int]) -> int | None:
    if (int(vendor_id), number) in live:
        return int(live[(int(vendor_id), number)])
    index = json.loads((OUT / "kimco-index.json").read_text())
    hits = [row for row in index["invoices"] if str(row.get("number") or "") == number]
    path = OUT / "created-ids.json"
    if path.exists():
        for row in json.loads(path.read_text()):
            if row.get("invoice") == number and row.get("bill_id"):
                hits.append({"id": row["bill_id"]})
    for hit in hits:
        record = client.get_item("ap_invoices", int(hit["id"]))
        values = record.get("values") or {}
        if lookup_id(values.get("Vendor")) == vendor_id and str(values.get("Invoice_Number") or "") == number:
            return int(hit["id"])
    return None


def header_for(client: KimcoClient, item: dict[str, Any], batch_id: int, sample: dict[str, Any]) -> int:
    # Local wrapper so a created id at or below 10566 aborts.
    created = _create_header_local(client, item, batch_id, sample)
    return created


def _create_header_local(client: KimcoClient, item: dict[str, Any], batch_id: int, sample: dict[str, Any]) -> int:
    from ap_clerk.rules import comments_for, due_date_from_terms, kimco_datetime

    day = item["day"]
    terms = "Due Upon Receipt" if item.get("due_on_receipt") or "upon receipt" in str(sample.get("terms") or "").lower() else str(sample["terms"] or "")
    if item["vendor_id"] == 448:
        terms = "Due Upon Receipt"
        due = day
    else:
        due = due_date_from_terms(day, terms)
    payload: dict[str, Any] = {
        "AP_Invoice_Batch": {"id": int(batch_id)},
        "Vendor": {"id": int(item["vendor_id"])},
        "Invoice_Number": item["number"],
        "Invoice_Type": 3 if item.get("po_id") else 4,
        "Invoice_Date": kimco_datetime(day),
        "Invoice_Verification_Amount": float(item["amount"]),
        "Invoice_Due_Date": kimco_datetime(due),
        "Terms_Code": {"id": int(sample["terms_id"])},
        "Currency": {"id": 3},
        "Remit_To_Address": {"id": int(sample["remit_id"])},
        "Transaction_Date": kimco_datetime(day),
        "Comments": comments_for("live"),
    }
    if item.get("po_id"):
        payload["Purchase_Order"] = {"id": int(item["po_id"])}
    created, _body, status, error = client.create("ap_invoices", payload)
    if created is None:
        raise SystemExit(f"Header create for {item['number']} failed HTTP {status}: {error}")
    created = int(created)
    if created <= MIN_NEW_ID:
        raise SystemExit(f"Create for {item['number']} returned existing id {created}. Not editing it.")
    record = client.get_item("ap_invoices", created)
    values = record.get("values") or {}
    if str(values.get("Invoice_Number") or "") != item["number"]:
        raise SystemExit(f"Created id {created} is not invoice {item['number']}")
    if values.get("Posted") not in (None, "", False):
        raise SystemExit(f"Created bill {created} is posted")
    if str(values.get("Comments") or "") != "API Agent":
        raise SystemExit(f"Created bill {created} is not stamped API Agent")
    if int(lookup_id(values.get("Vendor")) or 0) != int(item["vendor_id"]):
        raise SystemExit(f"Created bill {created} vendor is not {item['vendor_id']}")
    remember_created(item["number"], created, int(item["vendor_id"]))
    LOGGER.info("Created %s as %s", item["number"], created)
    return created


def note_html(text: str, owner: str) -> str:
    body = text.strip()
    if not body.startswith("AP Clerk:"):
        body = "AP Clerk: " + body
    if owner == "none":
        if "data-mention-id" in body or "@" in body:
            raise SystemExit("This note must not tag an owner")
        return f"<p>{body}</p>"
    if owner == "anthony":
        if "@Anthony" not in body:
            body = body.replace("AP Clerk:", "AP Clerk: @Anthony", 1)
        if "@Shawn McKibben" not in body:
            body = body.replace("@Anthony", "@Anthony @Shawn McKibben", 1)
        html = body.replace("@Shawn McKibben", SHAWN_MENTION_HTML, 1)
        if "@Anthony" not in html or 'data-mention-id="104"' not in html:
            raise SystemExit("Anthony note is missing the plain name or Shawn mention")
        return html if html.startswith("<") else f"<p>{html}</p>"
    if owner == "shawn":
        return ap_clerk_edit_note(body, action="shawn")
    return ap_clerk_edit_note(body, action="treyce")


def write_note(client: KimcoClient, invoice_id: int, html: str) -> int | None:
    before = totals(client.get_item("ap_invoices", invoice_id))
    existing = [note for note in before["notes"] if str(note.get("text") or "").startswith("AP Clerk:")]
    if existing:
        return int(existing[-1]["id"])
    status = put(client, invoice_id, added_comment_payload(invoice_id, html))
    if status >= 400:
        return None
    after = totals(client.get_item("ap_invoices", invoice_id))
    notes = [note for note in after["notes"] if str(note.get("text") or "").startswith("AP Clerk:")]
    return int(notes[-1]["id"]) if notes else None


def confirm_attachment(client: KimcoClient, invoice_id: int, number: str) -> tuple[int, str]:
    from scripts.sept25_30_attachment_audit import attachment_bytes

    QC.mkdir(parents=True, exist_ok=True)
    items = client.list_attachments(invoice_id)
    if len(items) != 1:
        return 0, f"attachment-count-{len(items)}"
    content = attachment_bytes(client, invoice_id, items[0])
    if not content:
        return 0, "attachment-download-failed"
    document = pymupdf.open(stream=content, filetype="pdf")
    for page in document:
        if number not in page.get_text("text"):
            return document.page_count, "page-missing-invoice-number"
    for index, page in enumerate(document, start=1):
        image = QC / f"{invoice_id}-p{index}.png"
        page.get_pixmap(matrix=pymupdf.Matrix(1.7, 1.7), alpha=False).save(str(image))
    return document.page_count, "ok"


def discover_pittsburg(client: KimcoClient) -> int | None:
    rows = client.list_items("purchase_lines", fields="Purchase_Order_Number,Purchase_Line_Number")
    line_id = None
    for row in rows:
        if re.search(r"PO59174\b", json.dumps(row.get("values") or {})):
            line_id = int(row["id"])
            break
    if line_id is None:
        LOGGER.info("PO 59174 line was not in the purchase-line list")
        return None
    record = client.get_item("purchase_lines", line_id)
    values = record.get("values") or {}
    LOGGER.info("PO 59174 line %s keys %s", line_id, sorted(values))
    for key, value in values.items():
        if "vendor" in key.lower() or "Vendor" in key:
            LOGGER.info("vendor field %s = %s", key, value)
            found = lookup_id(value)
            if found:
                return int(found)
    text = json.dumps(values)
    LOGGER.info("PO 59174 vendor text snippet %s", text[:500])
    return None


def piece_inch_match(item: dict[str, Any], facts: list[dict[str, Any]]) -> tuple[list[int], float, str]:
    """40 pieces billed against 40 inches received is a match when the dollars are within $75."""
    open_rows = [row for row in facts if row.get("id") and row.get("open") and not row.get("placeholder")]
    want = float((item.get("lines_match") or [{"qty": 0}])[0]["qty"])
    hits = [
        row
        for row in open_rows
        if abs(float(row.get("qty") or 0) - want) < 0.001 and "IN" in str(row.get("uom") or "").upper()
    ]
    if len(hits) != 1:
        detail = ", ".join(
            f"{row['id']} qty {row.get('qty')} {row.get('uom')} ${row.get('extended')}" for row in open_rows[:6]
        )
        return [], 0.0, (
            f"No single open inch receipt matches {want:g} pieces on PO {item.get('po')}. "
            f"Open receipts: {detail or 'none'}. Nothing was selected."
        )
    gap = round(float(item["merch"]) - float(hits[0].get("extended") or 0), 2)
    if abs(gap) >= 75:
        return [], gap, (
            f"The price gap is ${abs(gap):,.2f}, which is $75 or more. Receipts were not selected."
        )
    return [int(hits[0]["id"])], gap, ""


def enter_one(client, item, batch_id, samples, receipts, live) -> dict[str, Any]:
    number = item["number"]
    path = SRC / item["pdf"]
    row: dict[str, Any] = {
        "vendor": item["vendor"],
        "invoice": number,
        "invoice_date": item["day"].isoformat(),
        "received": item["received"],
        "printed_total": item["amount"],
        "po": item.get("po") or "",
        "bill_id": "",
        "result": "HOLD",
        "owner": "",
        "reason": "",
        "batch": "",
        "note_id": "",
        "note": "",
        "email_moved": "n",
        "pages": 0,
        "receipts": "",
    }
    already = existing_id(client, number, int(item["vendor_id"]), live)
    if already and already <= MIN_NEW_ID:
        row["bill_id"] = already
        row["result"] = "ALREADY"
        row["reason"] = f"Already in KIMCO as bill {already}. Not edited."
        return row
    sample = sample_for(client, int(item["vendor_id"]), samples)
    ppv = 0.0
    hold_reason = str(item.get("hold_without_lines") or "")
    receipt_ids: list[int] = []
    if item["kind"] == "po" and not hold_reason and item.get("piece_inch"):
        receipt_ids, ppv, hold_reason = piece_inch_match(item, list(receipts.get(item["po"], [])))
    elif item["kind"] == "po" and not hold_reason:
        facts = list(receipts.get(item["po"], []))
        real = [fact for fact in facts if fact.get("id") and not fact.get("placeholder")]
        if not item.get("po_id"):
            hold_reason = f"Purchase order {item['po']} was not found. Nothing was selected."
        elif not real:
            hold_reason = po_absence_reason(str(item["po"]), list(item.get("lines_match") or []))
        else:
            receipt_ids, ppv, hold_reason = match_lines(item, real)
    created = already or header_for(client, item, batch_id, sample)
    live[(int(item["vendor_id"]), number)] = int(created)
    row["bill_id"] = created
    before = totals(client.get_item("ap_invoices", created))
    has_lines = bool(before["lines"])
    receipts_already = [int(line["receipt_id"]) for line in before["lines"] if line.get("receipt_id")]
    target_amount = round(float(item["amount"]), 2)
    if item.get("hold_without_lines"):
        hold_reason = str(item["hold_without_lines"])
        line_status = "not-added"
        charge_status = ""
        select_status = ""
    elif item["kind"] == "misc":
        line_status = "already" if has_lines else add_misc(client, created, item)
        if line_status in {"added", "already"} and cents(before["covered"]) == target_amount and has_lines:
            charge_status = "already"
        else:
            charge_status = add_charges(client, created, item, 0.0) if line_status in {"added", "already"} else ""
        select_status = ""
    elif receipt_ids and not hold_reason:
        if receipts_already:
            select_status = "selected"
            charge_status = "already" if cents(before["covered"]) == target_amount else add_charges(client, created, item, ppv)
        else:
            try:
                select_status = client.try_select_receipts(created, receipt_ids)
            except KimcoError:
                select_status = "blocked"
            charge_status = add_charges(client, created, item, ppv) if select_status == "selected" else ""
        if select_status == "selected":
            chosen = {int(value) for value in receipt_ids}
            for fact in receipts.get(item.get("po") or "", []):
                if int(fact.get("id") or 0) in chosen:
                    fact["open"] = False
        line_status = ""
    else:
        select_status = "not-selected"
        charge_status = ""
        line_status = ""
    attach_status = "attached" if client.list_attachments(created) else attach(client, created, path)
    pages, attach_check = confirm_attachment(client, created, number)
    row["pages"] = pages
    live_totals = totals(client.get_item("ap_invoices", created))
    covered_ok = cents(live_totals["verification"]) == target_amount and live_totals["covered"] == target_amount
    passed = (
        not hold_reason
        and covered_ok
        and live_totals["posted"] in (None, "", False)
        and attach_status == "attached"
        and attach_check == "ok"
        and (item["kind"] == "misc" or select_status == "selected")
    )
    if not passed and not hold_reason:
        hold_reason = (
            f"Live lines, charges, and tax are ${live_totals['covered']:,.2f} against ${item['amount']:,.2f}. "
            f"Lines {line_status or select_status}, charges {charge_status or 'n/a'}, PDF {attach_status} {attach_check}."
        )
    target = batch_id if passed else TRANSFER
    move_batch(client, created, target)
    total = f"${item['amount']:,.2f}"
    if passed and abs(ppv) >= 0.005 and item["owner"] != "anthony":
        owner = "shawn"
    elif item["owner"] == "anthony":
        owner = "anthony"
    elif not passed and item["kind"] == "po":
        owner = "shawn"
    else:
        owner = item["owner"]
    extra = f" {item['extra']}" if item.get("extra") else ""
    if passed:
        if owner == "none":
            text = (
                f"AP Clerk: {item['vendor']} invoice {number} is entered and is not posted. "
                f"The PDF total is {total}. There is no KIMCO purchase order, so it is miscellaneous shop supplies."
            )
        elif abs(ppv) >= 0.005:
            text = (
                f"AP Clerk: {item['vendor']} invoice {number} is entered and is not posted. "
                f"The PDF total is {total}. Purchase order {item.get('po')} receipts were selected. "
                f"A purchase price variance of ${ppv:,.2f} covers the difference.{extra}"
            )
        else:
            freight_bit = f" Freight of ${item['freight']:,.2f} is an additional charge." if item.get("freight") else ""
            fee_bit = ""
            if item.get("fee"):
                fee_bit = f" The {item.get('fee_name') or 'fee'} of ${item['fee']:,.2f} is an additional charge."
            po_bit = f" Purchase order {item['po']} receipts were selected." if item.get("po") else " There is no purchase order."
            text = (
                f"AP Clerk: {item['vendor']} invoice {number} is entered and is not posted. "
                f"The PDF total is {total}.{po_bit}{freight_bit}{fee_bit}{extra}"
            )
        result = "PASS"
    else:
        text = (
            f"AP Clerk: {item['vendor']} invoice {number} is on hold and is not posted. "
            f"The PDF total is {total}. {hold_reason} The bill is in Transfer AP.{extra}"
        )
        result = "HOLD"
    html = note_html(text, owner)
    note_id = write_note(client, created, html)
    final = totals(client.get_item("ap_invoices", created))
    row.update(
        {
            "result": result,
            "owner": "" if owner == "none" else owner,
            "reason": "" if result == "PASS" else hold_reason,
            "batch": final.get("batch") or "",
            "batch_id": final.get("batch_id"),
            "note_id": note_id or "",
            "note": text,
            "header_total": final.get("verification"),
            "covered": final.get("covered"),
            "receipts": ",".join(str(value) for value in receipt_ids) if result == "PASS" else "",
            "ppv": ppv if result == "PASS" else "",
            "items": "; ".join(
                f"{line.get('item')} {line.get('description')} qty {line.get('qty')} @ {line.get('unit_price')}"
                for line in final["lines"]
            ),
            "gl": "; ".join(str(line.get("gl") or "") for line in final["lines"]),
        }
    )
    LOGGER.info("%s %s bill %s covered %s", result, number, created, final.get("covered"))
    return row


def file_emails(rows: list[dict[str, Any]]) -> None:
    listing = json.loads(LISTING.read_text())["rows"]
    by_received = {row["received"]: row for row in listing}
    entered_received = {}
    for row in rows:
        if row.get("result") in {"PASS", "HOLD"} and row.get("bill_id"):
            entered_received.setdefault(row["received"], []).append(row)
    creds = load_graph_credentials()
    graph = GraphClient.authenticate(creds.tenant_id or "", creds.client_id or "", creds.client_secret or "")
    dest = archive_folder(graph)
    days: dict[str, list] = {}
    cache: dict[str, dict] = {}
    moves = []
    for received, bills in entered_received.items():
        expected = by_received[received]
        # Every invoice we intended from this email must be entered.
        intended = [item for item in JOBS if item["received"] == received and item.get("vendor_id")]
        if len(bills) < len(intended):
            LOGGER.info("Leave %s in Inbox. %s of %s invoices entered", received, len(bills), len(intended))
            continue
        from datetime import date as date_cls

        day = date_cls.fromisoformat(received[:10])
        key = day.isoformat()
        if key not in days:
            days[key] = graph.list_messages(ALLOWED_MAILBOX, received_from=day, received_to=day)
        found = []
        for message in days[key]:
            if str(message.get("receivedDateTime") or "") != received:
                continue
            if sender_of(message).lower() != expected["sender"].lower():
                continue
            if plain(str(message.get("subject") or "")) != plain(expected["subject"]):
                continue
            full = graph.get_message(ALLOWED_MAILBOX, str(message.get("id") or ""), select=SELECT)
            if (
                str(full.get("receivedDateTime") or "") == received
                and sender_of(full).lower() == expected["sender"].lower()
                and plain(str(full.get("subject") or "")) == plain(expected["subject"])
            ):
                found.append(full)
        unique = {str(item.get("id") or ""): item for item in found if item.get("id")}
        result = {"received": received, "subject": expected["subject"], "moved": "n"}
        if len(unique) != 1:
            result["result"] = f"found-{len(unique)}"
            moves.append(result)
            continue
        message = next(iter(unique.values()))
        path = folder_path(graph, str(message.get("parentFolderId") or ""), cache)
        if path != "Inbox":
            result["result"] = f"not-inbox:{path}"
            moves.append(result)
            continue
        if stamp(message.get("lastModifiedDateTime")) != stamp(expected["lastModifiedDateTime"]):
            result["result"] = "lastModified-changed"
            moves.append(result)
            continue
        old_id = str(message.get("id") or "")
        response = graph.request(
            "PATCH",
            graph._messages_url(ALLOWED_MAILBOX, old_id),
            json={"categories": [ENTERED_IN_AI_CATEGORY]},
            headers={"Content-Type": "application/json"},
        )
        if response.status_code >= 400:
            result["result"] = f"category-http-{response.status_code}"
            moves.append(result)
            continue
        moved = graph.request(
            "POST",
            graph._messages_url(ALLOWED_MAILBOX, old_id, "move"),
            json={"destinationId": dest},
            headers={"Content-Type": "application/json"},
        )
        if moved.status_code >= 400:
            result["result"] = f"move-http-{moved.status_code}"
            moves.append(result)
            continue
        new_id = str((moved.json() or {}).get("id") or "")
        check = graph.get_message(ALLOWED_MAILBOX, new_id, select=SELECT)
        after = folder_path(graph, str(check.get("parentFolderId") or ""), cache)
        cats = [str(item) for item in (check.get("categories") or [])]
        result.update({"moved": "y" if after == "Inbox/9 - FORT WORTH ARCHIVE" and cats == [ENTERED_IN_AI_CATEGORY] else "n", "result": "moved", "folder_after": after, "categories_after": cats})
        if result["moved"] == "y":
            for bill in bills:
                bill["email_moved"] = "y"
        moves.append(result)
        LOGGER.info("Mail %s %s", received, result["moved"])
    (OUT / "mail-moves.json").write_text(json.dumps(moves, indent=2))


TOO_NEW = [
    ("2026-10-07T15:26:34Z", "JP Steel", "125667"),
    ("2026-10-07T15:44:10Z", "Primo Brands", "06J6709801835"),
    ("2026-10-07T18:23:02Z", "Houston Plating", "795289"),
    ("2026-10-07T18:30:12Z", "MSC", "60517341"),
    ("2026-10-07T19:23:29Z", "3P Industries", "142578"),
    ("2026-10-07T21:31:09Z", "A-1 Image", "67534"),
    ("2026-10-07T22:27:57Z", "Guerrero Plating", "77946"),
    ("2026-10-08T00:24:43Z", "Priority1", "18515820"),
    ("2026-10-08T03:09:25Z", "EMJ", "10/07 invoice pack"),
    ("2026-10-08T04:30:52Z", "Metal Supermarkets", "SI1092554"),
    ("2026-10-08T04:53:25Z", "Gas and Supply", "0040492629"),
    ("2026-10-08T04:53:25Z", "Gas and Supply", "0040492260"),
    ("2026-10-08T06:36:28Z", "Ryerson", "9900104085"),
    ("2026-10-08T06:38:49Z", "McMaster-Carr", "73214449"),
    ("2026-10-08T14:28:41Z", "Morgan Steel", "130248"),
    ("2026-10-08T14:37:33Z", "Luxor Staffing", "PDF not opened; mail is under 48 hours"),
    ("2026-10-08T16:45:11Z", "Katy Spring", "146855"),
    ("2026-10-08T18:28:35Z", "KIMCO Accounting", "consulting invoice, PDF not opened"),
    ("2026-10-08T19:37:19Z", "AQPC", "11084"),
    ("2026-10-08T19:37:38Z", "AQPC", "11076"),
    ("2026-10-08T19:37:53Z", "AQPC", "11079"),
    ("2026-10-08T20:15:26Z", "Trace Metal", "238595"),
    ("2026-10-09T01:11:53Z", "Purvis", "32688747"),
    ("2026-10-09T01:29:29Z", "Air Products", "0436684378"),
    ("2026-10-09T04:30:53Z", "O'Neal Steel", "10/9 email, PDF not opened"),
]

BACKLOG = [
    ("2026-10-01T13:15:16Z", "Stella Source", "INV-0001510", "No invoice PDF in the email"),
    ("2026-10-01T14:55:40Z", "KIMCO Holdings", "2390", "No KIMCO vendor"),
    ("2026-10-01T18:13:33Z", "AQPC", "11053", "Payment request, no invoice PDF"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179869", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179870", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179871", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179877", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179878", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179879", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179880", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179897", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179898", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179899", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179900", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179901", "September invoice, not in KIMCO"),
    ("2026-10-02T02:02:33Z", "Arrow Plating", "9179902", "September invoice, not in KIMCO"),
    ("2026-10-02T13:43:21Z", "AQPC", "11055", "Payment request, no invoice PDF"),
    ("2026-10-02T13:43:28Z", "AQPC", "11056", "Payment request, no invoice PDF"),
    ("2026-10-05T19:50:36Z", "3P Industries", "142539", "Scanned packing list, no printed invoice total"),
    ("2026-10-05T20:02:28Z", "3P Industries", "142540", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T13:42:21Z", "AQPC", "11068", "Payment request, no invoice PDF"),
    ("2026-10-06T13:44:14Z", "AQPC", "11061", "Payment request, no invoice PDF"),
    ("2026-10-06T14:37:04Z", "3P Industries", "142545", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T16:23:50Z", "AQPC", "11071", "Payment request, no invoice PDF"),
    ("2026-10-06T18:26:30Z", "3P Industries", "142550", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T18:36:30Z", "3P Industries", "142552", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T20:31:05Z", "3P Industries", "142544", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142560", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142561", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142562", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142563", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142564", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142565", "Scanned packing list, no printed invoice total"),
    ("2026-10-06T21:13:36Z", "3P Industries", "142566", "Scanned packing list, no printed invoice total"),
    ("2026-10-07T04:27:55Z", "O'Neal Steel", "885055", "Same printed invoice number as the 10/6 bill for $9,490.50. This one is $10,503.04 on PO 59137. Left in Inbox."),
]


AUTOPAY_RECEIVED = {
    "2026-10-01T05:00:34Z",
    "2026-10-01T17:05:16Z",
    "2026-10-05T12:29:16Z",
    "2026-10-05T12:37:21Z",
    "2026-10-05T17:04:35Z",
    "2026-10-06T15:26:22Z",
    "2026-10-06T15:27:27Z",
}


def autopay_folder(graph: GraphClient) -> str:
    response = graph.request(
        "GET",
        graph._user_url(ALLOWED_MAILBOX, "mailFolders/inbox/childFolders"),
        params={"$select": "id,displayName", "$top": 50},
    )
    if response.status_code != 200:
        raise SystemExit(f"Inbox child folders HTTP {response.status_code}")
    hits = [
        item
        for item in (response.json() or {}).get("value") or []
        if str(item.get("displayName") or "") == "AutoPay Archive" and item.get("id")
    ]
    if len(hits) != 1:
        raise SystemExit(f"AutoPay Archive was found {len(hits)} times. No autopay mail changed.")
    return str(hits[0]["id"])


def file_autopay() -> list[dict[str, Any]]:
    listing = json.loads(LISTING.read_text())["rows"]
    wanted = [row for row in listing if row["received"] in AUTOPAY_RECEIVED]
    creds = load_graph_credentials()
    graph = GraphClient.authenticate(creds.tenant_id or "", creds.client_id or "", creds.client_secret or "")
    dest = autopay_folder(graph)
    days: dict[str, list] = {}
    cache: dict[str, dict] = {}
    moves = []
    for expected in wanted:
        from datetime import date as date_cls

        day = date_cls.fromisoformat(expected["received"][:10])
        key = day.isoformat()
        if key not in days:
            days[key] = graph.list_messages(ALLOWED_MAILBOX, received_from=day, received_to=day)
        found = []
        for message in days[key]:
            if str(message.get("receivedDateTime") or "") != expected["received"]:
                continue
            if sender_of(message).lower() != expected["sender"].lower():
                continue
            if plain(str(message.get("subject") or "")) != plain(expected["subject"]):
                continue
            full = graph.get_message(ALLOWED_MAILBOX, str(message.get("id") or ""), select=SELECT)
            if (
                str(full.get("receivedDateTime") or "") == expected["received"]
                and sender_of(full).lower() == expected["sender"].lower()
                and plain(str(full.get("subject") or "")) == plain(expected["subject"])
            ):
                found.append(full)
        unique = {str(item.get("id") or ""): item for item in found if item.get("id")}
        result = {"received": expected["received"], "subject": expected["subject"], "moved": "n"}
        if len(unique) != 1:
            result["result"] = f"found-{len(unique)}"
            moves.append(result)
            continue
        message = next(iter(unique.values()))
        path = folder_path(graph, str(message.get("parentFolderId") or ""), cache)
        if path != "Inbox":
            result["result"] = f"not-inbox:{path}"
            moves.append(result)
            continue
        if stamp(message.get("lastModifiedDateTime")) != stamp(expected["lastModifiedDateTime"]):
            result["result"] = "lastModified-changed"
            moves.append(result)
            continue
        old_id = str(message.get("id") or "")
        response = graph.request(
            "PATCH",
            graph._messages_url(ALLOWED_MAILBOX, old_id),
            json={"categories": [LEGACY_AI_SKIPPED_CATEGORY]},
            headers={"Content-Type": "application/json"},
        )
        if response.status_code >= 400:
            result["result"] = f"category-http-{response.status_code}"
            moves.append(result)
            continue
        moved = graph.request(
            "POST",
            graph._messages_url(ALLOWED_MAILBOX, old_id, "move"),
            json={"destinationId": dest},
            headers={"Content-Type": "application/json"},
        )
        if moved.status_code >= 400:
            result["result"] = f"move-http-{moved.status_code}"
            moves.append(result)
            continue
        new_id = str((moved.json() or {}).get("id") or "")
        check = graph.get_message(ALLOWED_MAILBOX, new_id, select=SELECT)
        after = folder_path(graph, str(check.get("parentFolderId") or ""), cache)
        cats = [str(item) for item in (check.get("categories") or [])]
        result.update(
            {
                "moved": "y" if after == "Inbox/AutoPay Archive" and cats == [LEGACY_AI_SKIPPED_CATEGORY] else "n",
                "result": "moved",
                "folder_after": after,
                "categories_after": cats,
            }
        )
        moves.append(result)
        LOGGER.info("Autopay %s %s", expected["received"], result["moved"])
    (OUT / "autopay-moves.json").write_text(json.dumps(moves, indent=2))
    return moves


def write_workbook(rows: list[dict[str, Any]], batch: dict[str, Any], autopay: list[dict[str, Any]], sign_ins: int) -> None:
    passed = [row for row in rows if row["result"] == "PASS"]
    held = [row for row in rows if row["result"] == "HOLD"]
    backlog = list(BACKLOG)
    entered_numbers = {str(row["invoice"]) for row in rows if row["result"] in {"PASS", "HOLD"}}
    if "151291" not in entered_numbers:
        backlog.append(("2026-10-06T13:01:46Z", "Pittsburg Steel", "151291", "Vendor 1001-PITTSBURG STEEL is id 3 on PO 59174. No recent bill had remit and terms, and there is no open receipt. No header was created."))
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    oldest = min(backlog, key=lambda item: item[0])
    next_queue = "O'Neal Steel 885055 for $10,503.04 on PO 59137, received 2026-10-07T04:27:55Z. Same printed number as the 10/6 bill. Left in Inbox."
    values = [
        ("Item", "Value"),
        ("Run date", "Friday 2026-10-09 America/Chicago"),
        ("Batch", BATCH_NAME),
        ("Batch id", batch["id"]),
        ("Batch status", "Unposted"),
        ("Entered", len(passed) + len(held)),
        ("Passed", len(passed)),
        ("Held", len(held)),
        ("Skipped and archived (autopay)", sum(1 for row in autopay if row.get("moved") == "y")),
        ("Too new invoices", len(TOO_NEW)),
        ("Backlog invoices", len(backlog)),
        ("Oldest backlog received", f"{oldest[0]} {oldest[1]} {oldest[2]}"),
        ("Next in queue", next_queue),
        ("Sign-ins", f"Password worked. {sign_ins} sign-in(s) in this process. No password retry. No third sign-in."),
    ]
    for row_idx, pair in enumerate(values, start=1):
        summary.cell(row_idx, 1, pair[0])
        summary.cell(row_idx, 2, pair[1])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 88
    invoices = book.create_sheet("Invoices")
    headers = ["Vendor", "Invoice", "Invoice date", "Received date", "Printed total", "Bill id", "PASS/HOLD", "Owner", "Reason", "Batch", "Note id", "Email moved"]
    for col, name in enumerate(headers, start=1):
        cell = invoices.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for row_idx, row in enumerate(rows, start=2):
        data = [
            row["vendor"], row["invoice"], row["invoice_date"], row["received"], row["printed_total"],
            row.get("bill_id"), row["result"], row.get("owner") or "", row.get("reason") or "",
            row.get("batch") or "", row.get("note_id") or "", row.get("email_moved") or "n",
        ]
        for col, value in enumerate(data, start=1):
            invoices.cell(row_idx, col, value)
    invoices.freeze_panes = "C2"
    invoices.auto_filter.ref = f"A1:L{max(2, len(rows) + 1)}"
    widths = [28, 18, 14, 24, 14, 12, 12, 12, 70, 28, 12, 14]
    from openpyxl.utils import get_column_letter
    for idx, width in enumerate(widths, start=1):
        invoices.column_dimensions[get_column_letter(idx)].width = width
    too = book.create_sheet("Too new")
    for col, name in enumerate(("Received UTC", "Vendor", "Invoice"), start=1):
        cell = too.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for row_idx, item in enumerate(TOO_NEW, start=2):
        for col, value in enumerate(item, start=1):
            too.cell(row_idx, col, value)
    too.freeze_panes = "C2"
    too.column_dimensions["A"].width = 24
    too.column_dimensions["B"].width = 24
    too.column_dimensions["C"].width = 42
    backlog_sheet = book.create_sheet("Backlog")
    for col, name in enumerate(("Received UTC", "Vendor", "Invoice", "Why still open"), start=1):
        cell = backlog_sheet.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for row_idx, item in enumerate(backlog, start=2):
        for col, value in enumerate(item, start=1):
            backlog_sheet.cell(row_idx, col, value)
    backlog_sheet.freeze_panes = "C2"
    backlog_sheet.column_dimensions["A"].width = 24
    backlog_sheet.column_dimensions["B"].width = 24
    backlog_sheet.column_dimensions["C"].width = 18
    backlog_sheet.column_dimensions["D"].width = 78
    path = OUT / "AP-run-2026-10-09.xlsx"
    book.save(path)
    artifacts = Path("/workspace/artifacts")
    artifacts.mkdir(parents=True, exist_ok=True)
    book.save(artifacts / "AP-run-2026-10-09.xlsx")


def write_outputs(rows: list[dict[str, Any]], batch: dict[str, Any], autopay: list[dict[str, Any]], sign_ins: int) -> None:
    passed = [row for row in rows if row["result"] == "PASS"]
    held = [row for row in rows if row["result"] == "HOLD"]
    payload = {
        "run_date": "2026-10-09",
        "timezone": "America/Chicago",
        "batch_name": BATCH_NAME,
        "batch_id": batch["id"],
        "batch_status": "unposted",
        "entered": len(passed) + len(held),
        "passed": len(passed),
        "held": len(held),
        "autopay_archived": sum(1 for row in autopay if row.get("moved") == "y"),
        "too_new": len(TOO_NEW),
        "sign_ins": sign_ins,
        "rows": rows,
    }
    (OUT / "result.json").write_text(json.dumps(payload, indent=2, default=str))
    (OUT / "progress.json").write_text(json.dumps(rows, indent=2, default=str))
    write_workbook(rows, batch, autopay, sign_ins)
    with (QC / "readback.csv").open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["bill", "vendor", "invoice#", "printed total", "header total", "lines+charges sum", "gap", "items/GL", "receipts selected", "batch", "note id + text", "attachment page count", "email moved y/n"])
        for row in rows:
            printed = float(row["printed_total"])
            covered = row.get("covered")
            gap = "" if covered in (None, "") else round(printed - float(covered), 2)
            writer.writerow([
                row.get("bill_id"), row["vendor"], row["invoice"], f"{printed:.2f}", row.get("header_total"),
                covered, gap, f"{row.get('items') or ''} | {row.get('gl') or ''}", row.get("receipts"),
                row.get("batch"), f"{row.get('note_id')}: {row.get('note')}", row.get("pages"), row.get("email_moved"),
            ])


def main() -> None:
    prepare_pdfs()
    client = finish.login()
    finish.install_401_guard(client)
    pittsburg = discover_pittsburg(client)
    LOGGER.info("Pittsburg vendor id %s", pittsburg)
    for item in JOBS:
        if item["number"] == "151291" and pittsburg:
            item["vendor_id"] = pittsburg
    runnable = [item for item in JOBS if item["vendor_id"]]
    skipped_vendor = [item["number"] for item in JOBS if not item["vendor_id"]]
    if skipped_vendor:
        LOGGER.info("No vendor id, not creating %s", skipped_vendor)
    batches = client.list_items("ap_batches")
    batch = _find_or_create_batch(client, batches, BATCH_NAME)
    if batch["name"] != BATCH_NAME:
        raise SystemExit(f"Refusing batch {batch}")
    record = client.get_item("ap_batches", int(batch["id"]))
    status = (record.get("values") or {}).get("Status")
    if status not in (0, "0", None, ""):
        raise SystemExit(f"Batch {batch['id']} status is {status}. Not using a posted batch.")
    LOGGER.info("Batch %s id %s", BATCH_NAME, batch["id"])
    live = scan_live_bills(client)
    receipts = load_open_receipts(client, {item["po"] for item in runnable if item.get("po")})
    samples = {
        71: {"terms_id": 4, "terms": "F-N60-Net 60", "remit_id": 192},
        292: {"terms_id": 1, "terms": "Net 30-Net 30", "remit_id": 482},
        304: {"terms_id": 1, "terms": "Net 30-Net 30", "remit_id": 495},
        333: {"terms_id": 1, "terms": "Net 30-Net 30", "remit_id": 554},
        208: {"terms_id": 58, "terms": "F-0.5/10,N30-0.5% 10, Net 30", "remit_id": 337},
        137: {"terms_id": 4, "terms": "F-N60-Net 60", "remit_id": 258},
        322: {"terms_id": 1, "terms": "Net 30-Net 30", "remit_id": 536},
        448: {"terms_id": 62, "terms": "Due Upon Receipt", "remit_id": 734},
    }
    kept: list[dict[str, Any]] = []
    for item in runnable:
        try:
            sample_for(client, int(item["vendor_id"]), samples)
        except SystemExit as exc:
            LOGGER.info("Skipping %s. %s", item["number"], exc)
            item["vendor_id"] = 0
            continue
        kept.append(item)
    runnable = kept
    rows: list[dict[str, Any]] = []
    created = 0
    for item in runnable:
        if created >= CAP:
            break
        row = enter_one(client, item, int(batch["id"]), samples, receipts, live)
        rows.append(row)
        (OUT / "progress.json").write_text(json.dumps(rows, indent=2, default=str))
        if row["result"] in {"PASS", "HOLD"} and row.get("bill_id") and int(row["bill_id"]) > MIN_NEW_ID:
            created += 1
    file_emails(rows)
    autopay = file_autopay()
    write_outputs(rows, batch, autopay, finish.SIGN_INS)
    print(json.dumps({"batch": batch["id"], "entered": created, "sign_ins": finish.SIGN_INS, "results": [(r["invoice"], r["result"], r["bill_id"], r["email_moved"]) for r in rows]}))


if __name__ == "__main__":
    main()
