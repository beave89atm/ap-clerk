"""Gap-check and enter Aug 1+ invoices for four vendors.

Vendors: O'Neal Steel, AQPC, Legacy Wire, Gas & Supply.
Mailbox: accountspayable@kannonmfg.com only.
invent=false. No Mail.Send. Live KIMCO writes when --live (not --discover-only).

PPV gate is absolute dollars only: |PPV| under $75 may post; |PPV| >= $75
is price_variance HOLD, over-gate lines not selected, Transfer AP + @Shawn 104.
Packing-slip Success gate is suspended (Kyle 2026-09-23).
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from ap_clerk.auth import load_credentials
from ap_clerk.cli import _optional_graph_client, run_enter
from ap_clerk.graph import (
    ALLOWED_MAILBOX,
    ENTERED_IN_AI_CATEGORY,
    ENTERED_WITH_ISSUES_CATEGORY,
    AI_HOLD_CATEGORY,
    GraphError,
    has_process_category,
    is_fort_worth_inbox_folder,
)
from ap_clerk.inbox import (
    _pdfs_from_body_link,
    _prefer_printed_invoice_aliases,
    _safe_filename,
    sender_address,
    sender_name,
)
from ap_clerk.kimco import KimcoClient, KimcoError
from ap_clerk.misc_lines import misc_line_snapshot, type4_shop_supplies_lines_ok
from ap_clerk.pdf_invoice import (
    NON_INVOICE_ATTACHMENT_KINDS,
    classify_attachment,
    parse_invoice_pdf,
)
from ap_clerk.quality_v12 import COL_EXCEPTION_CATEGORY, COL_EXCEPTION_OWNER
from ap_clerk.receiving_owners import (
    missing_receipt_comments_1_html,
    missing_receipt_exception_owner,
    should_tag_missing_receipt,
)
from ap_clerk.rules import (
    AQPC_TOO_OLD_INVOICES,
    invoice_number_key,
    is_aqpc_vendor,
    is_gas_and_supply,
    lookup_id,
    lookup_text,
    money,
)
from ap_clerk.transfer_ap import apply_transfer_ap_batch_move
from gas_supply_0917 import apply_type4_misc_lines
import ap_clerk.cli as cli_mod
import ap_clerk.graph as graph_mod
import ap_clerk.rules as rules

LOGGER = logging.getLogger("ap_clerk.gap_four")

MIN_DAY = date(2026, 8, 1)
LIST_FROM = date(2026, 7, 15)
PPV_HOLD_AT = 75.0
# decide_ppv holds when abs(var) > PPV_MAX_ABS_ON_BILL. 74.99 makes >= 75.00 hold
# and leaves 74.99 in-gate. Percent gate is disabled (absolute dollars only).
PPV_ABS_THRESHOLD = 74.99

VENDORS = {
    "oneal": {
        "label": "O'Neal Steel",
        "canonical": "O'Neal Steel - Dallas (GP)",
        "vendor_id": 137,
        "batch_id": 722,
        "batch_name": "API Agent - 9/21/26 O'Neal",
        "new_batch": "API Agent - 9/23/26 O'Neal",
    },
    "aqpc": {
        "label": "AQPC",
        "canonical": "American Quality Powder Coating",
        "vendor_id": 22,
        "batch_id": 711,
        "batch_name": "API Agent - 9/15/26",
        "new_batch": "API Agent - 9/23/26 AQPC",
    },
    "legacy": {
        "label": "Legacy Wire",
        "canonical": "Legacy Wire Products",
        "vendor_id": 292,
        "batch_id": 717,
        "batch_name": "API Agent - 9/17/26 Legacy Wire",
        "new_batch": "API Agent - 9/23/26 Legacy Wire",
    },
    "gas": {
        "label": "Gas & Supply",
        "canonical": "Gas and Supply North Texas, LLC",
        "vendor_id": 71,
        "batch_id": 720,
        "batch_name": "API Agent - 9/17/26 Gas & Supply",
        "new_batch": "API Agent - 9/23/26 Gas & Supply",
    },
}

SHAWN_HTML_EMAIL = "Shawn.McKibben@kannonmfg.com"
ARTIFACT_JSON = Path("/opt/cursor/artifacts/AP-gap-oneal-aqpc-legacy-gas-2026-09-23.json")
ARTIFACT_XLSX = Path("/opt/cursor/artifacts/AP-gap-oneal-aqpc-legacy-gas-2026-09-23.xlsx")
RUN_JSON = ROOT / "runs" / "AP-gap-oneal-aqpc-legacy-gas-2026-09-23.json"
RUN_XLSX = ROOT / "runs" / "AP-gap-oneal-aqpc-legacy-gas-2026-09-23.xlsx"
PDF_DIR = ROOT / "runs" / "inbox-pdfs" / "gap-2026-09-23"

SKIP_FOLDER_NAMES = {
    "deleted items",
    "junk email",
    "drafts",
    "outbox",
    "sync issues",
    "conflicts",
    "local failures",
    "server failures",
    "conversation history",
    "rss feeds",
    "rss subscriptions",
}

NOISE_SUBJECT = re.compile(
    r"payment confirmation|market informer|payment status requested|"
    r"order pending payment|pending payment status|remittance advice",
    flags=re.I,
)

SEARCH_NEEDLES = (
    "O'Neal Steel",
    "onealsteel",
    "American Quality Powder",
    "AQPC",
    "Legacy Wire",
    "legacywire",
    "PS-INV",
    "Gas and Supply",
    "Gas & Supply",
    "gasandsupply",
)

SHEET_COLUMNS = [
    "Vendor",
    "Invoice #",
    "date",
    "PO",
    "Amount",
    "Result",
    "Why",
    "Exception category",
    "Exception owner",
    "KIMCO id",
    "Batch",
    "Outlook category",
    "Notes",
]


def configure_ppv_absolute() -> None:
    """Absolute |PPV| gate only. Percent-of-invoice does not HOLD under $75."""
    rules.PPV_MAX_ABS_ON_BILL = PPV_ABS_THRESHOLD
    rules.PPV_MAX_PCT_OF_INVOICE = 999.0

    def _too_old(**kwargs: Any) -> bool:
        number = invoice_number_key(kwargs.get("invoice_number"))
        if number not in AQPC_TOO_OLD_INVOICES:
            return False
        vendor = kwargs.get("vendor")
        vendor_id = kwargs.get("vendor_id")
        if not vendor and vendor_id in (None, ""):
            return True
        return bool(rules.is_aqpc_vendor(vendor, vendor_id))

    # Gap scope is invoice date OR email received on/after Aug 1.
    # Voided AQPC numbers stay blocked. Date-before-Aug-1 is not an automatic skip.
    rules.aqpc_invoice_too_old = _too_old
    cli_mod.aqpc_invoice_too_old = _too_old


def defer_fort_worth_move(*_args: Any, **_kwargs: Any) -> str:
    return "deferred-until-category"


def _compact(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def classify_vendor(
    *,
    subject: str = "",
    preview: str = "",
    from_name: str = "",
    from_addr: str = "",
) -> str | None:
    """Map a mailbox message to one of the four vendors. None if not ours."""
    blob = "\n".join([subject or "", preview or "", from_name or "", from_addr or ""]).lower()
    compact = _compact(blob)
    addr = (from_addr or "").lower()
    scores = {"oneal": 0, "aqpc": 0, "legacy": 0, "gas": 0}
    if "onealsteel.com" in addr or "onealsteel" in compact:
        scores["oneal"] += 5
    if "o'neal" in blob or "o’neal" in blob or "oneal steel" in blob:
        scores["oneal"] += 3
    if "quality powder" in blob or "aqpowder" in compact or "aqpc" in compact:
        scores["aqpc"] += 5
    if "american quality" in blob:
        scores["aqpc"] += 3
    if "legacywire.com" in addr or "legacywire" in compact or "legacy wire" in blob:
        scores["legacy"] += 5
    if "ps-inv" in blob and "legacy" in blob:
        scores["legacy"] += 2
    if "gasandsupply" in compact or "gas and supply" in blob or "gas & supply" in blob:
        scores["gas"] += 5
    if is_gas_and_supply(blob):
        scores["gas"] += 2
    best = max(scores, key=lambda key: scores[key])
    if scores[best] <= 0:
        return None
    # A second vendor must not steal a domain hit.
    winners = [key for key, score in scores.items() if score == scores[best] and score > 0]
    if len(winners) != 1:
        return None
    return best


def parse_day(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if len(text) < 10:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def in_scope(*, received: Any, invoice_date: Any) -> bool:
    """Invoice date OR email received on/after 2026-08-01."""
    rec = parse_day(received)
    inv = parse_day(invoice_date)
    if rec is not None and rec >= MIN_DAY:
        return True
    if inv is not None and inv >= MIN_DAY:
        return True
    return False


def vendor_text_key(text: str | None, vendor_id: Any = None) -> str | None:
    if is_aqpc_vendor(text, vendor_id):
        return "aqpc"
    blob = str(text or "")
    low = blob.lower()
    compact = _compact(blob)
    if vendor_id not in (None, "") and int(vendor_id) == 137:
        return "oneal"
    if vendor_id not in (None, "") and int(vendor_id) == 71:
        return "gas"
    if vendor_id not in (None, "") and int(vendor_id) == 292:
        return "legacy"
    if "oneal" in compact or "o'neal" in low or "o’neal" in low:
        return "oneal"
    if "legacy wire" in low or "legacywire" in compact:
        return "legacy"
    if "gas and supply" in low or "gas & supply" in low or "gasandsupply" in compact or is_gas_and_supply(blob):
        return "gas"
    return None


def _graph_pages(graph, url: str, *, params: dict[str, Any] | None = None, headers: dict[str, str] | None = None) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    first = True
    hops = 0
    while url and hops < 80:
        hops += 1
        kwargs: dict[str, Any] = {}
        if first and params:
            kwargs["params"] = params
        if headers:
            kwargs["headers"] = headers
        first = False
        response = graph.request("GET", url, **kwargs)
        if response.status_code != 200:
            raise GraphError(f"Graph HTTP {response.status_code} {url[:80]}")
        payload = response.json() or {}
        items.extend(item for item in (payload.get("value") or []) if isinstance(item, dict))
        url = payload.get("@odata.nextLink") or ""
    return items


def list_folders(graph) -> list[dict[str, Any]]:
    mailbox = ALLOWED_MAILBOX
    base = graph._user_url(mailbox, "mailFolders")
    top = _graph_pages(
        graph,
        base,
        params={"$select": "id,displayName,totalItemCount,childFolderCount", "$top": 50},
    )
    folders = list(top)
    for folder in list(top):
        if int(folder.get("childFolderCount") or 0) <= 0:
            continue
        child_url = graph._user_url(
            mailbox, f"mailFolders/{quote(str(folder.get('id')), safe='')}/childFolders"
        )
        try:
            children = _graph_pages(
                graph,
                child_url,
                params={"$select": "id,displayName,totalItemCount,childFolderCount", "$top": 50},
            )
        except GraphError as exc:
            LOGGER.info("Child folders skipped: %s", exc)
            continue
        for child in children:
            child["parentDisplayName"] = folder.get("displayName")
        folders.extend(children)
    return folders


def _skip_folder(name: str) -> bool:
    return (name or "").strip().lower() in SKIP_FOLDER_NAMES


def list_folder_messages(graph, folder_id: str) -> list[dict[str, Any]]:
    mailbox = ALLOWED_MAILBOX
    url = graph._user_url(mailbox, f"mailFolders/{quote(str(folder_id), safe='')}/messages")
    select = "id,subject,from,receivedDateTime,hasAttachments,flag,categories,bodyPreview,parentFolderId"
    try:
        return _graph_pages(
            graph,
            url,
            params={
                "$select": select,
                "$filter": f"receivedDateTime ge {LIST_FROM.isoformat()}T00:00:00Z",
                "$orderby": "receivedDateTime desc",
                "$top": 50,
            },
        )
    except GraphError:
        LOGGER.info("Folder filter failed; paging without filter")
    messages = _graph_pages(graph, url, params={"$select": select, "$top": 50})
    kept = []
    for msg in messages:
        day = parse_day(msg.get("receivedDateTime"))
        if day is None or day >= LIST_FROM:
            kept.append(msg)
    return kept


def search_messages(graph, needle: str) -> list[dict[str, Any]]:
    mailbox = ALLOWED_MAILBOX
    url = graph._messages_url(mailbox)
    return _graph_pages(
        graph,
        url,
        params={
            "$search": f'"{needle}"',
            "$select": "id,subject,from,receivedDateTime,hasAttachments,flag,categories,bodyPreview,parentFolderId",
            "$top": 50,
        },
        headers={"ConsistencyLevel": "eventual"},
    )


def collect_messages(graph) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    folders = list_folders(graph)
    folder_counts = [
        {
            "name": f.get("displayName"),
            "parent": f.get("parentDisplayName"),
            "items": f.get("totalItemCount"),
            "fort_worth": is_fort_worth_inbox_folder(str(f.get("displayName") or "")),
        }
        for f in folders
    ]
    seen: dict[str, dict[str, Any]] = {}
    scanned_folders = []
    for folder in folders:
        name = str(folder.get("displayName") or "")
        if _skip_folder(name):
            continue
        fid = str(folder.get("id") or "")
        if not fid:
            continue
        try:
            msgs = list_folder_messages(graph, fid)
        except GraphError as exc:
            LOGGER.info("Folder %s list failed: %s", name, exc)
            scanned_folders.append({"name": name, "error": str(exc)[:160]})
            continue
        scanned_folders.append({"name": name, "listed": len(msgs)})
        for msg in msgs:
            mid = str(msg.get("id") or "")
            if not mid:
                continue
            msg["_folder"] = name
            seen[mid] = msg
    for needle in SEARCH_NEEDLES:
        try:
            hits = search_messages(graph, needle)
        except GraphError as exc:
            LOGGER.info("Search %s failed: %s", needle, exc)
            continue
        for msg in hits:
            mid = str(msg.get("id") or "")
            if mid and mid not in seen:
                msg["_folder"] = msg.get("_folder") or "search"
                seen[mid] = msg
    return list(seen.values()), {"folders": folder_counts, "scanned": scanned_folders, "unique_messages": len(seen)}


def _stamp_bill(bill: dict[str, Any], message: dict[str, Any], vendor_key: str, pdf_path: str) -> dict[str, Any]:
    spec = VENDORS[vendor_key]
    bill = dict(bill)
    bill["vendor"] = spec["canonical"]
    bill["gap_vendor"] = vendor_key
    bill["graph_message_id"] = str(message.get("id") or "")
    bill["subject"] = str(message.get("subject") or "")
    bill["receivedDateTime"] = message.get("receivedDateTime")
    bill["from_name"] = sender_name(message)
    bill["categories"] = list(message.get("categories") or [])
    bill["folder"] = message.get("_folder")
    bill["action"] = "create"
    if not bill.get("pdf_path"):
        bill["pdf_path"] = pdf_path
    return bill


def bills_from_message(graph, message: dict[str, Any], vendor_key: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return (bills, noise notes). Does not invent invoice fields."""
    subject = str(message.get("subject") or "")
    preview = str(message.get("bodyPreview") or "")
    message_id = str(message.get("id") or "")
    from_name = sender_name(message)
    from_addr = sender_address(message)
    notes: list[dict[str, Any]] = []
    pdfs: list[tuple[str, bytes]] = []
    if message.get("hasAttachments"):
        pdfs = graph.download_pdf_attachments(ALLOWED_MAILBOX, message_id)
    link_hold = None
    if not pdfs:
        pdfs, link_hold = _pdfs_from_body_link(
            graph, ALLOWED_MAILBOX, message_id, preview, subject=subject
        )
    if not pdfs:
        if link_hold:
            inv = ""
            match = re.search(r"invoice\s+([A-Z0-9][A-Z0-9-]{2,})", subject, flags=re.I)
            if match:
                inv = match.group(1)
            bill = _stamp_bill(
                {
                    "invoice_number": inv,
                    "date": None,
                    "po": None,
                    "amount": None,
                    "hold_reason": "pdf-behind-link",
                    "pdf_behind_link": True,
                    "pdf_link_url": link_hold.get("url"),
                    "browser_tried": link_hold.get("browser_tried"),
                    "browser_failure": link_hold.get("browser_failure"),
                    "lines": [],
                    "fees": [],
                },
                message,
                vendor_key,
                "",
            )
            return [bill], notes
        notes.append(
            {
                "vendor": vendor_key,
                "reason": "no-pdf",
                "subject": subject[:180],
                "received": message.get("receivedDateTime"),
                "from": from_addr,
            }
        )
        return [], notes
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    chosen: list[dict[str, Any]] = []
    for filename, content in pdfs:
        kind = classify_attachment(filename=filename)
        dest = PDF_DIR / f"{_safe_filename(str(message.get('receivedDateTime') or '')[:10])}_{_safe_filename(filename)}"
        if dest.exists():
            dest = PDF_DIR / f"{len(chosen)}_{dest.name}"
        dest.write_bytes(content)
        parsed = parse_invoice_pdf(dest, subject=subject, from_name=from_name, from_address=from_addr)
        parsed_kind = str(parsed.get("attachment_class") or kind)
        if (
            parsed_kind in NON_INVOICE_ATTACHMENT_KINDS
            or parsed.get("is_purchase_order_doc")
            or parsed.get("is_receipt_scan_doc")
            or parsed.get("is_statement_doc")
            or kind in NON_INVOICE_ATTACHMENT_KINDS
        ):
            notes.append(
                {
                    "vendor": vendor_key,
                    "reason": "not-an-invoice",
                    "kind": parsed_kind or kind,
                    "filename": filename,
                    "subject": subject[:180],
                }
            )
            continue
        siblings = list(parsed.pop("siblings", []) or [])
        for bill in [parsed, *siblings]:
            number = invoice_number_key(bill.get("invoice_number"))
            if not number and bill.get("amount") in (None, ""):
                notes.append(
                    {
                        "vendor": vendor_key,
                        "reason": "unparsed-pdf",
                        "filename": filename,
                        "subject": subject[:180],
                    }
                )
                continue
            if number in AQPC_TOO_OLD_INVOICES and vendor_key == "aqpc":
                notes.append(
                    {
                        "vendor": vendor_key,
                        "reason": "aqpc-voided-too-old",
                        "invoice": number,
                        "subject": subject[:180],
                    }
                )
                continue
            if not in_scope(received=message.get("receivedDateTime"), invoice_date=bill.get("date")):
                notes.append(
                    {
                        "vendor": vendor_key,
                        "reason": "before-aug-1",
                        "invoice": number,
                        "date": bill.get("date"),
                        "received": message.get("receivedDateTime"),
                    }
                )
                continue
            chosen.append(_stamp_bill(bill, message, vendor_key, str(bill.get("pdf_path") or dest)))
    chosen = _prefer_printed_invoice_aliases(chosen)
    return chosen, notes


def index_kimco(client: KimcoClient) -> dict[str, dict[str, int]]:
    """Index by list-view Vendor_$_Display_Name. A fields=Vendor request drops the name."""
    found: dict[str, dict[str, int]] = {key: {} for key in VENDORS}
    try:
        items = client.list_items("ap_invoices", fields="Invoice_Number,Vendor_$_Display_Name")
    except KimcoError:
        items = client.list_items("ap_invoices")
    for item in items:
        vals = item.get("values") or {}
        vendor_id = lookup_id(vals.get("Vendor"))
        vendor_txt = str(lookup_text(vals.get("Vendor") or vals.get("Vendor_$_Display_Name")) or "")
        key = vendor_text_key(vendor_txt, vendor_id)
        if key not in found:
            continue
        number = invoice_number_key(vals.get("Invoice_Number"))
        if number and item.get("id") not in (None, ""):
            found[key][number] = int(item["id"])
    return found


def choose_batch(client: KimcoClient, vendor_key: str) -> dict[str, Any]:
    spec = VENDORS[vendor_key]
    try:
        record = client.get_item("ap_batches", int(spec["batch_id"]))
    except KimcoError as exc:
        record = {"error": str(exc)[:200]}
    vals = record.get("values") or {}
    name = str(vals.get("AP_Invoice_Batch_ID") or "")
    status = vals.get("Status")
    open_batch = name == spec["batch_name"] and status in (0, "0", None)
    if open_batch:
        return {
            "id": spec["batch_id"],
            "name": name,
            "status": status,
            "created": False,
            "source": "preferred-open",
        }
    return {
        "id": None,
        "name": spec["new_batch"],
        "status": status,
        "preferred_name": name,
        "preferred_id": spec["batch_id"],
        "created": None,
        "source": "create-new-dated",
        "preferred_open": False,
    }


def _po_text(bill: dict[str, Any]) -> str:
    pos = [str(p) for p in (bill.get("pos") or []) if p]
    if len(pos) > 1:
        return ", ".join(pos)
    if bill.get("po"):
        return str(bill.get("po"))
    return ""


def sheet_row_from_enter(vendor_key: str, bill: dict[str, Any], row: dict[str, Any]) -> dict[str, Any]:
    return {
        "Vendor": VENDORS[vendor_key]["label"],
        "Invoice #": row.get("Invoice #") or bill.get("invoice_number"),
        "date": row.get("date") or bill.get("date"),
        "PO": row.get("PO") or _po_text(bill),
        "Amount": row.get("Amount") if row.get("Amount") not in (None, "") else bill.get("amount"),
        "Result": row.get("Result"),
        "Why": row.get("Why"),
        "Exception category": row.get(COL_EXCEPTION_CATEGORY) or row.get("Exception category") or "",
        "Exception owner": row.get(COL_EXCEPTION_OWNER) or row.get("Exception owner") or "",
        "KIMCO id": row.get("KIMCO id") or "",
        "Batch": row.get("Batch") or "",
        "Outlook category": "",
        "Notes": row.get("Notes") or "",
        "_message_id": bill.get("graph_message_id"),
        "_vendor_key": vendor_key,
        "_pdf_path": bill.get("pdf_path"),
        "_hold_reason": bill.get("hold_reason"),
    }


def post_comments_1(client: KimcoClient, kimco_id: int, html: str, *, mention_id: int | None) -> dict[str, Any]:
    if not html or not kimco_id:
        return {"status": "skipped"}
    base = {
        "HtmlValue": html,
        "Entity": {"id": 203},
        "ObjectId": int(kimco_id),
        "FormId": 218,
    }
    attempts = [base]
    if mention_id:
        attempts.insert(0, {**base, "Mention": {"id": int(mention_id)}})
    last: dict[str, Any] = {"status": "blocked"}
    for values in attempts:
        payload = {
            "state": "Modified",
            "id": int(kimco_id),
            "lists": {"Comments_1": [{"state": "Added", "values": values}]},
        }
        try:
            _body, status, error = client.update("ap_invoices", int(kimco_id), payload)
        except KimcoError as exc:
            last = {"status": "blocked", "error": str(exc)[:180]}
            continue
        last = {"status": "put" if status < 400 else f"put-{status}", "http": status, "error": error}
        if status < 400:
            return last
    return last


def shawn_price_html(vendor_label: str, invoice: str, po: str) -> str:
    return (
        f'<p><span data-mention-id="104" data-mention-name="Shawn McKibben" '
        f'data-mention-email="{SHAWN_HTML_EMAIL}" class="prosemirror-mention-node">'
        f"@Shawn McKibben</span> HOLD price_variance on {vendor_label} invoice {invoice} "
        f"PO {po or 'n/a'}. |PPV| is $75 or more. Over-gate receipts were not selected. "
        f"No email.</p>"
    )


def apply_hold_followthrough(client: KimcoClient, row: dict[str, Any]) -> dict[str, Any]:
    """Transfer AP + Comments_1. Does not invent Success. Does not finish Transfer AP."""
    result = str(row.get("Result") or "")
    why = str(row.get("Why") or "")
    category = str(row.get("Exception category") or "")
    kid = row.get("KIMCO id")
    if result not in {"HOLD", "Incomplete"} or kid in (None, ""):
        return row
    vendor_key = row.get("_vendor_key") or ""
    spec = VENDORS.get(vendor_key) or {}
    why_l = why.lower()
    price = category == "price_variance" or "price does not match" in why_l or "price_variance" in why_l
    missing = category == "missing_receipt" or "no receipts" in why_l or "missing_receipt" in why_l
    if price and "missing_receipt" not in why_l:
        moved = apply_transfer_ap_batch_move(client, kimco_id=int(kid))
        row["Batch"] = f"Transfer AP ({moved.get('batch_id')})" if moved.get("batch_id") else row.get("Batch")
        row["Exception category"] = "price_variance"
        row["Exception owner"] = "Shawn McKibben"
        comment = post_comments_1(
            client,
            int(kid),
            shawn_price_html(spec.get("label") or vendor_key, str(row.get("Invoice #") or ""), str(row.get("PO") or "")),
            mention_id=104,
        )
        row["Notes"] = f"{row.get('Notes') or ''} Transfer AP {moved.get('status')}. Comments_1 {comment.get('status')}.".strip()
        row["Why"] = f"{why} Over-PPV → Transfer AP ({moved.get('status')}) + @Shawn McKibben mention-id 104.".strip()
        return row
    if missing:
        # cli already moves missing_receipt. Ensure the move, then owner-specific comment.
        moved = apply_transfer_ap_batch_move(client, kimco_id=int(kid))
        if moved.get("batch_id"):
            row["Batch"] = f"Transfer AP ({moved.get('batch_id')})"
        owner = missing_receipt_exception_owner(spec.get("canonical"))
        row["Exception category"] = "missing_receipt"
        row["Exception owner"] = owner
        if should_tag_missing_receipt(spec.get("canonical")):
            html = missing_receipt_comments_1_html(spec.get("canonical"))
            comment = post_comments_1(client, int(kid), html, mention_id=None)
            row["Notes"] = (
                f"{row.get('Notes') or ''} missing_receipt Transfer AP {moved.get('status')}. "
                f"Comments_1 {comment.get('status')} owner {owner}."
            ).strip()
        else:
            row["Notes"] = (
                f"{row.get('Notes') or ''} missing_receipt Transfer AP {moved.get('status')}. "
                "Receiving owner blank — no @tag."
            ).strip()
        row["Why"] = f"{why} NOTE-53 missing_receipt → Transfer AP ({moved.get('status')}).".strip()
    return row


def gas_lines_if_needed(client: KimcoClient, bill: dict[str, Any], row: dict[str, Any], cache: dict[str, Any]) -> dict[str, Any]:
    if row.get("_vendor_key") != "gas":
        return row
    if str(row.get("PO") or "").strip():
        return row
    kid = row.get("KIMCO id")
    if kid in (None, ""):
        return row
    if str(row.get("Result") or "") not in {"Success", "Incomplete", "HOLD"}:
        return row
    try:
        record = client.get_item("ap_invoices", int(kid))
    except KimcoError as exc:
        row["Notes"] = f"{row.get('Notes') or ''} Lines-K GET failed {exc}.".strip()
        return row
    vals = record.get("values") or {}
    if int(vals.get("Invoice_Type") or 0) != 4:
        return row
    existing = misc_line_snapshot(record)
    proof = {"invoice_type": 4, "misc_lines": existing}
    try:
        posted = apply_type4_misc_lines(
            client,
            parsed=bill,
            kimco_id=int(kid),
            vendor_id=71,
            proof=proof,
            lookup_cache=cache,
        )
    except (KimcoError, TypeError, ValueError) as exc:
        posted = {"status": "blocker", "why": str(exc)[:180]}
    row["Notes"] = f"{row.get('Notes') or ''} Lines-K {posted.get('status')} {posted.get('why') or ''}".strip()
    if str(row.get("Result")) == "Success" and posted.get("status") not in {"added", "already-ok", "skip-not-type4"}:
        row["Result"] = "Incomplete"
        row["Why"] = (
            f"{row.get('Why') or ''} Gas Misc Type 4 Lines-K not confirmed "
            f"({posted.get('status')}). Not Success."
        ).strip()
        row["Exception category"] = row.get("Exception category") or "other"
        row["Exception owner"] = row.get("Exception owner") or "AP clerk"
    elif posted.get("status") in {"added", "already-ok"} and str(row.get("Result")) == "Success":
        try:
            after = client.get_item("ap_invoices", int(kid))
        except KimcoError:
            after = record
        lines = misc_line_snapshot(after)
        if not type4_shop_supplies_lines_ok(lines):
            row["Result"] = "Incomplete"
            row["Why"] = f"{row.get('Why') or ''} Lines-K empty after post. Not Success.".strip()
    return row


def outlook_category_for(rows: list[dict[str, Any]]) -> str:
    results = [str(r.get("Result") or "") for r in rows]
    if results and all(r == "Success" or r == "already in KIMCO" for r in results):
        # already-in-KIMCO alone is handled by the caller with a live check.
        if all(r == "Success" for r in results):
            return ENTERED_IN_AI_CATEGORY
    if any(r in {"Success", "HOLD", "Incomplete", "already in KIMCO"} and row.get("KIMCO id") not in (None, "") for r, row in zip(results, rows)):
        if all(r == "Success" for r in results):
            return ENTERED_IN_AI_CATEGORY
        return ENTERED_WITH_ISSUES_CATEGORY
    return AI_HOLD_CATEGORY


def stamp_parent(graph, message_id: str, category: str, *, fort_worth_id: str | None, already_folder: str | None) -> dict[str, Any]:
    if category == ENTERED_IN_AI_CATEGORY:
        flag = graph.flag_matched(ALLOWED_MAILBOX, message_id)
    elif category == ENTERED_WITH_ISSUES_CATEGORY:
        flag = graph.flag_issues(ALLOWED_MAILBOX, message_id)
    elif category == AI_HOLD_CATEGORY:
        flag = graph.flag_hold(ALLOWED_MAILBOX, message_id)
    else:
        flag = "skipped"
    moved = "not-moved"
    if fort_worth_id and category in {ENTERED_IN_AI_CATEGORY, ENTERED_WITH_ISSUES_CATEGORY}:
        if already_folder and is_fort_worth_inbox_folder(already_folder):
            moved = "already-fort-worth"
        else:
            try:
                result = graph.move_message(ALLOWED_MAILBOX, message_id, fort_worth_id)
                moved = str(result.get("status") or result)
            except GraphError as exc:
                moved = f"move-failed {exc}"[:160]
    return {"flag": flag, "category": category, "move": moved}


def kimco_outlook_category(client: KimcoClient, kimco_id: int) -> str:
    """Entered in AI only when the live header already has lines. Else issues."""
    try:
        record = client.get_item("ap_invoices", int(kimco_id))
    except KimcoError:
        return ENTERED_WITH_ISSUES_CATEGORY
    vals = record.get("values") or {}
    batch = str(lookup_text(vals.get("AP_Invoice_Batch")) or "")
    lines = ((record.get("lists") or {}).get("APInvoiceLine") or [])
    if "transfer ap" in batch.lower():
        return ENTERED_WITH_ISSUES_CATEGORY
    if lines:
        return ENTERED_IN_AI_CATEGORY
    return ENTERED_WITH_ISSUES_CATEGORY


def write_xlsx(path: Path, rows: list[dict[str, Any]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "AP gap"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for col, name in enumerate(SHEET_COLUMNS, start=1):
        cell = sheet.cell(1, col, name)
        cell.font = header_font
        cell.fill = header_fill
    for ridx, row in enumerate(rows, start=2):
        for cidx, name in enumerate(SHEET_COLUMNS, start=1):
            value = row.get(name, "")
            cell = sheet.cell(ridx, cidx, "" if value is None else value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate([22, 18, 12, 16, 12, 18, 60, 20, 24, 12, 28, 22, 40], start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:M{max(1, len(rows) + 1)}"
    book.save(path)


def public_row(row: dict[str, Any]) -> dict[str, Any]:
    return {key: row.get(key, "") for key in SHEET_COLUMNS}


def build_proof(rows: list[dict[str, Any]], *, discovery: dict[str, Any], skipped: list[dict[str, Any]], batches: dict[str, Any]) -> dict[str, Any]:
    by_vendor: dict[str, dict[str, Any]] = {}
    for key, spec in VENDORS.items():
        mine = [r for r in rows if r.get("Vendor") == spec["label"]]
        entered = [r for r in mine if r.get("Result") not in {"already in KIMCO", ""} and r.get("KIMCO id") not in (None, "")]
        # newly entered = this run created them. already rows are tagged.
        new_rows = [r for r in mine if r.get("_new")]
        by_vendor[key] = {
            "emails": (discovery.get("email_counts") or {}).get(key, 0),
            "invoice_rows": len(mine),
            "already_in_kimco": sum(1 for r in mine if r.get("Result") == "already in KIMCO"),
            "newly_entered_success": sum(1 for r in new_rows if r.get("Result") == "Success"),
            "newly_entered_hold": sum(1 for r in new_rows if r.get("Result") == "HOLD"),
            "newly_entered_incomplete": sum(1 for r in new_rows if r.get("Result") == "Incomplete"),
            "newly_entered_fail": sum(1 for r in new_rows if r.get("Result") == "Fail"),
            "remaining_unprocessed": [
                r.get("Invoice #") for r in mine if r.get("Result") == "remaining"
            ],
        }
    return {
        "proof": "ap-gap-oneal-aqpc-legacy-gas-2026-09-23",
        "invent": False,
        "mail_send": False,
        "mailbox": ALLOWED_MAILBOX,
        "ppv_gate": "absolute dollars; |PPV| >= 75 HOLD price_variance; under 75 may post",
        "packing_slip_success_gate": "suspended-2026-09-23",
        "discovery": discovery,
        "batches": batches,
        "already_in_kimco": [public_row(r) for r in rows if r.get("Result") == "already in KIMCO"],
        "newly_entered": [public_row(r) for r in rows if r.get("_new")],
        "holds": [public_row(r) for r in rows if r.get("Result") == "HOLD"],
        "skipped_reasons": skipped,
        "remaining_unprocessed": [r.get("Invoice #") for r in rows if r.get("Result") == "remaining"],
        "per_vendor": by_vendor,
        "rows": [public_row(r) for r in rows],
    }


def save_outputs(proof: dict[str, Any], rows: list[dict[str, Any]]) -> None:
    text = json.dumps(proof, indent=2, default=str) + "\n"
    for path in (ARTIFACT_JSON, RUN_JSON):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text)
    public = [public_row(r) for r in rows]
    write_xlsx(ARTIFACT_XLSX, public)
    write_xlsx(RUN_XLSX, public)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--discover-only", action="store_true")
    parser.add_argument("--wave", type=int, default=0, help="Max new invoices per vendor. 0 = all.")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")
    if not args.live and not args.discover_only:
        raise SystemExit("Refusing: pass --live or --discover-only.")
    configure_ppv_absolute()
    cli_mod.apply_fort_worth_move_after_enter = defer_fort_worth_move

    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error or "live credentials not ready")
    client = KimcoClient.authenticate(
        creds.instance_url, creds.key or "", creds.password or "", target="live"
    )
    graph = _optional_graph_client()
    if graph is None:
        raise SystemExit("Graph client required")

    LOGGER.info("Indexing live KIMCO AP invoices")
    kimco_index = index_kimco(client)
    LOGGER.info(
        "KIMCO counts %s",
        {key: len(nums) for key, nums in kimco_index.items()},
    )
    batches = {key: choose_batch(client, key) for key in VENDORS}
    LOGGER.info("Collecting mailbox messages")
    messages, mail_meta = collect_messages(graph)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for msg in messages:
        key = classify_vendor(
            subject=str(msg.get("subject") or ""),
            preview=str(msg.get("bodyPreview") or ""),
            from_name=sender_name(msg),
            from_addr=sender_address(msg),
        )
        if key:
            grouped[key].append(msg)
    email_counts = {key: len(grouped.get(key) or []) for key in VENDORS}
    LOGGER.info("Vendor emails %s", email_counts)

    graph_mod.FORT_WORTH_FOLDER_ID = None
    fort = graph.resolve_fort_worth_folder(ALLOWED_MAILBOX)
    fort_id = fort.get("id")

    rows: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    pending: dict[str, list[dict[str, Any]]] = defaultdict(list)
    seen_invoice: set[tuple[str, str]] = set()

    for vendor_key in VENDORS:
        for msg in grouped.get(vendor_key) or []:
            subject = str(msg.get("subject") or "")
            if NOISE_SUBJECT.search(subject) and "invoice" not in subject.lower():
                skipped.append(
                    {
                        "vendor": vendor_key,
                        "reason": "not-an-invoice",
                        "subject": subject[:180],
                        "received": msg.get("receivedDateTime"),
                    }
                )
                continue
            bills, notes = bills_from_message(graph, msg, vendor_key)
            skipped.extend(notes)
            if not bills and not notes:
                skipped.append(
                    {
                        "vendor": vendor_key,
                        "reason": "no-bill",
                        "subject": str(msg.get("subject") or "")[:180],
                        "received": msg.get("receivedDateTime"),
                    }
                )
            for bill in bills:
                number = invoice_number_key(bill.get("invoice_number"))
                dedupe = (vendor_key, number)
                if number and dedupe in seen_invoice:
                    continue
                if number:
                    seen_invoice.add(dedupe)
                existing = kimco_index.get(vendor_key, {}).get(number) if number else None
                if existing:
                    cats = list(bill.get("categories") or [])
                    row = {
                        "Vendor": VENDORS[vendor_key]["label"],
                        "Invoice #": number,
                        "date": bill.get("date"),
                        "PO": _po_text(bill),
                        "Amount": bill.get("amount"),
                        "Result": "already in KIMCO",
                        "Why": f"Live KIMCO id {existing} already has this vendor invoice #. Not re-entered.",
                        "Exception category": "",
                        "Exception owner": "",
                        "KIMCO id": existing,
                        "Batch": "",
                        "Outlook category": ", ".join(cats),
                        "Notes": "",
                        "_message_id": bill.get("graph_message_id"),
                        "_vendor_key": vendor_key,
                        "_new": False,
                        "_folder": bill.get("folder"),
                        "_needs_category": not has_process_category({"categories": cats}),
                    }
                    rows.append(row)
                    continue
                if bill.get("hold_reason") == "pdf-behind-link":
                    row = {
                        "Vendor": VENDORS[vendor_key]["label"],
                        "Invoice #": number,
                        "date": bill.get("date"),
                        "PO": _po_text(bill),
                        "Amount": bill.get("amount"),
                        "Result": "HOLD",
                        "Why": (
                            "HOLD pdf-behind-link. Guest/unauth PDF download failed. "
                            f"browser_tried={bill.get('browser_tried')} "
                            f"browser_failure={bill.get('browser_failure')} "
                            f"url={bill.get('pdf_link_url')}."
                        ),
                        "Exception category": "pdf-behind-link",
                        "Exception owner": "AP clerk",
                        "KIMCO id": "",
                        "Batch": "",
                        "Outlook category": "",
                        "Notes": "AQPC portal" if vendor_key == "aqpc" else "",
                        "_message_id": bill.get("graph_message_id"),
                        "_vendor_key": vendor_key,
                        "_new": False,
                        "_folder": bill.get("folder"),
                        "_category_target": AI_HOLD_CATEGORY,
                    }
                    rows.append(row)
                    continue
                amount_value = money(bill.get("amount"))
                if amount_value is not None and amount_value < 0:
                    rows.append(
                        {
                            "Vendor": VENDORS[vendor_key]["label"],
                            "Invoice #": number,
                            "date": bill.get("date"),
                            "PO": _po_text(bill),
                            "Amount": bill.get("amount"),
                            "Result": "HOLD",
                            "Why": "Credit memo / negative amount. Not entered as a payable. Amount not invented.",
                            "Exception category": "other",
                            "Exception owner": "AP clerk",
                            "KIMCO id": "",
                            "Batch": "",
                            "Outlook category": "",
                            "Notes": "credit",
                            "_message_id": bill.get("graph_message_id"),
                            "_vendor_key": vendor_key,
                            "_new": False,
                            "_folder": bill.get("folder"),
                            "_category_target": AI_HOLD_CATEGORY,
                        }
                    )
                    continue
                if bill.get("amount") in (None, "") or not number or not bill.get("date"):
                    row = {
                        "Vendor": VENDORS[vendor_key]["label"],
                        "Invoice #": number,
                        "date": bill.get("date"),
                        "PO": _po_text(bill),
                        "Amount": bill.get("amount"),
                        "Result": "HOLD",
                        "Why": "HOLD parse-error. Invoice #, date, or after-tax amount missing from the PDF. Not invented.",
                        "Exception category": "parse-error",
                        "Exception owner": "AP clerk",
                        "KIMCO id": "",
                        "Batch": "",
                        "Outlook category": "",
                        "Notes": "",
                        "_message_id": bill.get("graph_message_id"),
                        "_vendor_key": vendor_key,
                        "_new": False,
                        "_folder": bill.get("folder"),
                        "_category_target": AI_HOLD_CATEGORY,
                    }
                    rows.append(row)
                    continue
                pending[vendor_key].append(bill)

    discovery = {
        "mail": mail_meta,
        "email_counts": email_counts,
        "kimco_invoice_counts": {key: len(nums) for key, nums in kimco_index.items()},
        "pending_new": {key: [invoice_number_key(b.get("invoice_number")) for b in pending[key]] for key in VENDORS},
        "fort_worth": {"id_present": bool(fort_id), "name": fort.get("displayName")},
    }
    proof = build_proof(rows, discovery=discovery, skipped=skipped, batches=batches)
    save_outputs(proof, rows)
    LOGGER.info("Discovery saved pending=%s", {k: len(v) for k, v in pending.items()})
    if args.discover_only or not args.live:
        print(json.dumps({"pending": discovery["pending_new"], "emails": email_counts, "invent": False}, indent=2))
        return 0

    gas_cache: dict[str, Any] = {}
    for vendor_key, bills in pending.items():
        cap = args.wave if args.wave and args.wave > 0 else len(bills)
        chosen = bills[:cap]
        rest = bills[cap:]
        for bill in rest:
            rows.append(
                {
                    "Vendor": VENDORS[vendor_key]["label"],
                    "Invoice #": invoice_number_key(bill.get("invoice_number")),
                    "date": bill.get("date"),
                    "PO": _po_text(bill),
                    "Amount": bill.get("amount"),
                    "Result": "remaining",
                    "Why": "Wave cap reached. Not entered this run.",
                    "Exception category": "",
                    "Exception owner": "",
                    "KIMCO id": "",
                    "Batch": "",
                    "Outlook category": "",
                    "Notes": "",
                    "_vendor_key": vendor_key,
                    "_new": False,
                    "_message_id": bill.get("graph_message_id"),
                }
            )
        if not chosen:
            continue
        batch_name = batches[vendor_key]["name"]
        LOGGER.info("Entering %s bills on %s", len(chosen), batch_name)
        entered = run_enter(
            client,
            chosen,
            batch_name=batch_name,
            pdf_dir=PDF_DIR,
            graph_client=graph,
            mailbox=ALLOWED_MAILBOX,
            flag_outlook=False,
        )
        for bill, enter_row in zip(chosen, entered):
            row = sheet_row_from_enter(vendor_key, bill, enter_row)
            row["_new"] = True
            row["_folder"] = bill.get("folder")
            row = gas_lines_if_needed(client, bill, row, gas_cache)
            row = apply_hold_followthrough(client, row)
            number = invoice_number_key(row.get("Invoice #"))
            if row.get("KIMCO id") not in (None, "") and number:
                kimco_index[vendor_key][number] = int(row["KIMCO id"])
            rows.append(row)
        proof = build_proof(rows, discovery=discovery, skipped=skipped, batches=batches)
        save_outputs(proof, rows)

    # Outlook category + Fort Worth, once per parent message.
    by_message: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        mid = str(row.get("_message_id") or "")
        if mid:
            by_message[mid].append(row)
    for mid, group in by_message.items():
        new_ones = [r for r in group if r.get("_new") or r.get("_category_target") or r.get("_needs_category")]
        if not new_ones:
            continue
        if any(r.get("_new") or r.get("_category_target") for r in group):
            targets = [r.get("_category_target") for r in group if r.get("_category_target")]
            if targets and not any(r.get("_new") for r in group):
                category = targets[0]
            else:
                fresh = [r for r in group if r.get("_new") or r.get("Result") in {"Success", "HOLD", "Incomplete", "Fail"}]
                category = outlook_category_for(fresh or group)
                if any(str(r.get("Result")) == "HOLD" and not r.get("KIMCO id") for r in group):
                    if not any(r.get("KIMCO id") for r in group if r.get("_new")):
                        category = AI_HOLD_CATEGORY
        else:
            # KIMCO exists, Outlook had no process category.
            kid = next((r.get("KIMCO id") for r in group if r.get("KIMCO id")), None)
            category = kimco_outlook_category(client, int(kid)) if kid else ENTERED_WITH_ISSUES_CATEGORY
        folder = next((r.get("_folder") for r in group if r.get("_folder")), None)
        # Move only after a header we created or an already-entered bill with lines category.
        do_move = category in {ENTERED_IN_AI_CATEGORY, ENTERED_WITH_ISSUES_CATEGORY} and any(
            r.get("KIMCO id") for r in group
        )
        stamped = stamp_parent(
            graph,
            mid,
            category,
            fort_worth_id=fort_id if do_move else None,
            already_folder=str(folder or ""),
        )
        for row in group:
            row["Outlook category"] = stamped.get("category") or row.get("Outlook category")
            row["Notes"] = f"{row.get('Notes') or ''} Outlook {stamped.get('flag')} move {stamped.get('move')}.".strip()

    proof = build_proof(rows, discovery=discovery, skipped=skipped, batches=batches)
    save_outputs(proof, rows)
    print(
        json.dumps(
            {
                "invent": False,
                "per_vendor": proof["per_vendor"],
                "newly_entered": [
                    {
                        "vendor": r.get("Vendor"),
                        "invoice": r.get("Invoice #"),
                        "kimco_id": r.get("KIMCO id"),
                        "result": r.get("Result"),
                    }
                    for r in proof["newly_entered"]
                ],
            },
            indent=2,
            default=str,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
