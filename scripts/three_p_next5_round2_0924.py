"""Enter the next 5 unentered 3P invoices on live batch 724.

Oldest by PDF date, then invoice number. invent=false. No Mail.Send. Do not post.
Do not touch 142231 / KIMCO 10186.

Freight-only: Miscellaneous Invoice_Type 4, blank PO, no lines, one Freight
External charge (lookup id 1) equal to that invoice's own PDF total.

Parts: Invoice_Type 3, header PO blank when the PDF says SEE BELOW (multi-PO).
Select receipts that match part and qty. PPV is invoice total minus selected
receipt extended. |PPV| under $75 is an Additional Charge. $75 or more is
HOLD price_variance, receipts on that line are not selected, then Transfer AP.
"""

from __future__ import annotations

import json
import sys
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from ap_clerk.auth import load_credentials
from ap_clerk.cli import _optional_graph_client
from ap_clerk.graph import ALLOWED_MAILBOX, FORT_WORTH_FOLDER_DISPLAY_NAME, GraphClient
from ap_clerk.kimco import (
    KimcoClient,
    fees_payload,
    receipt_line_values_from_records,
    select_receipts_payload,
)
from ap_clerk.rules import kimco_datetime, lookup_id, lookup_text, money
from ap_clerk.transfer_ap import apply_transfer_ap_batch_move
from gas_59081_retry import SHAWN_MENTION_ID, comments_1_shawn

BATCH_ID = 724
VENDOR_ID = 1
TERMS_ID = 4
CURRENCY_ID = 3
REMIT_ID = 120
DO_NOT_TOUCH = {"142231": 10186}
PPV_HOLD = 75.0

PDF_DIR = Path("/tmp/3p-round2/pdfs")
PRIOR_XLSX = ROOT / "artifacts" / "AP-3P-updated-next5-2026-09-24.xlsx"
OUT_JSON = ROOT / "artifacts" / "3p-next5-round2-2026-09-24.json"
OUT_XLSX = ROOT / "artifacts" / "AP-3P-updated-round2-2026-09-24.xlsx"

# Confirmed by reading each invoice page. Packing lists, emails, and photos
# stay off the bill. Freight pages are the single invoice page in the 9/9 pack.
BILLS = [
    {
        "invoice": "142176",
        "kind": "parts",
        "date": date(2026, 8, 21),
        "due": date(2026, 10, 20),
        "total": 364.92,
        "pdf_name_needle": "142176",
        "pages": [0],
        "pos": ["58952", "58925"],
        "po_on_pdf": "SEE BELOW",
        "lines": [
            {"po": "58952", "part": "21678-1", "qty": 3, "unit": 91.23, "amount": 273.69},
            {"po": "58925", "part": "21678-1", "qty": 1, "unit": 91.23, "amount": 91.23},
        ],
        "receipts": [23429, 23417, 23428],
        "skip_receipts": [],
        "message_subject_needle": "INV # 142176",
    },
    {
        "invoice": "142238",
        "kind": "freight",
        "date": date(2026, 8, 21),
        "due": date(2026, 10, 20),
        "total": 500.00,
        "pdf_name_needle": "142231",
        "pages": [12],
        "pos": [],
        "po_on_pdf": "FREIGHT",
        "freight_for": "142176",
        "pdf_line": "1 FREIGHT FOR INVOICE # 142176 @ 500.00 = 500.00",
        "message_subject_needle": "142238",
    },
    {
        "invoice": "142179",
        "kind": "parts",
        "date": date(2026, 8, 25),
        "due": date(2026, 10, 24),
        "total": 1855.80,
        "pdf_name_needle": "142179",
        "pages": [0],
        "pos": ["58952", "58925"],
        "po_on_pdf": "SEE BELOW",
        "lines": [
            {"po": "58952", "part": "21678-1", "qty": 1, "unit": 91.23, "amount": 91.23},
            {"po": "58925", "part": "10095-1", "qty": 20, "unit": 7.76, "amount": 155.20},
            {"po": "58952", "part": "21912-1", "qty": 7, "unit": 39.35, "amount": 275.45},
            {"po": "58952", "part": "34368-1", "qty": 5, "unit": 39.35, "amount": 196.75},
            {"po": "58952", "part": "14597-1", "qty": 25, "unit": 12.30, "amount": 307.50},
            {"po": "58952", "part": "1020586-1", "qty": 11, "unit": 39.97, "amount": 439.67},
            {"po": "58925", "part": "1007044-1", "qty": 4, "unit": 97.50, "amount": 390.00},
        ],
        "receipts": [23547, 23521, 23522, 23420, 23421, 23422, 23423, 23418, 23419],
        "skip_receipts": [
            {
                "id": 23414,
                "part": "1007044-1",
                "po": "58925",
                "qty": 4,
                "po_unit": 361.46,
                "po_ext": 1445.84,
                "invoice_unit": 97.50,
                "invoice_ext": 390.00,
            }
        ],
        "message_subject_needle": "INV # 142179",
    },
    {
        "invoice": "142239",
        "kind": "freight",
        "date": date(2026, 8, 25),
        "due": date(2026, 10, 24),
        "total": 500.00,
        "pdf_name_needle": "142231",
        "pages": [14],
        "pos": [],
        "po_on_pdf": "FREIGHT",
        "freight_for": "142179",
        "pdf_line": "1 FREIGHT FOR INVOICE # 142179 @ 500.00 = 500.00",
        "message_subject_needle": "142239",
    },
    {
        "invoice": "142188",
        "kind": "parts",
        "date": date(2026, 8, 27),
        "due": date(2026, 10, 26),
        "total": 2553.27,
        "pdf_name_needle": "142188",
        "pages": [0],
        "pos": ["58972", "59002", "58952", "58925"],
        "po_on_pdf": "SEE BELOW",
        "lines": [
            {"po": "58972", "part": "21678-1", "qty": 5, "unit": 91.23, "amount": 456.15},
            {"po": "59002", "part": "21678-1", "qty": 3, "unit": 91.23, "amount": 273.69},
            {"po": "59002", "part": "15058-1", "qty": 32, "unit": 3.65, "amount": 116.80},
            {"po": "58972", "part": "14597-1", "qty": 23, "unit": 12.30, "amount": 282.90},
            {"po": "58952", "part": "1020586-1", "qty": 1, "unit": 39.97, "amount": 39.97},
            {"po": "59002", "part": "1008270-1", "qty": 6, "unit": 83.02, "amount": 498.12},
            {"po": "58925", "part": "1007044-1", "qty": 2, "unit": 97.50, "amount": 195.00},
            {"po": "58925", "part": "10095-1", "qty": 89, "unit": 7.76, "amount": 690.64},
        ],
        "receipts": [23531, 23532, 23537, 23535, 23536, 23540, 23541, 23542, 23546, 23538, 23539, 23544, 23545],
        "skip_receipts": [
            {
                "id": 23543,
                "part": "1007044-1",
                "po": "58925",
                "qty": 2,
                "po_unit": 361.46,
                "po_ext": 722.92,
                "invoice_unit": 97.50,
                "invoice_ext": 195.00,
            }
        ],
        "message_subject_needle": "INV # 142188",
    },
]


def find_pdf(needle: str) -> Path:
    hits = sorted(PDF_DIR.glob(f"*{needle}*.pdf"))
    if len(hits) != 1:
        raise SystemExit(f"expected one PDF for {needle}, found {hits}")
    return hits[0]


def page_bytes(pdf_path: Path, page_indexes: list[int]) -> bytes:
    reader = PdfReader(str(pdf_path))
    writer = PdfWriter()
    for index in page_indexes:
        writer.add_page(reader.pages[index])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def existing_3p(client: KimcoClient) -> dict[str, int]:
    found: dict[str, int] = {}
    url = client._url("ap_invoices")
    offset = 0
    total = None
    while total is None or offset < total:
        response = client.request("GET", url, params={"pageSize": 2000, "offset": offset})
        if response.status_code != 200:
            raise SystemExit(f"AP list HTTP {response.status_code}")
        payload = response.json()
        items = payload.get("items") or []
        total = int(payload.get("totalCount") or 0)
        for item in items:
            values = item.get("values") or {}
            vendor = str(values.get("Vendor_$_Display_Name") or "")
            number = str(values.get("Invoice_Number") or "").strip()
            if "3P" not in vendor.upper():
                continue
            if number and item.get("id") not in (None, ""):
                found[number] = int(item["id"])
        if not items:
            break
        offset += len(items)
    return found


def snapshot(client: KimcoClient, kimco_id: int) -> dict[str, Any]:
    rec = client.get_item("ap_invoices", kimco_id)
    values = rec.get("values") or {}
    lists = rec.get("lists") or {}
    lines = []
    for line in lists.get("APInvoiceLine") or []:
        lv = line.get("values") or {}
        lines.append(
            {
                "id": line.get("id"),
                "qty": lv.get("Quantity"),
                "price": money(lv.get("Unit_Price")),
                "ext": money(lv.get("Extended_Amount")),
                "receipt": lookup_id(lv.get("Receipt")),
                "po": lookup_text(lv.get("Purchase_Order_Number")),
                "po_line": lookup_text(lv.get("Purchase_Order_Line")),
            }
        )
    charges = []
    for charge in lists.get("InvoiceAdditionalCharges") or []:
        cv = charge.get("values") or {}
        charges.append(
            {
                "id": charge.get("id"),
                "kind": cv.get("Additional_Charges"),
                "amount": money(cv.get("Amount")),
                "price": money(cv.get("Price")),
                "qty": cv.get("Quantity"),
            }
        )
    comments = []
    for comment in lists.get("Comments_1") or []:
        comments.append(
            {
                "id": comment.get("id"),
                "html": ((comment.get("values") or {}).get("HtmlValue") or "")[:1200],
            }
        )
    try:
        atts = client.list_attachments(kimco_id)
    except Exception:
        atts = []
    return {
        "id": kimco_id,
        "invoice": values.get("Invoice_Number"),
        "type": values.get("Invoice_Type"),
        "date": values.get("Invoice_Date"),
        "due": values.get("Invoice_Due_Date"),
        "amount": money(values.get("Invoice_Amount")),
        "verification": money(values.get("Invoice_Verification_Amount")),
        "status": values.get("Status"),
        "posted": values.get("Posted"),
        "po": lookup_text(values.get("Purchase_Order")),
        "po_id": lookup_id(values.get("Purchase_Order")),
        "batch": lookup_text(values.get("AP_Invoice_Batch")),
        "batch_id": lookup_id(values.get("AP_Invoice_Batch")),
        "vendor": lookup_text(values.get("Vendor")),
        "lines": lines,
        "charges": charges,
        "comments_1": comments,
        "attachments": [a.get("name") for a in atts],
    }


def header_payload(spec: dict[str, Any]) -> dict[str, Any]:
    invoice_type = 4 if spec["kind"] == "freight" else 3
    return {
        "AP_Invoice_Batch": {"id": BATCH_ID},
        "Vendor": {"id": VENDOR_ID},
        "Invoice_Number": spec["invoice"],
        "Invoice_Type": invoice_type,
        "Invoice_Date": kimco_datetime(spec["date"]),
        "Invoice_Verification_Amount": float(spec["total"]),
        "Invoice_Due_Date": kimco_datetime(spec["due"]),
        "Terms_Code": {"id": TERMS_ID},
        "Currency": {"id": CURRENCY_ID},
        "Remit_To_Address": {"id": REMIT_ID},
        "Transaction_Date": kimco_datetime(spec["date"]),
        "Comments": "API Agent",
    }


def assert_receipts_open(client: KimcoClient, receipt_ids: list[int]) -> None:
    for rid in receipt_ids:
        rec = client.get_item("receipts", int(rid))
        values = rec.get("values") or {}
        ap = str(values.get("AP_Invoice_Number") or "").strip()
        if values.get("Invoiced") is True or ap:
            raise SystemExit(f"receipt {rid} already billed on {ap or 'an invoice'}")
        qty = money(values.get("Quantity_Received"))
        if qty is None or qty <= 0:
            raise SystemExit(f"receipt {rid} qty {qty} is not selectable")


def select_open_receipts(client: KimcoClient, kimco_id: int, receipt_ids: list[int]) -> dict[str, Any]:
    assert_receipts_open(client, receipt_ids)
    invoice = client.get_item("ap_invoices", kimco_id)
    lines = []
    for rid in receipt_ids:
        receipt = client.get_item("receipts", int(rid))
        lines.append(receipt_line_values_from_records(invoice, receipt))
    payload = select_receipts_payload(lines, invoice_id=kimco_id)
    put = client.request("PUT", client._record_url("ap_invoices", kimco_id), json=payload)
    return {
        "http": put.status_code,
        "error": "" if put.status_code < 400 else put.text[:600],
        "count": len(receipt_ids),
    }


def price_note(spec: dict[str, Any]) -> str:
    skipped = spec.get("skip_receipts") or []
    if not skipped:
        return ""
    bits = []
    for row in skipped:
        gap = round(float(row["invoice_ext"]) - float(row["po_ext"]), 2)
        bits.append(
            f"qty {row['qty']} of part {row['part']} on PO {row['po']} is "
            f"${row['invoice_unit']:.2f} on the invoice (${row['invoice_ext']:.2f}) "
            f"and receipt {row['id']} is already at PO price ${row['po_unit']:.2f} "
            f"(${row['po_ext']:.2f}). |PPV| ${abs(gap):.2f} is $75 or more, so that "
            "receipt was not selected and the unit price was not changed. "
            f"Unreceive, set PO {row['po']} part {row['part']} to ${row['invoice_unit']:.2f}, "
            f"and re-receive qty {row['qty']}."
        )
    return (
        "AP Clerk: @Shawn McKibben HOLD price_variance on 3P "
        f"{spec['invoice']}. " + " ".join(bits)
    )


def price_html(spec: dict[str, Any]) -> str:
    text = price_note(spec)
    if not text:
        return ""
    mention = (
        f'<span data-mention-id="{SHAWN_MENTION_ID}" '
        'data-mention-name="Shawn McKibben" '
        'data-mention-email="Shawn.McKibben@kannonmfg.com" '
        f'class="prosemirror-mention-node">@Shawn McKibben</span>'
    )
    body = text.replace("@Shawn McKibben", mention, 1)
    return f"<p>{body}</p>"


def archive_email(graph: GraphClient, needle: str) -> dict[str, Any]:
    hits = graph.search_messages(ALLOWED_MAILBOX, needle, top=5)
    chosen = None
    for msg in hits:
        subject = str(msg.get("subject") or "")
        if needle in subject:
            chosen = msg
            break
    if chosen is None and hits:
        chosen = hits[0]
    if chosen is None:
        return {"status": "not-found", "folder": ""}
    message_id = str(chosen.get("id") or "")
    full = graph.get_message(
        ALLOWED_MAILBOX,
        message_id,
        select="id,subject,parentFolderId,categories",
    )
    folder = graph.resolve_fort_worth_folder(ALLOWED_MAILBOX)
    fort_id = str(folder.get("id") or "")
    parent = str(full.get("parentFolderId") or "")
    if fort_id and parent == fort_id:
        return {
            "status": "already",
            "folder": folder.get("displayName") or FORT_WORTH_FOLDER_DISPLAY_NAME,
            "moved_this_run": False,
            "subject": full.get("subject"),
        }
    if not fort_id:
        return {"status": "no-folder", "folder": "", "subject": full.get("subject")}
    moved = graph.move_message(ALLOWED_MAILBOX, message_id, fort_id)
    new_id = str(moved.get("new_id") or message_id)
    after = graph.get_message(
        ALLOWED_MAILBOX,
        new_id,
        select="id,subject,parentFolderId,categories",
    )
    after_parent = str(after.get("parentFolderId") or "")
    in_folder = after_parent == fort_id
    return {
        "status": "moved" if in_folder else moved.get("status"),
        "folder": (folder.get("displayName") or FORT_WORTH_FOLDER_DISPLAY_NAME) if in_folder else "",
        "moved_this_run": bool(in_folder),
        "http": moved.get("http"),
        "subject": after.get("subject") or full.get("subject"),
    }


def enter_one(client: KimcoClient, graph: GraphClient, spec: dict[str, Any]) -> dict[str, Any]:
    number = spec["invoice"]
    if number in DO_NOT_TOUCH:
        raise SystemExit(f"refusing to touch {number}")
    pdf = page_bytes(find_pdf(spec["pdf_name_needle"]), spec["pages"])
    created_id, _body, status, error = client.create("ap_invoices", header_payload(spec))
    out: dict[str, Any] = {
        "invoice": number,
        "kind": spec["kind"],
        "pdf_date": spec["date"].isoformat(),
        "pdf_due": spec["due"].isoformat(),
        "pdf_total": spec["total"],
        "po": None,
        "po_on_pdf": spec.get("po_on_pdf"),
        "pos": spec.get("pos") or [],
        "create_http": status,
        "create_error": error,
        "kimco_id": created_id,
    }
    if created_id is None:
        out["status"] = "Fail"
        out["why"] = f"header create HTTP {status}"
        return out
    if int(created_id) == DO_NOT_TOUCH["142231"]:
        raise SystemExit("refusing to edit KIMCO 10186")
    attach = client.try_official_attach(
        created_id,
        name=f"{number}.pdf",
        content_type="application/pdf",
        size=len(pdf),
        content=pdf,
    )
    out["attach"] = attach
    email = archive_email(graph, spec["message_subject_needle"])
    out["email"] = email
    if spec["kind"] == "freight":
        fee_body = fees_payload(
            [{"amount": spec["total"], "freight_external": True}],
            invoice_id=created_id,
            freight_external=True,
        )
        put = client.request("PUT", client._record_url("ap_invoices", created_id), json=fee_body)
        out["charge_http"] = put.status_code
        out["charge_error"] = "" if put.status_code < 400 else put.text[:400]
        out["freight_for"] = spec.get("freight_for")
        out["pdf_line"] = spec.get("pdf_line")
    else:
        selected = select_open_receipts(client, created_id, list(spec["receipts"]))
        out["select"] = selected
        if selected["http"] >= 400:
            out["status"] = "Fail"
            out["why"] = f"select receipts HTTP {selected['http']}"
            out["live"] = snapshot(client, created_id)
            return out
    live = snapshot(client, created_id)
    out["live"] = live
    receipt_ext = round(sum(line["ext"] or 0 for line in live["lines"]), 2)
    charge_sum = round(sum(c["amount"] or 0 for c in live["charges"]), 2)
    out["receipts_selected"] = [
        {
            "id": line["receipt"],
            "qty": line["qty"],
            "price": line["price"],
            "ext": line["ext"],
            "po": line["po"],
            "po_line": line["po_line"],
        }
        for line in live["lines"]
    ]
    out["receipt_extended"] = receipt_ext
    out["charges_sum"] = charge_sum
    if spec["kind"] == "freight":
        freight_ok = charge_sum == spec["total"] and any(
            "freight external" in str((c.get("kind") or {}).get("text") or "").lower()
            for c in live["charges"]
        )
        amount_ok = live["amount"] == spec["total"] and live["verification"] == spec["total"]
        if amount_ok and freight_ok and not live["lines"] and live["posted"] in (None, "", False) and attach == "attached":
            out["status"] = "Success"
            out["why"] = (
                f"Freight-only. Invoice_Amount {live['amount']:.2f} matches the PDF total. "
                "Freight External lookup id 1. No lines. Not posted."
            )
        else:
            out["status"] = "HOLD"
            out["why"] = (
                f"freight check failed amount={live['amount']} charges={charge_sum} "
                f"lines={len(live['lines'])} attach={attach}"
            )
        out["comments_1"] = ""
        return out

    skipped = spec.get("skip_receipts") or []
    if skipped:
        gap = round(spec["total"] - receipt_ext, 2)
        over = abs(gap) >= PPV_HOLD
        note = price_note(spec)
        comment = comments_1_shawn(client, created_id, price_html(spec))
        moved = apply_transfer_ap_batch_move(client, kimco_id=created_id)
        live = snapshot(client, created_id)
        out["live"] = live
        out["comments_1"] = note
        out["comments_1_put"] = comment
        out["transfer_ap"] = moved
        out["ppv"] = None
        out["price_gap"] = gap
        if over and live["batch_id"] == 375 and comment.get("status") == "persisted" and attach == "attached":
            out["status"] = "HOLD"
            out["exception"] = "price_variance"
            out["why"] = note
        else:
            out["status"] = "HOLD"
            out["exception"] = "price_variance"
            out["why"] = (
                f"{note} transfer={moved.get('status')} comment={comment.get('status')} "
                f"batch={live['batch_id']} attach={attach}"
            )
        return out

    ppv = round(spec["total"] - receipt_ext, 2)
    out["ppv"] = ppv if ppv else None
    if abs(ppv) >= PPV_HOLD:
        out["status"] = "HOLD"
        out["exception"] = "price_variance"
        out["why"] = f"unexpected |PPV| {ppv} on a bill with no skipped line"
        return out
    if ppv and abs(ppv) >= 0.01:
        ppv_put = client.try_post_ppv(created_id, ppv)
        out["ppv_post"] = ppv_put
        live = snapshot(client, created_id)
        out["live"] = live
    amount_ok = live["amount"] == spec["total"] and live["verification"] == spec["total"]
    lines_ok = len(live["lines"]) == len(spec["receipts"])
    if amount_ok and lines_ok and live["posted"] in (None, "", False) and attach == "attached" and live["type"] == 3:
        out["status"] = "Success"
        out["why"] = (
            f"Parts. Invoice_Amount {live['amount']:.2f} matches the PDF total. "
            f"Selected {len(live['lines'])} receipts. "
            + ("No PPV." if not ppv else f"PPV {ppv:.2f}.")
            + " Header PO blank (PDF SEE BELOW). Not posted."
        )
    else:
        out["status"] = "HOLD"
        out["why"] = (
            f"parts check failed amount={live['amount']} verification={live['verification']} "
            f"lines={len(live['lines'])} expected={len(spec['receipts'])} ppv={ppv} attach={attach}"
        )
    out["comments_1"] = ""
    return out


def charge_label(row: dict[str, Any]) -> str:
    live = row.get("live") or {}
    bits = []
    for charge in live.get("charges") or []:
        kind = charge.get("kind") or {}
        text = kind.get("text") if isinstance(kind, dict) else str(kind)
        bits.append(f"{text} {charge.get('amount')}")
    return "; ".join(bits)


def receipt_label(row: dict[str, Any]) -> str:
    bits = []
    for line in row.get("receipts_selected") or []:
        bits.append(f"{line.get('id')} qty {line.get('qty')} @ {line.get('price')}")
    return "; ".join(bits) if bits else "none"


def write_outputs(rows: list[dict[str, Any]]) -> None:
    payload = {
        "run": "3p-next5-round2-2026-09-24",
        "host": "https://live.kimcoerp.com",
        "vendor": "999-3P INDUSTRIES",
        "vendor_id": 1,
        "batch_id": 724,
        "batch_name": "API Agent - 9/22/26 3P",
        "posted": False,
        "emails_sent": False,
        "invent": False,
        "untouched": {"invoice": "142231", "kimco_id": 10186},
        "selection": (
            "Five oldest unentered 3P invoices by PDF date, then invoice number. "
            "Live list had no header for 142176, 142238, 142179, 142239, or 142188."
        ),
        "invoices": [],
    }
    for row in rows:
        live = row.get("live") or {}
        email = row.get("email") or {}
        payload["invoices"].append(
            {
                "invoice": row["invoice"],
                "pdf_date": row.get("pdf_date"),
                "pdf_total": row.get("pdf_total"),
                "po": row.get("po"),
                "po_on_pdf": row.get("po_on_pdf"),
                "pos": row.get("pos") or [],
                "kimco_id": row.get("kimco_id"),
                "status": row.get("status"),
                "exception": row.get("exception"),
                "receipts": row.get("receipts_selected") or [],
                "charges": [
                    {
                        "code": (c.get("kind") or {}).get("text") if isinstance(c.get("kind"), dict) else c.get("kind"),
                        "id": c.get("id"),
                        "amount": c.get("amount"),
                    }
                    for c in (live.get("charges") or [])
                ],
                "invoice_amount_after": live.get("amount"),
                "verification_amount": live.get("verification"),
                "comments_1": row.get("comments_1") or "",
                "email_folder": email.get("folder") or "",
                "email_moved_this_run": email.get("moved_this_run"),
                "why": row.get("why"),
                "batch_id": live.get("batch_id"),
                "batch": live.get("batch"),
                "posted": live.get("posted"),
                "invoice_type": live.get("type"),
                "attachments": live.get("attachments"),
            }
        )
    OUT_JSON.write_text(json.dumps(payload, indent=2, default=str) + "\n")

    prior = load_workbook(PRIOR_XLSX)
    updated = prior["3P updated"]
    headers = [cell.value for cell in next(updated.iter_rows(min_row=1, max_row=1))]
    index = {str(headers[i]): i for i in range(len(headers)) if headers[i]}
    by_invoice = {row["invoice"]: row for row in rows}
    seen = set()
    for excel_row in updated.iter_rows(min_row=2):
        number = str(excel_row[index["Invoice #"]].value or "")
        if number not in by_invoice:
            continue
        seen.add(number)
        row = by_invoice[number]
        live = row.get("live") or {}
        excel_row[index["date"]].value = row["pdf_date"]
        excel_row[index["PO"]].value = live.get("po")
        excel_row[index["Amount"]].value = live.get("amount")
        excel_row[index["Result"]].value = row.get("status")
        excel_row[index["Why"]].value = row.get("why")
        excel_row[index["Exception category"]].value = row.get("exception")
        excel_row[index["Exception owner"]].value = "Shawn McKibben" if row.get("exception") == "price_variance" else None
        excel_row[index["KIMCO id"]].value = row.get("kimco_id")
        batch_id = live.get("batch_id")
        batch_name = live.get("batch") or ""
        excel_row[index["Batch"]].value = f"{batch_name} ({batch_id})" if batch_id else None
        excel_row[index["Notes"]].value = (
            f"PDF total {row['pdf_total']:.2f}. PO on PDF {row.get('po_on_pdf')}. "
            f"Receipts: {receipt_label(row) or 'none'}. Charges: {charge_label(row) or 'none'}. "
            f"Email folder: {(row.get('email') or {}).get('folder')}. Not posted."
        )
        excel_row[index["Recheck verdict"]].value = row.get("status")
        excel_row[index["Action taken"]].value = (
            f"Created KIMCO {row.get('kimco_id')} on batch 724, then "
            + ("moved to Transfer AP 375. " if batch_id == 375 else "left on batch 724. ")
            + "Attached the invoice page only. Did not post."
        )
    if seen != set(by_invoice):
        raise SystemExit(f"3P updated sheet missing invoices: {set(by_invoice) - seen}")

    wb = Workbook()
    first = wb.active
    first.title = "This run (5)"
    first_headers = [
        "Invoice #",
        "Verdict",
        "PDF date",
        "PDF total",
        "PO",
        "KIMCO id",
        "Receipts",
        "Charges",
        "Invoice_Amount after",
        "Comments_1",
        "Email folder",
    ]
    first.append(first_headers)
    for row in rows:
        live = row.get("live") or {}
        first.append(
            [
                row["invoice"],
                row.get("status"),
                row.get("pdf_date"),
                row.get("pdf_total"),
                ", ".join(row.get("pos") or []) or row.get("po_on_pdf"),
                row.get("kimco_id"),
                receipt_label(row),
                charge_label(row),
                live.get("amount"),
                row.get("comments_1") or "",
                (row.get("email") or {}).get("folder") or "",
            ]
        )
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in first[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for excel_row in first.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [14, 12, 14, 12, 18, 12, 42, 28, 18, 55, 28]
    for i, width in enumerate(widths, 1):
        first.column_dimensions[get_column_letter(i)].width = width
    first.auto_filter.ref = f"A1:K{first.max_row}"
    first.freeze_panes = "A2"
    first.row_dimensions[1].height = 22
    for r in range(2, first.max_row + 1):
        first.row_dimensions[r].height = 48

    # Copy refreshed 3P updated, plus a lines sheet and summary.
    dest_updated = wb.create_sheet("3P updated")
    for excel_row in updated.iter_rows(values_only=True):
        dest_updated.append(list(excel_row))
    for cell in dest_updated[1]:
        cell.font = Font(bold=True)

    lines_ws = wb.create_sheet("Invoice PDF lines")
    lines_ws.append(
        ["Invoice #", "PDF page", "Date", "Due", "PO", "Part", "Qty", "Price", "Line amount", "PDF total", "Entered this run", "KIMCO id"]
    )
    for row in rows:
        spec = next(b for b in BILLS if b["invoice"] == row["invoice"])
        if row["kind"] == "freight":
            lines_ws.append(
                [
                    row["invoice"],
                    spec["pages"][0] + 1,
                    row["pdf_date"],
                    row["pdf_due"],
                    "FREIGHT",
                    row.get("pdf_line"),
                    1,
                    row["pdf_total"],
                    row["pdf_total"],
                    row["pdf_total"],
                    "yes",
                    row.get("kimco_id"),
                ]
            )
            continue
        for line in spec["lines"]:
            lines_ws.append(
                [
                    row["invoice"],
                    spec["pages"][0] + 1,
                    row["pdf_date"],
                    row["pdf_due"],
                    line["po"],
                    line["part"],
                    line["qty"],
                    line["unit"],
                    line["amount"],
                    row["pdf_total"],
                    "yes",
                    row.get("kimco_id"),
                ]
            )

    summary = wb.create_sheet("Summary")
    success = [r for r in rows if r.get("status") == "Success"]
    holds = [r for r in rows if r.get("status") != "Success"]
    summary.append(["Run", "3P next 5 round 2, 2026-09-24"])
    summary.append(["Vendor", "999-3P INDUSTRIES (id 1)"])
    summary.append(["Batch started", "724 API Agent - 9/22/26 3P"])
    summary.append(["Entered this run", len(rows)])
    summary.append(["Success", len(success)])
    summary.append(["HOLD", len(holds)])
    summary.append(["Bills posted", "no"])
    summary.append(["Email sent", "no"])
    summary.append(["142231 / KIMCO 10186", "not touched"])
    for row in rows:
        live = row.get("live") or {}
        summary.append(
            [
                row["invoice"],
                f"{row.get('status')} KIMCO {row.get('kimco_id')} PDF {row.get('pdf_total')} "
                f"Invoice_Amount {live.get('amount')} batch {live.get('batch_id')}",
            ]
        )
    for row in holds:
        summary.append([f"Blocker {row['invoice']}", row.get("why")])

    wb.save(OUT_XLSX)
    print("WROTE", OUT_JSON)
    print("WROTE", OUT_XLSX)


def main() -> int:
    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error or "live credentials missing")
    client = KimcoClient.authenticate(
        creds.instance_url, creds.key or "", creds.password or "", target="live"
    )
    graph = _optional_graph_client()
    if graph is None:
        raise SystemExit("Graph client required")
    batch = client.get_item("ap_batches", BATCH_ID)
    bvals = batch.get("values") or {}
    print("BATCH", BATCH_ID, bvals.get("AP_Invoice_Batch_ID"), flush=True)
    already = existing_3p(client)
    if already.get("142231") != 10186:
        raise SystemExit(f"142231 id changed: {already.get('142231')}")
    results = []
    for spec in BILLS:
        number = spec["invoice"]
        if number in already:
            print("STOP already in KIMCO", number, already[number], flush=True)
            results.append(
                {
                    "invoice": number,
                    "status": "skipped-already-in-kimco",
                    "kimco_id": already[number],
                    "pdf_date": spec["date"].isoformat(),
                    "pdf_due": spec["due"].isoformat(),
                    "pdf_total": spec["total"],
                    "kind": spec["kind"],
                    "po": None,
                    "po_on_pdf": spec.get("po_on_pdf"),
                    "pos": spec.get("pos") or [],
                    "live": {},
                    "email": {},
                    "why": "already had a live header; not entered again",
                }
            )
            break
        print("ENTER", number, spec["kind"], spec["total"], flush=True)
        row = enter_one(client, graph, spec)
        results.append(row)
        live = row.get("live") or {}
        print(
            number,
            row.get("status"),
            row.get("kimco_id"),
            "amount",
            live.get("amount"),
            "batch",
            live.get("batch_id"),
            "folder",
            (row.get("email") or {}).get("folder"),
            flush=True,
        )
        if row.get("status") not in {"Success", "HOLD"}:
            print("STOP", number, row.get("why"), flush=True)
            break
        if spec["kind"] == "parts" and not spec.get("skip_receipts") and row.get("status") != "Success":
            print("STOP parts success expected", number, row.get("why"), flush=True)
            break
        if spec["kind"] == "freight" and row.get("status") != "Success":
            print("STOP freight success expected", number, row.get("why"), flush=True)
            break
    if len(results) == len(BILLS):
        write_outputs(results)
    else:
        Path("/tmp/3p-round2/partial-results.json").write_text(json.dumps(results, indent=2, default=str))
        print("PARTIAL", len(results), flush=True)
    ok = len(results) == len(BILLS) and all(r.get("status") in {"Success", "HOLD"} for r in results)
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
