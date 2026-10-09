"""Rebuild runs/email-undo-2026-10-06/plan.json from repo files only (offline, no Graph, no KIMCO).
Sources: runs/AP-sept-catchup-2026-10-06.xlsx, runs/sept-catchup-progress.json, qc/sept-catchup-mail-restore.json.
Reproduces the 476 rows of the 10/6 audit: 233 restored by the 'AI Needs Review' restore + 243 moved/re-tagged by the catch-up only."""
import base64, collections, csv, json, re, sys
from datetime import datetime, timedelta, timezone
from pathlib import Path
import openpyxl
R = Path(__file__).resolve().parents[1]
OUT = R / "runs" / "email-undo-2026-10-06"
def dec(i):
    s = i.replace('-', '+').replace('_', '/'); s += '=' * (-len(s) % 4); return base64.b64decode(s)
def enc(b): return base64.b64encode(b).decode().replace('+', '-').replace('/', '_')
def folder_id(mid):
    b = dec(mid); hdr = b[:41]; ln = int.from_bytes(b[41:43], 'little'); e = b[43:43 + ln]
    assert ln == 0x46 and e[20:22] == b'\x07\x00'
    fe = b'\x00\x00\x00\x00' + e[4:20] + b'\x01\x00' + e[22:44] + b'\x00\x00'
    return enc(hdr + len(fe).to_bytes(2, 'little') + fe)
FID = {'00000000010c': 'Inbox', '0000831fecb5': '9 - FORT WORTH ARCHIVE', '000000000109': 'Sent Items',
       '000376dc5ed3': 'UNNAMED (FID 0x376dc5ed3)', '000000000113': 'UNNAMED (FID 0x0113)'}
def pre_folder(mid): h = dec(mid)[43 + 38:43 + 44].hex(); return FID.get(h, 'UNNAMED (FID 0x' + h.lstrip('0') + ')')
CT = timezone(timedelta(hours=-5))
def ct(z): return datetime.fromisoformat(z.replace('Z', '+00:00')).astimezone(CT).strftime('%Y-%m-%d %H:%M')
LEAVE_RE = re.compile(r"\bstatement\b|credit\s*(memo|note)|\bpast[\s-]*due\b|proof of delivery|\bpods?\b|payment submitted|payment was successfully|automatic payment|approved automatic payment|payment declined|scheduled payment|payment due|payment inquiry|notice of intent to cancel|(?:check|ck)\s*#|cancell?ed check|lost check", re.I)
SKIP = {"no-attachment", "no-pdf", "unreadable-or-not-a-bill"}
def decision(s, pr):
    why = str(pr.get('Why') or '')
    if s.get('Result') == 'Hold' and not s.get('Bill id'):
        return 'category-only' if str(pr.get('Folder move') or '').startswith('left') else 'move'
    if s.get('Result') == 'Skipped' and why.split('.', 1)[0].strip() in SKIP:
        return 'leave-genuine' if LEAVE_RE.search(f"{pr.get('subject') or ''} {s.get('Vendor') or ''}") else 'move'
    return 'keep'
REMOVABLE = ['AI Needs Review', 'Entered in AI', 'AI HOLD', 'Entered with issues', 'AI Skipped 2', 'AI Skipped']
def main():
    ws = openpyxl.load_workbook(R / 'runs/AP-sept-catchup-2026-10-06.xlsx', data_only=True)['September catch-up']
    h = [c.value for c in ws[1]]; sheet = [dict(zip(h, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    prog = json.load(open(R / 'runs/sept-catchup-progress.json'))['rows']
    assert len(sheet) == len(prog) == 620
    man = json.load(open(R / 'qc/sept-catchup-mail-restore.json'))
    lr = man['leftover_restored']; assert folder_id(lr[0]['id']) == lr[0]['parent']
    hints = collections.defaultdict(list)
    for x in lr: hints[x['receivedDateTime']].append(x['id'])
    g = collections.defaultdict(list)
    for s, pr in zip(sheet, prog): g[pr['graph_message_id']].append((s, pr, decision(s, pr)))
    rows = []
    for mid, it in g.items():
        prs = [pr for _, pr, _ in it]; ds = {d for *_, d in it}
        restore = 'move' if 'move' in ds else ('category-only' if 'category-only' in ds else None)
        fm = str(prs[0].get('Folder move') or '')
        if not restore and fm.startswith('left'): continue   # never touched by finalize_mail
        recv = {pr['receivedDateTime'] for pr in prs}; assert len(recv) == 1; ru = recv.pop()
        subj = sorted({pr['subject'] for pr in prs if pr.get('subject')})
        inv = sorted({str(pr['Invoice #']) for pr in prs if pr.get('Invoice #') and str(pr['Invoice #']).lower() not in ('unknown', 'none')})
        cat = ', '.join(sorted({str(pr.get('Flag status') or '') for pr in prs}))
        pre = pre_folder(mid)
        if restore:
            exp_folder, exp_cats, tab = 'Inbox', 'AI Needs Review', 'Undo plan'
        else:
            exp_folder = 'AutoPay Archive' if fm.startswith('AutoPay') else ('9 - FORT WORTH ARCHIVE' if fm.startswith('9 -') else fm)
            exp_cats, tab = cat, 'Not restored'
        rows.append(dict(source_tab=tab, received_ct=ct(ru), received_utc=ru, sender_recorded=str(prs[0].get('Vendor') or ''),
            subject_recorded=(subj[0] if subj else ''), subjects_logged=' | '.join(subj), invoice_recorded=' | '.join(inv),
            pre_run_folder=pre, pre_run_folder_id=folder_id(mid), expected_current_folder=exp_folder, expected_current_categories=exp_cats,
            catchup_category=cat, planned_destination=pre, planned_move='no' if pre == exp_folder else 'yes',
            categories_to_remove=';'.join(REMOVABLE), restore_action=restore or 'not restored',
            post_restore_id_hint=';'.join(hints.get(ru, [])) if restore else '', pre_run_graph_id=mid))
    rows.sort(key=lambda x: (x['source_tab'] != 'Undo plan', x['received_utc']))
    n = collections.Counter()
    for x in rows:
        n[x['source_tab']] += 1; x['row_id'] = ('U' if x['source_tab'] == 'Undo plan' else 'N') + f"{n[x['source_tab']]:03d}"
    assert n['Undo plan'] == 233 and n['Not restored'] == 243, n
    return rows
if __name__ == '__main__':
    rows = main()
    print(collections.Counter((r['planned_destination'], r['planned_move']) for r in rows))
    if '--write' in sys.argv:
        OUT.mkdir(parents=True, exist_ok=True); json.dump(rows, open(OUT / 'plan.rebuilt.json', 'w'), indent=1)
