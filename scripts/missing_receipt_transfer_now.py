"""Kyle 2026-09-22: move current missing_receipt HOLDs to Transfer AP NOW.

O'Neal 10160/10163/10175/10178 (722→375).
Gas 10173/10181 (720→375).

GET first. Only move if still missing_receipt HOLD. PUT batch via NOTE-53
helper (lookup by name; 375 hint only). Keep/ensure Comments_1. No Select.
No Mail.Send. invent=false. Do not watch Transfer AP after the move.
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook

from ap_clerk.auth import load_credentials
from ap_clerk.kimco import KimcoClient, KimcoError
from ap_clerk.rules import SHAWN_MCKIBBEN, TRANSFER_AP_BATCH_NAME, lookup_id, lookup_text, money
from ap_clerk.transfer_ap import apply_missing_receipt_transfer_ap, should_transfer_ap_missing_receipt
from gas_0040438052 import SHAWN_MENTION_ID
from gas_59081_retry import comments_1_shawn
from gas_supply_0917 import exact_invoice_number

LOGGER = logging.getLogger("ap_clerk.missing_receipt_transfer_now")

ONEAL_SHEET = ROOT / "runs" / "AP-run-2026-09-21-oneal.xlsx"
ONEAL_JSON = ROOT / "runs" / "AP-run-2026-09-21-oneal.json"
GAS_SHEET = ROOT / "runs" / "AP-run-2026-09-17-gas-supply.xlsx"
GAS_JSON = ROOT / "runs" / "AP-run-2026-09-17-gas-supply.json"
PROOF = ROOT / "runs" / "missing-receipt-transfer-now-2026-09-22.json"

TARGETS: tuple[dict[str, Any], ...] = (
    {
        "kimco_id": 10160,
        "invoice": "15469453",
        "po": "59156",
        "vendor": "O'Neal",
        "from_batch": 722,
        "owner": "Anthony",
        "mention": "plain",
        "amount": 23151.42,
    },
    {
        "kimco_id": 10163,
        "invoice": "15464707",
        "po": "58912",
        "vendor": "O'Neal",
        "from_batch": 722,
        "owner": "Anthony",
        "mention": "plain",
        "amount": 649.53,
    },
    {
        "kimco_id": 10175,
        "invoice": "15460995",
        "po": "59085",
        "vendor": "O'Neal",
        "from_batch": 722,
        "owner": "Anthony",
        "mention": "plain",
        "amount": 6954.21,
    },
    {
        "kimco_id": 10178,
        "invoice": "15476365",
        "po": "59070",
        "vendor": "O'Neal",
        "from_batch": 722,
        "owner": "Anthony",
        "mention": "plain",
        "amount": 1115.68,
    },
    {
        "kimco_id": 10173,
        "invoice": "0040438052",
        "po": "59081",
        "vendor": "Gas",
        "from_batch": 720,
        "owner": "Shawn McKibben",
        "mention": "shawn",
        "amount": 5178.55,
    },
    {
        "kimco_id": 10181,
        "invoice": "0040437952",
        "po": "59006",
        "vendor": "Gas",
        "from_batch": 720,
        "owner": "Shawn McKibben",
        "mention": "shawn",
        "amount": 241.46,
    },
)


def snapshot(client: KimcoClient, kid: int) -> dict[str, Any]:
    rec = client.get_item("ap_invoices", int(kid))
    vals = rec.get("values") or {}
    comments = (rec.get("lists") or {}).get("Comments_1") or []
    items: list[dict[str, Any]] = []
    for row in comments if isinstance(comments, list) else []:
        html = str((row.get("values") or {}).get("HtmlValue") or "")
        items.append({"id": row.get("id"), "html": html[:500]})
    return {
        "kimco_id": int(kid),
        "invoice": exact_invoice_number(vals.get("Invoice_Number"))
        or str(vals.get("Invoice_Number") or ""),
        "amount": money(vals.get("Invoice_Amount")),
        "verification": money(vals.get("Invoice_Verification_Amount")),
        "type": vals.get("Invoice_Type"),
        "po": lookup_text(vals.get("PO") or vals.get("Purchase_Order")),
        "batch_id": lookup_id(vals.get("AP_Invoice_Batch")),
        "batch": lookup_text(vals.get("AP_Invoice_Batch")),
        "receipt_n": len((rec.get("lists") or {}).get("APInvoiceLine") or []),
        "comments": items,
    }


def still_missing_receipt_hold(snap: dict[str, Any], spec: dict[str, Any]) -> tuple[bool, str]:
    """GET-only gate. Never invent Success. Never treat selected receipts as missing."""
    want = exact_invoice_number(spec["invoice"]) or spec["invoice"]
    got = exact_invoice_number(snap.get("invoice")) or snap.get("invoice")
    if want != got:
        return False, "invoice-mismatch"
    why = " ".join(str(c.get("html") or "") for c in (snap.get("comments") or []))
    if "price-does-not-match" in why.lower() and "HOLD (receipt)" not in why:
        return False, "price_variance-not-missing-receipt"
    receipt_n = int(snap.get("receipt_n") or 0)
    amount = snap.get("amount")
    verification = snap.get("verification")
    if receipt_n > 0 and amount not in (None, 0) and verification not in (None, 0):
        if abs(float(amount) - float(verification)) < 0.02:
            return False, "looks-finished-has-receipts"
        return False, "receipts-already-selected"
    if receipt_n > 0:
        return False, "receipts-already-selected"
    if not should_transfer_ap_missing_receipt(
        result="HOLD",
        category="missing_receipt",
        why="HOLD (receipt): no open receipt leftover",
    ):
        return False, "rule-false"
    return True, "missing_receipt-no-receipts"


def comments_present(snap: dict[str, Any], spec: dict[str, Any]) -> list[dict[str, Any]]:
    inv = exact_invoice_number(spec["invoice"]) or spec["invoice"]
    hits = []
    for row in snap.get("comments") or []:
        html = str(row.get("html") or "")
        if inv not in html or "HOLD (receipt)" not in html:
            continue
        if spec.get("mention") == "shawn" and f'data-mention-id="{SHAWN_MENTION_ID}"' not in html:
            continue
        hits.append(row)
    return hits


def comments_1_plain(client: KimcoClient, kimco_id: int, html: str) -> dict[str, Any]:
    payload = {
        "state": "Modified",
        "id": int(kimco_id),
        "lists": {
            "Comments_1": [
                {
                    "state": "Added",
                    "values": {
                        "HtmlValue": html,
                        "Entity": {"id": 203},
                        "ObjectId": int(kimco_id),
                        "FormId": 218,
                    },
                }
            ]
        },
    }
    try:
        _body, status, error = client.update("ap_invoices", int(kimco_id), payload)
    except KimcoError as exc:
        return {"status": "blocked", "error": str(exc)[:200], "invent": False}
    return {"status": "persisted" if status < 400 else f"put-{status}", "put": status, "error": error, "invent": False}


def ensure_comments(client: KimcoClient, spec: dict[str, Any], snap: dict[str, Any]) -> dict[str, Any]:
    existing = comments_present(snap, spec)
    if existing:
        return {"status": "already-on-tab", "comments_1": existing, "invent": False}
    amt = money(spec.get("amount"))
    amt_txt = f"{amt:.2f}" if amt is not None else "unknown"
    inv = spec["invoice"]
    po = spec["po"]
    if spec.get("mention") == "shawn":
        body = (
            f"{SHAWN_MCKIBBEN} HOLD (receipt) on Gas and Supply {inv} "
            f"PO {po} PDF ${amt_txt}. Open leftover merch is not received "
            "or does not match PDF lines. Kyle 2026-09-22 NOTE-53: moved to "
            f"{TRANSFER_AP_BATCH_NAME}. Receiving / purchasing: receive the PO "
            "so AP can Select Receipts. No email."
        )
        html = (
            f'<p><span data-mention-id="{SHAWN_MENTION_ID}" '
            f'data-mention-name="Shawn McKibben" '
            f'data-mention-email="Shawn.McKibben@kannonmfg.com" '
            f'class="prosemirror-mention-node">{SHAWN_MCKIBBEN}</span> '
            f"{body[len(SHAWN_MCKIBBEN):].strip()}</p>"
        )
        put = comments_1_shawn(client, int(spec["kimco_id"]), html)
    else:
        html = (
            f"<p>Anthony HOLD (receipt) on O'Neal Steel {inv} PO {po} "
            f"PDF ${amt_txt}. Open leftover merch is not received or does not "
            f"match PDF lines. Kyle 2026-09-22 NOTE-53: moved to "
            f"{TRANSFER_AP_BATCH_NAME}. Mention-id for Anthony is not proven "
            "(invent=false). No email.</p>"
        )
        put = comments_1_plain(client, int(spec["kimco_id"]), html)
    after = snapshot(client, int(spec["kimco_id"]))
    return {
        "status": "added" if comments_present(after, spec) else put.get("status"),
        "put": put,
        "comments_1": comments_present(after, spec),
        "invent": False,
    }


def move_header(client: KimcoClient, spec: dict[str, Any], *, dry_run: bool) -> dict[str, Any]:
    before = snapshot(client, int(spec["kimco_id"]))
    ok, reason = still_missing_receipt_hold(before, spec)
    out: dict[str, Any] = {
        "kimco_id": spec["kimco_id"],
        "invoice": spec["invoice"],
        "vendor": spec["vendor"],
        "before_batch": before.get("batch_id"),
        "before_batch_name": before.get("batch"),
        "before": before,
        "eligible": ok,
        "reason": reason,
        "mail_send": False,
        "select": False,
        "invent": False,
    }
    if not ok:
        out["status"] = f"skipped-{reason}"
        out["after_batch"] = before.get("batch_id")
        return out
    if dry_run:
        out["status"] = "dry-run"
        out["after_batch"] = before.get("batch_id")
        return out
    if before.get("batch_id") == 375 or str(before.get("batch") or "").casefold() == TRANSFER_AP_BATCH_NAME.casefold():
        comments = ensure_comments(client, spec, before)
        after = snapshot(client, int(spec["kimco_id"]))
        out["status"] = "already-on-transfer-ap"
        out["comments"] = comments
        out["after"] = after
        out["after_batch"] = after.get("batch_id")
        return out
    moved = apply_missing_receipt_transfer_ap(client, kimco_id=int(spec["kimco_id"]))
    after_move = snapshot(client, int(spec["kimco_id"]))
    comments = ensure_comments(client, spec, after_move)
    after = snapshot(client, int(spec["kimco_id"]))
    out["move"] = moved
    out["comments"] = comments
    out["after"] = after
    out["after_batch"] = after.get("batch_id")
    out["status"] = moved.get("status")
    return out


def _patch_sidecar_rows(
    path: Path,
    results: list[dict[str, Any]],
    *,
    proof_key: str,
) -> list[dict[str, Any]]:
    sidecar: dict[str, Any] = {}
    if path.is_file():
        sidecar = json.loads(path.read_text())
    prior = [r for r in (sidecar.get("rows") or []) if isinstance(r, dict)]
    by_inv = {exact_invoice_number(r.get("invoice")): r for r in results}
    for row in prior:
        inv = exact_invoice_number(row.get("Invoice #"))
        hit = by_inv.get(inv)
        if not hit:
            continue
        after_bid = hit.get("after_batch")
        if after_bid in (375, "375") or hit.get("status") in {"moved", "already-on-transfer-ap"}:
            row["Batch"] = f"{TRANSFER_AP_BATCH_NAME} (375)" if after_bid in (None, 375, "375") else f"{TRANSFER_AP_BATCH_NAME} ({after_bid})"
            if after_bid not in (None, ""):
                row["Batch"] = f"{TRANSFER_AP_BATCH_NAME} ({after_bid})"
        row["Exception category"] = "missing_receipt"
        row["Exception owner"] = "Anthony" if hit.get("vendor") == "O'Neal" else "Shawn McKibben"
        extra = (
            f"Kyle 2026-09-22 NOTE-53 missing_receipt → Transfer AP "
            f"{hit.get('before_batch')}→{hit.get('after_batch')} status={hit.get('status')}."
        )
        why = str(row.get("Why") or "")
        if extra not in why:
            row["Why"] = f"{why} {extra}".strip()
        notes = str(row.get("Notes") or "")
        if extra not in notes:
            row["Notes"] = f"{notes} {extra}".strip()
    sidecar["rows"] = prior
    sidecar[proof_key] = results
    path.write_text(json.dumps(sidecar, indent=2, default=str) + "\n")
    return prior


def _patch_xlsx(path: Path, results: list[dict[str, Any]]) -> None:
    if not path.is_file():
        return
    wb = load_workbook(path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    by_inv = {exact_invoice_number(r.get("invoice")): r for r in results}

    def col(name: str) -> int | None:
        return headers.index(name) if name in headers else None

    inv_i = col("Invoice #")
    if inv_i is None:
        return
    batch_i = col("Batch")
    cat_i = col("Exception category")
    owner_i = col("Exception owner")
    why_i = col("Why")
    notes_i = col("Notes")
    for excel_row in ws.iter_rows(min_row=2):
        inv = exact_invoice_number(excel_row[inv_i].value)
        hit = by_inv.get(inv)
        if not hit:
            continue
        after_bid = hit.get("after_batch") or 375
        extra = (
            f"Kyle 2026-09-22 NOTE-53 missing_receipt → Transfer AP "
            f"{hit.get('before_batch')}→{hit.get('after_batch')} status={hit.get('status')}."
        )
        if batch_i is not None and hit.get("status") in {"moved", "already-on-transfer-ap"}:
            excel_row[batch_i].value = f"{TRANSFER_AP_BATCH_NAME} ({after_bid})"
        if cat_i is not None:
            excel_row[cat_i].value = "missing_receipt"
        if owner_i is not None:
            excel_row[owner_i].value = "Anthony" if hit.get("vendor") == "O'Neal" else "Shawn McKibben"
        if why_i is not None:
            current = str(excel_row[why_i].value or "")
            if extra not in current:
                excel_row[why_i].value = f"{current} {extra}".strip()
        if notes_i is not None:
            current = str(excel_row[notes_i].value or "")
            if extra not in current:
                excel_row[notes_i].value = f"{current} {extra}".strip()
    wb.save(path)


def patch_sheets(results: list[dict[str, Any]]) -> None:
    oneal = [r for r in results if r.get("vendor") == "O'Neal"]
    gas = [r for r in results if r.get("vendor") == "Gas"]
    _patch_sidecar_rows(ONEAL_JSON, oneal, proof_key="missing_receipt_transfer_0922")
    _patch_xlsx(ONEAL_SHEET, oneal)
    _patch_sidecar_rows(GAS_JSON, gas, proof_key="missing_receipt_transfer_0922")
    _patch_xlsx(GAS_SHEET, gas)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")
    if not args.live:
        raise SystemExit("Refusing: pass --live.")
    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error or "live credentials not ready")
    client = KimcoClient.authenticate(
        creds.instance_url, creds.key or "", creds.password or "", target="live"
    )
    proof: dict[str, Any] = {
        "proof": "missing-receipt-transfer-now-2026-09-22",
        "note": 53,
        "invent": False,
        "mail_send": False,
        "select": False,
        "dry_run": args.dry_run,
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "targets": [spec["kimco_id"] for spec in TARGETS],
    }
    results = [move_header(client, spec, dry_run=args.dry_run) for spec in TARGETS]
    proof["results"] = results
    proof["moved_count"] = sum(1 for r in results if r.get("status") == "moved")
    proof["already_count"] = sum(1 for r in results if r.get("status") == "already-on-transfer-ap")
    proof["skipped_count"] = sum(1 for r in results if str(r.get("status") or "").startswith("skipped"))
    if not args.dry_run:
        patch_sheets(results)
        proof["oneal_sheet"] = str(ONEAL_SHEET)
        proof["gas_sheet"] = str(GAS_SHEET)
    PROOF.write_text(json.dumps(proof, indent=2, default=str) + "\n")
    print(json.dumps(proof, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
