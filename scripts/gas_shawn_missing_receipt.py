"""Kyle 2026-09-22 override: ping Shawn on Gas missing_receipt HOLDs.

10173 / 0040438052 PO 59081 ~$5178.55 batch 720
10181 / 0040437952 PO 59006 ~$241.46 batch 720

Comments_1 @Shawn McKibben mention-id 104 (Entity 203, FormId 218).
Ask Shawn to receive the PO so AP can Select Receipts. Stay on 720.
Do NOT Transfer AP. Do NOT Select leftovers. Do NOT Mail.Send.
Do NOT invent Success. Type 3 PO bills (Lines-K N/A).
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook

from ap_clerk.auth import load_credentials
from ap_clerk.kimco import KimcoClient, KimcoError
from ap_clerk.quality_v12 import COL_EXCEPTION_OWNER
from ap_clerk.report import write_report
from ap_clerk.rules import SHAWN_MCKIBBEN, lookup_id, lookup_text, money
from gas_0040438052 import SHAWN_MENTION_ID, _sheet_row_for
from gas_59081_retry import comments_1_shawn
from gas_supply_0917 import (
    KNOWN_BATCH_ID,
    PREFERRED_BATCH_NAME,
    exact_invoice_number,
    prior_rows_from_sidecar,
    verify_batch,
)

LOGGER = logging.getLogger("ap_clerk.gas_shawn_missing_receipt")

TARGETS = (
    {
        "kimco_id": 10173,
        "invoice": "0040438052",
        "po": "59081",
        "amount": 5178.55,
    },
    {
        "kimco_id": 10181,
        "invoice": "0040437952",
        "po": "59006",
        "amount": 241.46,
    },
)
SHEET = ROOT / "runs" / "AP-run-2026-09-17-gas-supply.xlsx"
SHEET_JSON = ROOT / "runs" / "AP-run-2026-09-17-gas-supply.json"
PROOF = ROOT / "runs" / "gas-shawn-missing-receipt-2026-09-22.json"
SHAWN_OWNER = "Shawn McKibben"


def receipt_hold_comment(*, invoice: str, po: str, amount: Any) -> str:
    amt = money(amount)
    amt_txt = f"{amt:.2f}" if amt is not None else "unknown"
    inv = exact_invoice_number(invoice) or str(invoice or "")
    return (
        f"{SHAWN_MCKIBBEN} HOLD (receipt) on Gas and Supply {inv} "
        f"PO {po} PDF ${amt_txt}. Open leftover merch is not received "
        "or does not match PDF lines. Kyle 2026-09-22 override: ping Shawn "
        "(Gas receiving owner was blank). Header stays on "
        f"{PREFERRED_BATCH_NAME} ({KNOWN_BATCH_ID}) — do not Transfer AP. "
        "Receiving / purchasing: receive the PO so AP can Select Receipts. "
        "No email."
    )


def mention_html(body: str) -> str:
    text = (body or "").strip()
    if text.startswith(SHAWN_MCKIBBEN):
        text = text[len(SHAWN_MCKIBBEN) :].strip()
    return (
        f'<p><span data-mention-id="{SHAWN_MENTION_ID}" '
        f'data-mention-name="{SHAWN_OWNER}" '
        f'data-mention-email="Shawn.McKibben@kannonmfg.com" '
        f'class="prosemirror-mention-node">{SHAWN_MCKIBBEN}</span> {text}</p>'
    )


def snapshot(client: KimcoClient, kid: int) -> dict[str, Any]:
    rec = client.get_item("ap_invoices", int(kid))
    vals = rec.get("values") or {}
    comments = (rec.get("lists") or {}).get("Comments_1") or []
    items: list[dict[str, Any]] = []
    for row in comments if isinstance(comments, list) else []:
        html = str((row.get("values") or {}).get("HtmlValue") or "")
        items.append({"id": row.get("id"), "html": html[:400]})
    return {
        "kimco_id": int(kid),
        "invoice": exact_invoice_number(vals.get("Invoice_Number")),
        "amount": money(vals.get("Invoice_Amount")),
        "verification": money(vals.get("Invoice_Verification_Amount")),
        "type": vals.get("Invoice_Type"),
        "po": lookup_text(vals.get("PO") or vals.get("Purchase_Order")),
        "batch_id": lookup_id(vals.get("AP_Invoice_Batch")),
        "batch": lookup_text(vals.get("AP_Invoice_Batch")),
        "receipt_n": len((rec.get("lists") or {}).get("APInvoiceLine") or []),
        "comments": items,
    }


def receipt_comment_items(comments: list[dict[str, Any]], invoice: str) -> list[dict[str, Any]]:
    inv = exact_invoice_number(invoice) or invoice
    hits = []
    for row in comments:
        html = str(row.get("html") or "")
        if (
            "HOLD (receipt)" in html
            and inv in html
            and f'data-mention-id="{SHAWN_MENTION_ID}"' in html
        ):
            hits.append(row)
    return hits


def move_back_to_720(client: KimcoClient, kid: int) -> dict[str, Any]:
    """Restore missing_receipt to the API Agent batch. Never move onto 375."""
    payload = {
        "state": "Modified",
        "id": int(kid),
        "values": {"AP_Invoice_Batch": {"id": int(KNOWN_BATCH_ID)}},
    }
    try:
        _body, status, error = client.update("ap_invoices", int(kid), payload)
    except KimcoError as exc:
        return {"status": "blocked", "error": str(exc)[:200]}
    return {"status": "moved" if status < 400 else f"put-{status}", "put": status, "error": error}


def ping_header(client: KimcoClient, spec: dict[str, Any]) -> dict[str, Any]:
    kid = int(spec["kimco_id"])
    before = snapshot(client, kid)
    report: dict[str, Any] = {
        "kimco_id": kid,
        "invoice": spec["invoice"],
        "po": spec["po"],
        "before": before,
        "mail_send": False,
        "transfer_ap": False,
        "select": False,
    }
    if before.get("invoice") != spec["invoice"]:
        report["status"] = "invoice-mismatch"
        report["after"] = before
        return report
    if before.get("batch_id") == 375:
        # Prior false over-PPV left 10181 on Transfer AP. Kyle: missing_receipt
        # stays on 720. Restore only 375 → 720; never the reverse.
        report["restore_720"] = move_back_to_720(client, kid)
        before = snapshot(client, kid)
        report["after_restore"] = {
            "batch_id": before.get("batch_id"),
            "batch": before.get("batch"),
        }
    if before.get("batch_id") != KNOWN_BATCH_ID:
        report["status"] = "refused-wrong-batch"
        report["after"] = before
        return report
    existing = receipt_comment_items(before.get("comments") or [], spec["invoice"])
    if existing:
        report["status"] = "already-on-tab"
        report["comments_1"] = existing
        report["after"] = before
        return report
    body = receipt_hold_comment(
        invoice=spec["invoice"], po=spec["po"], amount=spec["amount"]
    )
    html = mention_html(body)
    report["html"] = html
    put = comments_1_shawn(client, kid, html)
    after = snapshot(client, kid)
    hits = receipt_comment_items(after.get("comments") or [], spec["invoice"])
    report["put"] = put
    report["after"] = after
    report["comments_1"] = hits
    report["batch_still_720"] = after.get("batch_id") == KNOWN_BATCH_ID
    if after.get("batch_id") != KNOWN_BATCH_ID:
        report["status"] = "batch-moved-unexpected"
    elif hits:
        report["status"] = "added"
    else:
        report["status"] = "put-tab-missing"
    return report


def patch_sheet(results: list[dict[str, Any]]) -> list[dict[str, Any]]:
    prior = prior_rows_from_sidecar(SHEET_JSON)
    by_inv = {exact_invoice_number(r.get("Invoice #")): r for r in results}
    updated: list[dict[str, Any]] = []
    for row in prior:
        inv = exact_invoice_number(row.get("Invoice #"))
        hit = by_inv.get(inv)
        if not hit:
            updated.append(row)
            continue
        out = dict(row)
        out[COL_EXCEPTION_OWNER] = SHAWN_OWNER
        out["Exception category"] = "missing_receipt"
        comment_ids = [
            c.get("id") for c in (hit.get("comments_1") or []) if c.get("id") not in (None, "")
        ]
        out["Why"] = (
            f"category=missing_receipt; owner={SHAWN_OWNER}. "
            f"Kyle 2026-09-22 override: ping Shawn on Gas missing_receipt "
            f"(owner was blank). Comments_1 status={hit.get('status')} "
            f"item_ids={comment_ids} mention-id={SHAWN_MENTION_ID}. "
            f"Stay on {PREFERRED_BATCH_NAME} ({KNOWN_BATCH_ID}). "
            "Do not Transfer AP. Do not invent Success. No Mail.Send. "
            f"{out.get('Why') or ''}"
        )
        updated.append(out)
    write_report(SHEET, updated)
    # write_report remaps missing_receipt owner to Ruben — force Shawn on these two.
    wb = load_workbook(SHEET)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    inv_i = headers.index("Invoice #")
    owner_i = headers.index("Exception owner")
    wanted = {spec["invoice"] for spec in TARGETS}
    for row in ws.iter_rows(min_row=2):
        inv = exact_invoice_number(row[inv_i].value)
        if inv in wanted:
            row[owner_i].value = SHAWN_OWNER
    wb.save(SHEET)
    sidecar = json.loads(SHEET_JSON.read_text()) if SHEET_JSON.is_file() else {}
    sidecar["proof"] = "gas-supply-0917-shawn-missing-receipt"
    sidecar["rows"] = updated
    sidecar["shawn_missing_receipt_0922"] = results
    SHEET_JSON.write_text(json.dumps(sidecar, indent=2, default=str) + "\n")
    return updated


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")
    if not args.live:
        raise SystemExit("Refusing: pass --live.")
    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error or "live credentials not ready")
    client = KimcoClient.authenticate(
        creds.instance_url, creds.key or "", creds.password or "", target="live"
    )
    verified = verify_batch(client, KNOWN_BATCH_ID, PREFERRED_BATCH_NAME)
    if not verified.get("matches_expected"):
        raise SystemExit(f"Batch 720 mismatch: {verified}")
    prior = prior_rows_from_sidecar(SHEET_JSON)
    proof: dict[str, Any] = {
        "proof": "gas-shawn-missing-receipt-2026-09-22",
        "invent": False,
        "mail_send": False,
        "transfer_ap": False,
        "select": False,
        "kyle_override": "ping Shawn on Gas missing_receipt HOLDs",
        "mention_id": SHAWN_MENTION_ID,
        "batch": KNOWN_BATCH_ID,
        "sheet_before": [_sheet_row_for(prior, spec["invoice"]) for spec in TARGETS],
        "dry_run": args.dry_run,
        "checked_at": datetime.now(timezone.utc).isoformat(),
    }
    if args.dry_run:
        proof["would"] = [snapshot(client, spec["kimco_id"]) for spec in TARGETS]
        PROOF.write_text(json.dumps(proof, indent=2, default=str) + "\n")
        print(json.dumps(proof, indent=2, default=str))
        return 0
    results = [ping_header(client, spec) for spec in TARGETS]
    rows = patch_sheet(results)
    proof["results"] = results
    proof["sheet_after"] = [_sheet_row_for(rows, spec["invoice"]) for spec in TARGETS]
    proof["sheet"] = str(SHEET)
    PROOF.write_text(json.dumps(proof, indent=2, default=str) + "\n")
    print(json.dumps(proof, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
