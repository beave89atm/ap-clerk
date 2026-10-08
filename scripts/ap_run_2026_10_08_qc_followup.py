"""Correct the 10/8 PO-not-found holds and the UniFirst four-cent tax hold.

Does not post, close a batch, auto-pay, delete a bill or a note, or send mail.
One sign-in, then at most one re-sign-in after a 401. Notes are added only
after each bill's final batch and total are known. New comments must be
authored as API Agent user 175.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from ap_clerk.kimco import KimcoClient, added_comment_payload
from ap_clerk.rules import lookup_id, lookup_text, money, ppv_limit
from scripts.ap_run_2026_10_08_enter import OUT, QC, move_batch, note_html
from scripts.sept25_30_finish_2026_10_06 import install_401_guard, login
from scripts.sept_missed_entry_2026_10_07 import put, totals

RUN_BATCH = 750
TRANSFER = 375
PRINT_TOTAL = 928.54
PRINT_TAX = 69.22
TOUCHED = (10552, 10555, 10556, 10557, 10558, 10559, 10562, 10564, 10565)


def plain(html: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", str(html or ""))).strip()


def comments_of(record: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for comment in (record.get("lists") or {}).get("Comments_1") or []:
        values = comment.get("values") or {}
        creator = values.get("CreatorId")
        rows.append(
            {
                "id": comment.get("id"),
                "html": str(values.get("HtmlValue") or ""),
                "text": plain(values.get("HtmlValue")),
                "creator_id": creator.get("id") if isinstance(creator, dict) else creator,
                "creator_name": creator.get("text") if isinstance(creator, dict) else "",
            }
        )
    return rows


def tax_rows(record: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for tax in (record.get("lists") or {}).get("APInvoiceTaxCodes") or []:
        values = tax.get("values") or {}
        rows.append(
            {
                "id": tax.get("id"),
                "amount": money(values.get("Tax_Amount")),
                "taxable": money(values.get("Taxable_Amount")),
                "rate": values.get("Tax_Rate"),
                "manual": values.get("Manual_Calculation"),
                "code": lookup_id(values.get("Tax_Code")),
            }
        )
    return rows


def state_of(record: dict[str, Any]) -> dict[str, Any]:
    info = totals(record)
    values = record.get("values") or {}
    notes = comments_of(record)
    return {
        "bill": int(record.get("id") or 0),
        "invoice": str(values.get("Invoice_Number") or ""),
        "vendor": lookup_text(values.get("Vendor")) or "",
        "posted": values.get("Posted"),
        "header": info.get("verification"),
        "covered": info.get("covered"),
        "batch": info.get("batch"),
        "batch_id": info.get("batch_id"),
        "lines": info.get("lines"),
        "charges": info.get("charges"),
        "taxes": info.get("taxes"),
        "po": lookup_text(values.get("Purchase_Order")) or "",
        "notes": notes,
        "receipts": [line.get("receipt_id") for line in info["lines"] if line.get("receipt_id")],
    }


def require_unposted(client: KimcoClient, invoice_id: int) -> dict[str, Any]:
    record = client.get_item("ap_invoices", invoice_id)
    if (record.get("values") or {}).get("Posted") not in (None, "", False):
        raise SystemExit(f"Bill {invoice_id} is posted. Not editing it.")
    return record


def add_note(client: KimcoClient, invoice_id: int, text: str, owner: str) -> dict[str, Any]:
    record = require_unposted(client, invoice_id)
    html = note_html(text, owner)
    prior = {int(row["id"]) for row in comments_of(record) if row.get("id") not in (None, "")}
    status = put(client, invoice_id, added_comment_payload(invoice_id, html))
    if status >= 400:
        raise SystemExit(f"Note on bill {invoice_id} failed HTTP {status}")
    after = require_unposted(client, invoice_id)
    added = [row for row in comments_of(after) if row.get("id") not in (None, "") and int(row["id"]) not in prior]
    if len(added) != 1:
        raise SystemExit(f"Bill {invoice_id} added {len(added)} notes")
    row = added[0]
    if int(row.get("creator_id") or 0) != 175 or str(row.get("creator_name") or "") != "API Agent":
        raise SystemExit(
            f"Bill {invoice_id} note {row.get('id')} author is not API Agent user 175. Stopping."
        )
    if not str(row.get("text") or "").startswith("AP Clerk:"):
        raise SystemExit(f"Bill {invoice_id} note does not start with AP Clerk:")
    if owner == "shawn" and 'data-mention-id="104"' not in str(row.get("html") or ""):
        raise SystemExit(f"Bill {invoice_id} note is missing Shawn")
    if owner == "treyce" and 'data-mention-id="33"' not in str(row.get("html") or ""):
        raise SystemExit(f"Bill {invoice_id} note is missing Treyce")
    kept = [int(item["id"]) for item in comments_of(after) if item.get("id") not in (None, "") and int(item["id"]) in prior]
    if len(kept) != len(prior):
        raise SystemExit(f"Bill {invoice_id} lost an existing note")
    return row


def bill_for_invoice(client: KimcoClient, invoice_id: int) -> dict[str, Any]:
    record = client.get_item("ap_invoices", invoice_id)
    values = record.get("values") or {}
    info = totals(record)
    return {
        "id": invoice_id,
        "invoice": str(values.get("Invoice_Number") or ""),
        "vendor": lookup_text(values.get("Vendor")) or "",
        "vendor_id": lookup_id(values.get("Vendor")),
        "posted": values.get("Posted"),
        "receipts": [line.get("receipt_id") for line in info["lines"] if line.get("receipt_id")],
    }


def fix_unifirst(client: KimcoClient) -> dict[str, Any]:
    limit = ppv_limit()
    if limit != 75.0:
        raise SystemExit(f"PPV limit is {limit}, not $75. Stopping.")
    record = require_unposted(client, 10562)
    before_tax = tax_rows(record)
    if len(before_tax) != 1 or before_tax[0]["code"] != 2:
        raise SystemExit(f"UniFirst tax rows are not one Sales Tax line: {before_tax}")
    covered = totals(record)["covered"]
    method = "already"
    tax_error = ""
    if abs(float(covered) - PRINT_TOTAL) >= 0.001:
        # Taxes-tab rows are not editable ("The list does not allow items to be edited").
        _body, status, error = client.update(
            "ap_invoices",
            10562,
            {
                "id": 10562,
                "state": "Modified",
                "lists": {
                    "APInvoiceTaxCodes": [
                        {
                            "id": before_tax[0]["id"],
                            "state": "Modified",
                            "values": {"Manual_Calculation": True, "Tax_Amount": PRINT_TAX},
                        }
                    ]
                },
            },
        )
        tax_error = (error or "")[:180]
        if status < 400:
            method = "tax"
            refreshed = require_unposted(client, 10562)
            if len(tax_rows(refreshed)) != 1:
                raise SystemExit(f"UniFirst tax update changed the tax row count: {tax_rows(refreshed)}")
            covered = totals(refreshed)["covered"]
    if abs(float(covered) - PRINT_TOTAL) >= 0.001:
        # KIMCO keeps Tax_Amount = taxable × the tax-code rate ($69.26).
        # The printed tax is $69.22. The four cents are inside the PPV limit.
        gap = round(PRINT_TOTAL - float(covered), 2)
        existing_ppv = 0.0
        for charge in totals(require_unposted(client, 10562))["charges"]:
            name = str(charge.get("name") or "")
            if "Price Variance" in name or name == "Purchase Price Variance":
                existing_ppv = round(existing_ppv + float(money(charge.get("amount")) or 0), 2)
        if abs(existing_ppv - gap) >= 0.001:
            ppv_status = client.try_post_ppv(10562, gap)
            if ppv_status != "posted":
                raise SystemExit(f"UniFirst PPV was not accepted ({ppv_status})")
        method = "ppv"
        covered = totals(require_unposted(client, 10562))["covered"]
    if abs(float(covered) - PRINT_TOTAL) >= 0.001:
        raise SystemExit(f"UniFirst still sums to {covered}, not {PRINT_TOTAL}")
    batch = client.get_item("ap_batches", RUN_BATCH)
    batch_status = (batch.get("values") or {}).get("Status")
    if batch_status not in (0, "0", None, ""):
        raise SystemExit(f"Batch 750 status is {batch_status}. Not moving the bill onto it.")
    move_batch(client, 10562, RUN_BATCH)
    after = totals(require_unposted(client, 10562))
    if int(after.get("batch_id") or 0) != RUN_BATCH:
        raise SystemExit("UniFirst did not land on batch 750")
    if after.get("posted") not in (None, "", False):
        raise SystemExit("UniFirst became posted")
    ppv = 0.0
    for charge in after["charges"]:
        name = str(charge.get("name") or "")
        if "Price Variance" in name or "PPV" in name:
            ppv = round(ppv + float(money(charge.get("amount")) or 0), 2)
    return {
        "method": method,
        "covered": after["covered"],
        "ppv": ppv,
        "tax": tax_rows(client.get_item("ap_invoices", 10562)),
        "tax_error": tax_error,
    }


def main() -> None:
    client = login()
    install_401_guard(client)
    if ppv_limit() != 75.0:
        raise SystemExit(f"PPV limit is {ppv_limit()}, not $75. Stopping.")
    transfer = client.get_item("ap_batches", TRANSFER)
    transfer_name = str((transfer.get("values") or {}).get("AP_Invoice_Batch_ID") or "")
    if transfer_name != "TRANSFER AP":
        raise SystemExit("Batch 375 is not TRANSFER AP. Stopping.")
    before = {bill_id: state_of(require_unposted(client, bill_id)) for bill_id in TOUCHED}
    invoiced_66 = bill_for_invoice(client, 10491)
    if invoiced_66["invoice"] != "18564" or int(invoiced_66["vendor_id"] or 0) != 183:
        raise SystemExit(f"Bill 10491 is not TPI invoice 18564: {invoiced_66}")
    selected_62 = bill_for_invoice(client, 10554)
    if selected_62["invoice"] != "18663" or 24987 not in selected_62["receipts"]:
        raise SystemExit(f"Bill 10554 does not hold receipt 24987: {selected_62}")

    findings = {
        10552: "PO 59271 is open Legacy Wire Products id 7273. Lines 01-02 are 17 ea and 33 ea 75-10-201007 at $41, received 0. No receipt on the line or in the receipts list.",
        10557: "PO 59043 is open TPI id 7045. Line 01 is 100 ea A-04421-000 at $108.66, received 0. No receipt recorded.",
        10558: "PO 59043 is open TPI id 7045. Same line, received 0. No receipt for 1 ea A-04421-000.",
        10559: "PO 59061 is open TPI id 7063. Line 01 is 100 ea A-04421-000 at $108.66, received 0. No receipt recorded.",
        10564: "PO 59283 is open Legacy Wire Products id 7285. Lines 01-02 are 27 ea and 17 ea 75-10-201007 at $41, received 0. No receipt recorded.",
        10565: "PO 59303 is open Legacy Wire Products id 7305. Line 02 is 10 ea 75-10-201007 at $41, received 0. No receipt recorded.",
        10555: (
            f"PO 58931 receipt 24982 is 66 ea A-04421-000 for $7,171.56 and is already invoiced "
            f"by bill {invoiced_66['id']} invoice {invoiced_66['invoice']} ({invoiced_66['vendor']}). "
            "No uninvoiced qty-66 receipt."
        ),
        10556: (
            "PO 58966 has no qty-38 receipt. Receipt 24987 is 62 ea A-04421-000 and is already "
            "invoiced by bill 10554 invoice 18663. Remaining quantity 38 is not received."
        ),
    }
    notes = {
        10552: "AP Clerk: Legacy Wire Products invoice PS-INV104080 stays on hold and is not posted. The PDF total is $2,506.50. Purchase order 59271 is open for Legacy Wire Products. No receipt recorded yet on PO 59271 for 17 ea and 33 ea 75-10-201007. Received quantity is 0. Nothing was selected. The bill stays in Transfer AP.",
        10557: "AP Clerk: Telecom Products invoice 18677 stays on hold and is not posted. The PDF total is $10,757.34. Purchase order 59043 is open for TPI. No receipt recorded yet on PO 59043 for 99 ea A-04421-000. Received quantity is 0 of 100. Nothing was selected. The bill stays in Transfer AP.",
        10558: "AP Clerk: Telecom Products invoice 18678 stays on hold and is not posted. The PDF total is $108.66. Purchase order 59043 is open for TPI. No receipt recorded yet on PO 59043 for 1 ea A-04421-000. Received quantity is 0. Nothing was selected. The bill stays in Transfer AP.",
        10559: "AP Clerk: Telecom Products invoice 18679 stays on hold and is not posted. The PDF total is $325.98. Purchase order 59061 is open for TPI. No receipt recorded yet on PO 59061 for 3 ea A-04421-000. Received quantity is 0 of 100. Nothing was selected. The bill stays in Transfer AP.",
        10564: "AP Clerk: Legacy Wire Products invoice PS-INV104085 stays on hold and is not posted. The PDF total is $2,205.72. Purchase order 59283 is open for Legacy Wire Products. No receipt recorded yet on PO 59283 for 27 ea and 17 ea 75-10-201007. Received quantity is 0. Nothing was selected. The bill stays in Transfer AP.",
        10565: "AP Clerk: Legacy Wire Products invoice PS-INV104086 stays on hold and is not posted. The PDF total is $501.30. Purchase order 59303 is open for Legacy Wire Products. No receipt recorded yet on PO 59303 for 10 ea 75-10-201007. Received quantity is 0. Nothing was selected. The bill stays in Transfer AP.",
        10555: (
            "AP Clerk: Telecom Products invoice 18664 stays on hold and is not posted. "
            f"The PDF total is $7,171.56. The qty-66 receipt 24982 on PO 58931 is 66 ea A-04421-000 "
            f"at $7,171.56 and is already invoiced by bill {invoiced_66['id']} invoice {invoiced_66['invoice']}. "
            "No open uninvoiced receipt matches. Nothing was selected. The bill stays in Transfer AP."
        ),
        10556: "AP Clerk: Telecom Products invoice 18676 stays on hold and is not posted. The PDF total is $4,129.08. No qty-38 receipt exists on PO 58966. Receipt 24987 is 62 ea A-04421-000 and is already invoiced by bill 10554 invoice 18663. The remaining 38 ea are not received. Nothing was selected. The bill stays in Transfer AP.",
    }
    for bill_id in (10552, 10555, 10556, 10557, 10558, 10559, 10564, 10565):
        live = before[bill_id]
        if int(live["batch_id"] or 0) != TRANSFER:
            raise SystemExit(f"Bill {bill_id} is on batch {live['batch_id']}, not Transfer AP. Not moving it.")
        if live["receipts"]:
            raise SystemExit(f"Bill {bill_id} already has receipts {live['receipts']}. Not selecting more.")

    unifirst = fix_unifirst(client)
    if unifirst["method"] == "tax":
        notes[10562] = (
            "AP Clerk: UniFirst invoice 2810822429 is entered and is not posted. "
            "The PDF total is $928.54. The printed sales tax of $69.22 is on the Taxes tab. "
            "The four-cent KIMCO rounding difference is not a hold. "
            "The bill is in batch API Agent - 10/8/26."
        )
        findings[10562] = "Tax override stored the printed $69.22. No receipt. Moved back to batch 750."
    else:
        notes[10562] = (
            "AP Clerk: UniFirst invoice 2810822429 is entered and is not posted. "
            "The PDF total is $928.54. KIMCO tax stays $69.26 because the tax code computes "
            "taxable $859.32 at 8.06 percent. A purchase price variance of -$0.04 makes the bill "
            "match the printed total. The four cents are under the $75 PPV limit. "
            "The bill is in batch API Agent - 10/8/26."
        )
        findings[10562] = (
            "Taxes tab does not allow edits. Tax stays $69.26 on taxable $859.32 at 8.06%. "
            f"PPV {unifirst['ppv']} makes the bill ${unifirst['covered']}. Moved to batch 750. No receipts."
        )

    written: dict[int, dict[str, Any]] = {}
    for bill_id, text in notes.items():
        owner = "treyce" if bill_id == 10562 else "shawn"
        existing = [
            row
            for row in comments_of(require_unposted(client, bill_id))
            if text.split(".", 1)[0] in row["text"]
        ]
        if existing:
            written[bill_id] = existing[-1]
            continue
        written[bill_id] = add_note(client, bill_id, text, owner)

    after = {bill_id: state_of(require_unposted(client, bill_id)) for bill_id in TOUCHED}
    if int(after[10562]["batch_id"] or 0) != RUN_BATCH or abs(float(after[10562]["covered"]) - PRINT_TOTAL) >= 0.001:
        raise SystemExit(f"UniFirst readback is not batch 750 at {PRINT_TOTAL}: {after[10562]['batch_id']} {after[10562]['covered']}")
    for bill_id in TOUCHED:
        if bill_id == 10562:
            continue
        if int(after[bill_id]["batch_id"] or 0) != TRANSFER:
            raise SystemExit(f"Bill {bill_id} left Transfer AP")
        if after[bill_id]["receipts"]:
            raise SystemExit(f"Bill {bill_id} gained receipts")

    QC.mkdir(parents=True, exist_ok=True)
    path = OUT / "qc-followup.csv"
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "bill",
                "invoice#",
                "before state",
                "after state",
                "PO lookup finding",
                "receipts selected",
                "PPV",
                "batch",
                "note id + text",
            ]
        )
        for bill_id in TOUCHED:
            old = before[bill_id]
            new = after[bill_id]
            note = written[bill_id]
            before_state = f"{old['batch']} covered {old['covered']}"
            after_state = "PASS" if bill_id == 10562 else "HOLD"
            after_state = f"{after_state} {new['batch']} covered {new['covered']}"
            ppv = ""
            if bill_id == 10562:
                ppv = f"{unifirst['ppv']:.2f}"
            writer.writerow(
                [
                    bill_id,
                    new["invoice"],
                    before_state,
                    after_state,
                    findings[bill_id],
                    "",
                    ppv,
                    new["batch"],
                    f"{note['id']}: {note['text']}",
                ]
            )

    book = load_workbook(OUT / "AP-run-2026-10-08.xlsx")
    sheet = book["Invoices"]
    headers = [cell.value for cell in sheet[1]]
    index = {name: pos for pos, name in enumerate(headers)}
    passed = 0
    held = 0
    for row in sheet.iter_rows(min_row=2):
        bill_id = row[index["Bill id"]].value
        if row[index["PASS/HOLD"]].value == "PASS":
            passed += 1
        elif row[index["PASS/HOLD"]].value == "HOLD":
            held += 1
        if bill_id not in written:
            continue
        note = written[int(bill_id)]
        new = after[int(bill_id)]
        if int(bill_id) == 10562:
            row[index["PASS/HOLD"]].value = "PASS"
            row[index["Owner"]].value = "treyce"
            row[index["Reason"]].value = ""
            row[index["Batch"]].value = new["batch"]
            passed += 1
            held -= 1
        else:
            row[index["Reason"]].value = findings[int(bill_id)]
        row[index["Note id"]].value = int(note["id"])
    summary = book["Summary"]
    for row in summary.iter_rows(min_row=2):
        if row[0].value == "Passed":
            row[1].value = passed
        elif row[0].value == "Held":
            row[1].value = held
    book.save(OUT / "AP-run-2026-10-08.xlsx")

    result_path = OUT / "result.json"
    if result_path.exists():
        payload = json.loads(result_path.read_text())
        for row in payload.get("rows") or []:
            bill_id = int(row.get("bill_id") or 0)
            if bill_id not in written:
                continue
            new = after[bill_id]
            row["batch"] = new["batch"]
            row["batch_id"] = new["batch_id"]
            row["covered"] = new["covered"]
            row["note_id"] = int(written[bill_id]["id"])
            row["note"] = written[bill_id]["text"]
            if bill_id == 10562:
                row["result"] = "PASS"
                row["reason"] = ""
                row["ppv"] = unifirst["ppv"]
            else:
                row["result"] = "HOLD"
                row["reason"] = findings[bill_id]
        result_path.write_text(json.dumps(payload, indent=2))
    print(json.dumps({"unifirst": unifirst, "notes": {str(k): v["id"] for k, v in written.items()}}, default=str))


if __name__ == "__main__":
    main()
