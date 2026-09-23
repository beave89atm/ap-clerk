"""Live enter: first 5 open 3P Industries invoices (Kyle 2026-09-22).

New batch `API Agent - 9/22/26 3P`. Multi-PO: blank header PO, Select
Receipts per PO. PPV |total| under $75 → finish path; $75+ → Transfer AP
+ Comments_1 @Shawn mention-id 104. missing_receipt → Transfer AP 375 +
plain @Ruben Perez (mention-id unknown). Success needs vendor PDF + a
signed receiving@ packing slip. Fort Worth move after header+attach.

No Mail.Send. invent=false. 3P only.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from openpyxl import load_workbook

import ap_clerk.rules as rules
from ap_clerk.auth import load_credentials
from ap_clerk.cli import _find_or_create_batch, _optional_graph_client, run_enter
from ap_clerk.graph import (
    ALLOWED_MAILBOX,
    FORT_WORTH_FOLDER_DISPLAY_NAME,
    FORT_WORTH_FOLDER_ID,
    GraphClient,
    GraphError,
    is_fort_worth_inbox_folder,
    load_graph_credentials,
)
from ap_clerk.inbox import sender_address, sender_name
from ap_clerk.kimco import KimcoClient, KimcoError
from ap_clerk.outlook_finish import promote_ap_outlook_after_success
from ap_clerk.pdf_invoice import (
    extract_3p_lines,
    extract_pdf_text,
    first_invoice_page,
    parse_date_value,
    parse_invoice_pdf,
    parse_invoice_text,
    parse_money,
)
from ap_clerk.quality_v12 import (
    COL_EXCEPTION_CATEGORY,
    COL_EXCEPTION_OWNER,
    apply_exception_category_owner,
)
from ap_clerk.receiving_attach import (
    attach_slip_to_existing_header,
    header_already_has_slip,
    index_live_ap_invoice,
    unique_invoice_for_slip,
    verify_header_packing_slip,
)
from ap_clerk.receiving_mail import ReceivingGraph, stamp_receiving_ai_completed
from ap_clerk.receiving_owners import (
    lookup_receiving_owner,
    missing_receipt_comments_1_html,
    missing_receipt_exception_owner,
)
from ap_clerk.receiving_probe import RECEIVING_MAILBOX
from ap_clerk.report import write_report
from ap_clerk.rules import (
    SHAWN_MCKIBBEN,
    extract_subject_invoice_number,
    extract_subject_pos,
    invoice_number_key,
    lookup_id,
    lookup_text,
    money,
    names_match,
    never_skip_vendor_invoice,
    parse_iso_date,
)
from ap_clerk.transfer_ap import apply_transfer_ap_batch_move
from gas_59081_retry import comments_1_shawn
from gas_shawn_finish import NOTE47_PPV_MAX_ABS
from missing_receipt_transfer_now import comments_1_plain

LOGGER = logging.getLogger("ap_clerk.three_p_0922")

PREFERRED_BATCH_NAME = "API Agent - 9/22/26 3P"
VENDOR_NAME = "3P"
VENDOR_ID_HINT = 1
CAP = 5
MIN_INVOICE_DATE = date(2026, 8, 1)
SHEET = ROOT / "runs" / "AP-run-2026-09-22-3p.xlsx"
SHEET_JSON = ROOT / "runs" / "AP-run-2026-09-22-3p.json"
PROOF = ROOT / "runs" / "three-p-0922.json"
PDF_DIR = ROOT / "runs" / "inbox-pdfs"
SHAWN_MENTION_ID = 104
RUBEN_TAG = "@Ruben Perez"

THREE_P_DOMAINS = ("3pindustries.com",)
THREE_P_NAME_NEEDLES = ("3p industries", "3p industry", "3p ind")
OTHER_VENDOR_NEEDLES = (
    "mcmaster",
    "o'neal",
    "oneal",
    "gasandsupply",
    "gas & supply",
    "gas and supply",
    "legacy wire",
    "jp steel",
    "fastenal",
)
RFQ_ONLY_RE = re.compile(r"\brfq\b", flags=re.I)
INVOICE_HINT_RE = re.compile(r"\b(inv(?:oice)?|cpl)\b", flags=re.I)
FREIGHT_FYI_RE = re.compile(r"freight\s+fyi|fyi\s+freight", flags=re.I)
STATEMENT_RE = re.compile(r"\b(statement|account\s+statement)\b", flags=re.I)
CORRECTED_RE = re.compile(r"\bcorrected\s+invoice\b", flags=re.I)
THREE_P_INV_RE = re.compile(r"\b(14\d{4})\b")
THREE_P_FALSE_INV = frozenset({"11942"})
THREE_P_DATE_INV_RE = re.compile(r"\b(\d{1,2}/\d{1,2}/20\d{2})\s+(14\d{4})\b")
THREE_P_TOTAL_RE = re.compile(r"\bTotal\s*[S$]?\s*([\d,]+\.\d{2})\b", flags=re.I)
THREE_P_LINE_RE = re.compile(
    r"(?m)^(?:\s*)(\d+(?:\.\d+)?)\s+(?:LINE\s+\d+:\s*)?([A-Z0-9]+(?:-[A-Z0-9]+)*)\b"
    r".+?\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})\s*$",
    flags=re.I,
)
THREE_P_PO_HEAD_RE = re.compile(r"(?m)^PO\s*#\s*(\d{5,6})\s*$", flags=re.I)
THREE_P_HEADER_PO_RE = re.compile(
    r"P\.?O\.?\s*Number\b.{0,80}?\b(5[7-9]\d{3})\b",
    flags=re.I | re.S,
)


def exact_invoice_number(value: Any) -> str:
    return invoice_number_key(value) or str(value or "").strip()


def prefer_3p_invoice_number(
    *,
    subject: str = "",
    parsed: str | None = None,
    pdf_text: str = "",
) -> str:
    """Never use 11942 (Tyler street). Prefer subject / Date+Invoice # 14xxxx."""
    subject_inv = exact_invoice_number(extract_subject_invoice_number(subject))
    parsed_inv = exact_invoice_number(parsed)
    dated = ""
    hit = THREE_P_DATE_INV_RE.search(pdf_text or "")
    if hit:
        dated = exact_invoice_number(hit.group(2))
    for candidate in (subject_inv, dated, parsed_inv, *THREE_P_INV_RE.findall(pdf_text or "")):
        if candidate in THREE_P_FALSE_INV:
            continue
        if re.fullmatch(r"14\d{4}", candidate or ""):
            return candidate
    if parsed_inv and parsed_inv not in THREE_P_FALSE_INV:
        return parsed_inv
    return subject_inv


def recover_3p_fields(bill: dict[str, Any], pdf_text: str) -> dict[str, Any]:
    """Fill invoice # / date / amount / lines from OCR without inventing."""
    page = first_invoice_page(pdf_text or "")
    number = prefer_3p_invoice_number(
        subject=str(bill.get("subject") or ""),
        parsed=str(bill.get("invoice_number") or ""),
        pdf_text=page,
    )
    if number:
        bill["invoice_number"] = number
    dated = THREE_P_DATE_INV_RE.search(page)
    if dated and not invoice_day(bill.get("date")):
        bill["date"] = parse_date_value(dated.group(1))
    if bill.get("amount") in (None, ""):
        totals = [parse_money(m) for m in THREE_P_TOTAL_RE.findall(page)]
        totals = [a for a in totals if a not in (None, 0)]
        if totals:
            bill["amount"] = totals[-1]
    pos = [str(p) for p in (bill.get("pos") or []) if p]
    header_po = THREE_P_HEADER_PO_RE.search(page)
    if header_po and header_po.group(1) not in pos:
        pos.append(header_po.group(1))
    lines = list(bill.get("lines") or [])
    if not lines:
        lines = extract_3p_lines(page)
    if not lines:
        current_po = pos[0] if len(pos) == 1 else None
        recovered: list[dict[str, Any]] = []
        for raw in page.splitlines():
            headed = THREE_P_PO_HEAD_RE.match(raw.strip())
            if headed:
                current_po = headed.group(1)
                if current_po not in pos:
                    pos.append(current_po)
                continue
            match = THREE_P_LINE_RE.match(raw.strip())
            if not match:
                continue
            qty, part, each, amt = match.groups()
            recovered.append(
                {
                    "part": part,
                    "qty": parse_money(qty),
                    "amount": parse_money(amt),
                    "unit_price": parse_money(each),
                    "po": current_po,
                    "po_line": None,
                    "wo": None,
                    "label": raw.strip()[:80],
                    "description": raw.strip()[:120],
                }
            )
        lines = recovered
    if lines:
        bill["lines"] = lines
        line_pos = [str(item.get("po") or "") for item in lines if item.get("po")]
        for po in line_pos:
            if po not in pos:
                pos.append(po)
        if bill.get("amount") in (None, ""):
            line_sum = round(
                sum(float(item["amount"]) for item in lines if item.get("amount") not in (None, "")),
                2,
            )
            if line_sum:
                bill["amount"] = line_sum
    pos = [p for p in pos if p]
    bill["pos"] = pos
    bill["multi_po"] = len(pos) > 1
    bill["po"] = None if bill["multi_po"] else (bill.get("po") or (pos[0] if pos else None))
    return bill


def invoice_day(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    iso = parse_date_value(str(value).strip())
    if not iso:
        return None
    try:
        return date.fromisoformat(iso)
    except ValueError:
        return None


def is_3p_vendor_text(value: str | None) -> bool:
    blob = str(value or "")
    if names_match(blob, "3P") or names_match(blob, "3P Industries"):
        return True
    low = blob.lower()
    return "3p industries" in low or re.search(r"\b3p\b", low) is not None


def is_other_vendor_email(msg: dict[str, Any]) -> bool:
    blob = " ".join(
        [
            sender_address(msg),
            sender_name(msg),
            str(msg.get("subject") or ""),
        ]
    ).lower()
    return any(n in blob for n in OTHER_VENDOR_NEEDLES)


def is_3p_email(msg: dict[str, Any]) -> bool:
    """Rachel/Dennis @ 3pindustries.com or 3P Industries subject. Never other vendors."""
    if is_other_vendor_email(msg):
        return False
    addr = sender_address(msg).lower()
    name = sender_name(msg).lower()
    subject = str(msg.get("subject") or "").lower()
    preview = str(msg.get("bodyPreview") or "").lower()
    if any(addr.endswith(d) or d in addr for d in THREE_P_DOMAINS):
        return True
    if any(n in name or n in subject for n in THREE_P_NAME_NEEDLES):
        return True
    if re.search(r"\b3p\b", f"{name} {subject} {preview}") and INVOICE_HINT_RE.search(subject):
        return True
    if "rachel" in name and INVOICE_HINT_RE.search(subject) and "3p" in f"{name} {subject} {preview}":
        return True
    return False


def skip_reason_for_message(msg: dict[str, Any]) -> str | None:
    subject = str(msg.get("subject") or "")
    if RFQ_ONLY_RE.search(subject) and not INVOICE_HINT_RE.search(subject):
        return "rfq"
    if FREIGHT_FYI_RE.search(subject) and not INVOICE_HINT_RE.search(subject):
        return "freight-fyi"
    if STATEMENT_RE.search(subject) and not INVOICE_HINT_RE.search(subject):
        return "statement"
    return None


def pick_oldest_open(
    bills: list[dict[str, Any]],
    entered: dict[str, int],
    *,
    cap: int = CAP,
    min_date: date = MIN_INVOICE_DATE,
) -> list[dict[str, Any]]:
    """Oldest invoice date, then received. Skip already-entered. Cap invoices not emails."""
    open_bills: list[dict[str, Any]] = []
    for bill in bills:
        number = exact_invoice_number(bill.get("invoice_number"))
        if not number:
            continue
        if number in entered:
            continue
        inv_day = invoice_day(bill.get("date"))
        if inv_day is not None and inv_day < min_date:
            continue
        open_bills.append(bill)

    def sort_key(bill: dict[str, Any]) -> tuple:
        inv_day = invoice_day(bill.get("date")) or date.max
        received = str(bill.get("receivedDateTime") or "")
        return (inv_day.isoformat(), received, exact_invoice_number(bill.get("invoice_number")))

    open_bills.sort(key=sort_key)
    return open_bills[: max(0, int(cap))]


def message_in_fort_worth(graph: GraphClient, message_id: str) -> bool:
    if not message_id:
        return False
    try:
        rec = graph.get_message(
            ALLOWED_MAILBOX,
            message_id,
            select="id,parentFolderId,categories,subject",
        )
    except GraphError:
        return False
    folder_id = str(rec.get("parentFolderId") or "")
    if FORT_WORTH_FOLDER_ID and folder_id == FORT_WORTH_FOLDER_ID:
        return True
    return False


def find_3p_messages(graph: GraphClient) -> list[dict[str, Any]]:
    seen: dict[str, dict[str, Any]] = {}
    needles = (
        "3P Industries",
        "3pindustries.com",
        "Rachel Bailey",
        "INV #",
        "CPL #",
    )
    for needle in needles:
        try:
            hits = graph.search_messages(ALLOWED_MAILBOX, needle, top=50)
        except Exception as exc:  # noqa: BLE001
            LOGGER.info("Graph search %s failed: %s", needle, type(exc).__name__)
            continue
        for msg in hits:
            if not is_3p_email(msg):
                continue
            mid = str(msg.get("id") or "")
            if mid:
                seen[mid] = msg
    try:
        listed = graph.list_messages(
            ALLOWED_MAILBOX,
            received_from=date(2026, 8, 1),
            oldest_first=True,
        )
    except Exception as exc:  # noqa: BLE001
        LOGGER.info("Graph list failed: %s", type(exc).__name__)
        listed = []
    for msg in listed:
        if not is_3p_email(msg):
            continue
        mid = str(msg.get("id") or "")
        if mid:
            seen[mid] = msg
    return list(seen.values())


def kimco_3p_numbers(client: KimcoClient) -> dict[str, int]:
    found: dict[str, int] = {}
    for item in client.list_items("ap_invoices"):
        vals = item.get("values") or {}
        vendor_txt = str(lookup_text(vals.get("Vendor") or vals.get("Vendor_$_Display_Name")) or "")
        posted_id = lookup_id(vals.get("Vendor"))
        if posted_id != VENDOR_ID_HINT and not is_3p_vendor_text(vendor_txt):
            continue
        number = exact_invoice_number(vals.get("Invoice_Number"))
        if number and item.get("id") not in (None, ""):
            found[number] = int(item["id"])
    return found


def bills_from_message(graph: GraphClient, message: dict[str, Any], pdf_dir: Path) -> list[dict[str, Any]]:
    subject = str(message.get("subject") or "")
    message_id = str(message.get("id") or "")
    from_name = sender_name(message)
    skip = skip_reason_for_message(message)
    pdfs: list[tuple[str, bytes]] = []
    if message.get("hasAttachments"):
        pdfs = graph.download_pdf_attachments(ALLOWED_MAILBOX, message_id)
    if skip and not pdfs:
        return [
            {
                "vendor": VENDOR_NAME,
                "invoice_number": extract_subject_invoice_number(subject) or "",
                "skip_reason": skip,
                "graph_message_id": message_id,
                "subject": subject,
                "receivedDateTime": message.get("receivedDateTime"),
            }
        ]
    bills: list[dict[str, Any]] = []
    pdf_dir.mkdir(parents=True, exist_ok=True)
    for name, content in pdfs:
        dest = pdf_dir / f"3p-{exact_invoice_number(extract_subject_invoice_number(subject) or name)}-{name}"
        dest.write_bytes(content)
        pdf_text = extract_pdf_text(dest)
        parsed = parse_invoice_text(
            pdf_text,
            subject=subject,
            from_name=from_name,
            from_address=sender_address(message),
            filename=name,
        )
        extras = list(parsed.pop("siblings", []) or [])
        rows = [parsed, *extras]
        for row in rows:
            recovered = recover_3p_fields(
                {
                    "invoice_number": row.get("invoice_number"),
                    "date": row.get("date"),
                    "po": row.get("po"),
                    "pos": list(row.get("pos") or []),
                    "amount": row.get("amount"),
                    "lines": list(row.get("lines") or []),
                    "subject": subject,
                },
                pdf_text,
            )
            row.update(
                {
                    k: recovered[k]
                    for k in ("invoice_number", "date", "po", "pos", "amount", "lines", "multi_po")
                    if k in recovered
                }
            )
            number = prefer_3p_invoice_number(
                subject=subject,
                parsed=row.get("invoice_number") or parsed.get("invoice_number"),
                pdf_text=pdf_text,
            )
            pos = list(row.get("pos") or parsed.get("pos") or extract_subject_pos(subject) or [])
            pos = [str(p) for p in pos if p]
            multi_po = len(pos) > 1 or bool(row.get("multi_po") or parsed.get("multi_po"))
            bill = {
                "vendor": VENDOR_NAME,
                "invoice_number": number,
                "date": row.get("date") or parsed.get("date"),
                "po": None if multi_po else (row.get("po") or parsed.get("po") or (pos[0] if pos else None)),
                "pos": pos,
                "multi_po": multi_po,
                "amount": row.get("amount") if row.get("amount") not in (None, "") else parsed.get("amount"),
                "lines": list(row.get("lines") or parsed.get("lines") or []),
                "fees": list(row.get("fees") or parsed.get("fees") or []),
                "pdf_path": str(dest),
                "graph_message_id": message_id,
                "subject": subject,
                "receivedDateTime": message.get("receivedDateTime"),
                "from_name": from_name,
                "field_sources": row.get("field_sources") or parsed.get("field_sources") or {},
            }
            if number:
                bills.append(bill)
    if not bills:
        number = exact_invoice_number(extract_subject_invoice_number(subject) or "")
        if number and never_skip_vendor_invoice(
            subject=subject,
            from_name=from_name,
            preview=str(message.get("bodyPreview") or ""),
        ):
            bills.append(
                {
                    "vendor": VENDOR_NAME,
                    "invoice_number": number,
                    "date": None,
                    "po": None,
                    "pos": extract_subject_pos(subject) or [],
                    "multi_po": len(extract_subject_pos(subject) or []) > 1,
                    "amount": None,
                    "hold_reason": "parse-error",
                    "graph_message_id": message_id,
                    "subject": subject,
                    "receivedDateTime": message.get("receivedDateTime"),
                    "from_name": from_name,
                    "pdf_unavailable": True,
                }
            )
    return bills


def comments_1_ruben(client: KimcoClient, kimco_id: int, *, invoice: str, po: str | None) -> dict[str, Any]:
    """Plain @Ruben Perez. Do not invent a mention-id node."""
    html = missing_receipt_comments_1_html(VENDOR_NAME) or (
        f"<p>{RUBEN_TAG} HOLD missing_receipt on 3P {invoice} "
        f"PO {po or 'n/a'}. Receiving must receive the PO so AP can "
        "Select Receipts. Mention-id for Ruben Perez is not proven. No email.</p>"
    )
    if "data-mention-id" in html:
        html = html.replace("data-mention-id", "data-mention-skipped")
    return comments_1_plain(client, int(kimco_id), html)


def comments_1_price(client: KimcoClient, kimco_id: int, html: str) -> dict[str, Any]:
    return comments_1_shawn(client, int(kimco_id), html)


def try_packing_slip(
    kimco: KimcoClient,
    recv: ReceivingGraph | None,
    row: dict[str, Any],
) -> dict[str, Any]:
    kid = row.get("KIMCO id")
    if kid in (None, ""):
        return {"status": "no-header", "attached": False}
    try:
        atts = kimco.list_attachments(int(kid))
    except KimcoError:
        atts = []
    if header_already_has_slip(atts):
        return {"status": "already-on-header", "attached": True}
    if recv is None:
        return {"status": "receiving-unavailable", "attached": False}
    pos = [str(p) for p in (row.get("_pos") or []) if p]
    po = str(row.get("PO") or "")
    if po:
        pos.append(po)
    invoice = exact_invoice_number(row.get("Invoice #"))
    needles = [n for n in (*pos, invoice) if n]
    hits: list[dict[str, Any]] = []
    for needle in needles[:6]:
        try:
            found = recv.list_messages(top=25)
        except Exception as exc:  # noqa: BLE001
            return {"status": f"receiving-list-failed:{type(exc).__name__}", "attached": False}
        for msg in found:
            blob = f"{msg.get('subject') or ''} {msg.get('bodyPreview') or ''}"
            if needle and needle in blob:
                hits.append(msg)
        if hits:
            break
    if not hits:
        return {"status": "no-matching-receiving-mail", "attached": False}
    invoice_index = [index_live_ap_invoice(kimco.get_item("ap_invoices", int(kid)))]
    attached = False
    last_mid = ""
    for msg in hits[:5]:
        mid = str(msg.get("id") or "")
        last_mid = mid
        try:
            pdfs = recv.download_pdf_attachments(RECEIVING_MAILBOX, mid)
        except Exception as exc:  # noqa: BLE001
            LOGGER.info("receiving download failed: %s", type(exc).__name__)
            continue
        for name, content in pdfs:
            slip = {"po": pos[0] if pos else "", "invoice_number": invoice, "source_filename": name}
            match = unique_invoice_for_slip(slip, [r for r in invoice_index if r])
            if match.get("status") != "matched":
                continue
            put = attach_slip_to_existing_header(
                kimco,
                {"id": int(kid), "invoice_number": invoice},
                content=content,
                filename=name,
            )
            if put.get("status") in {"attached", "already"}:
                attached = True
    if attached and last_mid:
        try:
            stamp_receiving_ai_completed(recv, last_mid)
        except Exception as exc:  # noqa: BLE001
            LOGGER.info("AI Completed stamp failed: %s", type(exc).__name__)
    try:
        atts = kimco.list_attachments(int(kid))
    except KimcoError:
        atts = []
    verified = bool(header_already_has_slip(atts) or verify_header_packing_slip(atts).get("ok"))
    return {"status": "attached" if attached or verified else "not-attached", "attached": bool(attached or verified)}


def apply_post_enter(
    client: KimcoClient,
    row: dict[str, Any],
    *,
    recv: ReceivingGraph | None,
    graph: GraphClient | None,
    siblings: dict[str, str],
) -> dict[str, Any]:
    kid = row.get("KIMCO id")
    out = dict(row)
    invoice = exact_invoice_number(row.get("Invoice #"))
    if kid in (None, ""):
        return apply_exception_category_owner(out)
    slip = try_packing_slip(client, recv, out)
    out["packing_slip"] = "Y" if slip.get("attached") else "N"
    out["packing_slip_status"] = slip.get("status")
    result = str(out.get("Result") or "")
    category = str(out.get(COL_EXCEPTION_CATEGORY) or "")
    why = str(out.get("Why") or "").lower()
    ppv = money(str(out.get("PPV") or "").replace("$", "").replace(",", "") or None)
    if ppv is None:
        m = re.search(r"(-?\$?\d+(?:\.\d+)?)", str(out.get("PPV") or ""))
        ppv = money(m.group(1)) if m else None

    if result == "Success" and not slip.get("attached"):
        out["Result"] = "HOLD"
        out[COL_EXCEPTION_CATEGORY] = out.get(COL_EXCEPTION_CATEGORY) or "other"
        extra = (
            "HOLD incomplete: vendor PDF attached but no signed receiving@ "
            "packing slip on the header. Do not invent Success."
        )
        out["Why"] = f"{out.get('Why') or ''} {extra}".strip()
        result = "HOLD"

    if result == "HOLD" and (
        category == "missing_receipt" or "no receipts" in why or "no open receipt" in why
    ):
        moved = apply_transfer_ap_batch_move(client, kimco_id=int(kid))
        out["Batch"] = f"TRANSFER AP ({moved.get('batch_id') or 375})"
        out["transfer_ap"] = moved.get("status")
        out[COL_EXCEPTION_CATEGORY] = "missing_receipt"
        out[COL_EXCEPTION_OWNER] = missing_receipt_exception_owner(VENDOR_NAME)
        comment = comments_1_ruben(client, int(kid), invoice=invoice, po=str(out.get("PO") or ""))
        out["Comments_1"] = comment
        if "data-mention-id" in str((comment or {}).get("html") or ""):
            out["Comments_1"] = {**comment, "status": "refused-invented-mention"}
        extra = (
            f"NOTE-53 missing_receipt → Transfer AP ({moved.get('status')}). "
            f"Comments_1 {RUBEN_TAG} plain (mention-id unknown)."
        )
        out["Why"] = f"{out.get('Why') or ''} {extra}".strip()
        out["Notes"] = f"{out.get('Notes') or ''} {extra}".strip()

    over_ppv = (ppv is not None and abs(float(ppv)) >= NOTE47_PPV_MAX_ABS) or category == "price_variance"
    if result == "HOLD" and over_ppv and category != "missing_receipt":
        moved = apply_transfer_ap_batch_move(client, kimco_id=int(kid))
        out["Batch"] = f"TRANSFER AP ({moved.get('batch_id') or 375})"
        out["transfer_ap"] = moved.get("status")
        out[COL_EXCEPTION_CATEGORY] = "price_variance"
        out[COL_EXCEPTION_OWNER] = "Shawn McKibben"
        html = (
            f'<p><span data-mention-id="{SHAWN_MENTION_ID}" '
            f'data-mention-name="Shawn McKibben" '
            f'data-mention-email="Shawn.McKibben@kannonmfg.com" '
            f'class="prosemirror-mention-node">{SHAWN_MCKIBBEN}</span> '
            f"HOLD (price-does-not-match) on 3P {invoice} PO {out.get('PO') or 'n/a'}. "
            f"|bill PPV| is $75 or more. Receipts were NOT selected on over-gate lines. "
            "No email.</p>"
        )
        out["Comments_1"] = comments_1_price(client, int(kid), html)
        extra = f"Over-PPV → Transfer AP ({moved.get('status')}) + @Shawn mention-id {SHAWN_MENTION_ID}."
        out["Why"] = f"{out.get('Why') or ''} {extra}".strip()

    if str(out.get("Result") or "") == "Success" and graph is not None:
        sibling_results = dict(siblings)
        sibling_results[invoice] = "Success"
        try:
            promote_ap_outlook_after_success(
                graph,
                message_id=str(out.get("_message_id") or ""),
                sibling_results=sibling_results,
            )
        except Exception as exc:  # noqa: BLE001
            LOGGER.info("NOTE-51 promote failed: %s", type(exc).__name__)
    return apply_exception_category_owner(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s %(message)s")
    if not args.live:
        raise SystemExit("Refusing: pass --live.")

    rules.PPV_MAX_ABS_ON_BILL = NOTE47_PPV_MAX_ABS

    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error or "live credentials not ready")
    client = KimcoClient.authenticate(
        creds.instance_url, creds.key or "", creds.password or "", target="live"
    )
    graph = _optional_graph_client()
    if graph is None:
        raise SystemExit("Graph client required for 3P inbox enter")

    entered = kimco_3p_numbers(client)
    messages = find_3p_messages(graph)
    skipped: list[dict[str, Any]] = []
    catalog: list[dict[str, Any]] = []
    bills: list[dict[str, Any]] = []
    for msg in messages:
        mid = str(msg.get("id") or "")
        in_fw = message_in_fort_worth(graph, mid)
        skip = skip_reason_for_message(msg)
        catalog.append(
            {
                "id": mid[-20:] if mid else "",
                "subject": str(msg.get("subject") or "")[:160],
                "from": sender_address(msg),
                "received": msg.get("receivedDateTime"),
                "in_fort_worth": in_fw,
                "skip": skip,
            }
        )
        if in_fw:
            skipped.append({"reason": "already-fort-worth", "subject": msg.get("subject")})
            continue
        if skip:
            skipped.append({"reason": skip, "subject": msg.get("subject")})
            continue
        received_day = invoice_day(str(msg.get("receivedDateTime") or "")[:10])
        if received_day is not None and received_day < MIN_INVOICE_DATE:
            skipped.append({"reason": "received-before-aug1", "subject": msg.get("subject")})
            continue
        bills.extend(bills_from_message(graph, msg, PDF_DIR))

    wanted = pick_oldest_open(bills, entered, cap=CAP)
    proof: dict[str, Any] = {
        "proof": "three-p-0922",
        "invent": False,
        "mail_send": False,
        "batch_name": PREFERRED_BATCH_NAME,
        "ppv_gate": NOTE47_PPV_MAX_ABS,
        "receiving_owner": lookup_receiving_owner(VENDOR_NAME),
        "kimco_already": entered,
        "catalog": catalog,
        "skipped": skipped,
        "wanted": [exact_invoice_number(b.get("invoice_number")) for b in wanted],
        "dry_run": args.dry_run,
    }
    if args.dry_run:
        PROOF.write_text(json.dumps(proof, indent=2, default=str) + "\n")
        print(json.dumps(proof, indent=2, default=str))
        return 0

    batches = client.list_items("ap_batches")
    batch = _find_or_create_batch(client, batches, PREFERRED_BATCH_NAME)
    proof["batch"] = {"id": batch.get("id"), "name": batch.get("name"), "created": batch.get("created")}

    recv = None
    try:
        gcreds = load_graph_credentials()
        if gcreds.token:
            recv = ReceivingGraph(gcreds.token)
            recv.list_messages(top=1)
    except Exception as exc:  # noqa: BLE001
        LOGGER.info("receiving@ unavailable: %s", type(exc).__name__)
        recv = None
        proof["receiving"] = "unavailable"

    rows = run_enter(
        client,
        wanted,
        batch_name=PREFERRED_BATCH_NAME,
        pdf_dir=PDF_DIR,
        graph_client=graph,
        mailbox=ALLOWED_MAILBOX,
        flag_outlook=True,
    )
    siblings = {
        exact_invoice_number(r.get("Invoice #")): str(r.get("Result") or "") for r in rows
    }
    finished: list[dict[str, Any]] = []
    for row, bill in zip(rows, wanted):
        row["_pos"] = bill.get("pos") or []
        row["_message_id"] = bill.get("graph_message_id")
        row["outlook_move"] = "Y" if "NOTE-52" in str(row.get("Why") or "") or "Fort Worth" in str(row.get("Notes") or row.get("Why") or "") else ""
        finished.append(apply_post_enter(client, row, recv=recv, graph=graph, siblings=siblings))

    for row in finished:
        kid = row.get("KIMCO id")
        if kid in (None, ""):
            continue
        try:
            live = client.get_item("ap_invoices", int(kid))
        except KimcoError:
            continue
        vals = live.get("values") or {}
        row["_live_batch"] = lookup_text(vals.get("AP_Invoice_Batch"))
        row["_live_batch_id"] = lookup_id(vals.get("AP_Invoice_Batch"))
        row["_live_amount"] = money(vals.get("Invoice_Amount"))
        row["_live_verification"] = money(vals.get("Invoice_Verification_Amount"))
        try:
            atts = client.list_attachments(int(kid))
        except KimcoError:
            atts = []
        row["invoice_pdf"] = "Y" if any(str(a.get("name") or "").lower().endswith(".pdf") for a in atts) else "N"
        if not row.get("packing_slip"):
            row["packing_slip"] = "Y" if header_already_has_slip(atts) else "N"
        comments = (live.get("lists") or {}).get("Comments_1") or []
        row["Comments_1_ids"] = [c.get("id") for c in comments if isinstance(c, dict)]
        if row.get("_live_batch_id") == 375:
            row["Batch"] = f"TRANSFER AP (375)"
        elif row.get("_live_batch_id"):
            row["Batch"] = f"{row.get('_live_batch')} ({row.get('_live_batch_id')})"

    write_report(SHEET, finished)
    wb = load_workbook(SHEET)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if "Notes" in headers:
        notes_i = headers.index("Notes")
        inv_i = headers.index("Invoice #")
        by_inv = {exact_invoice_number(r.get("Invoice #")): r for r in finished}
        for excel_row in ws.iter_rows(min_row=2):
            hit = by_inv.get(exact_invoice_number(excel_row[inv_i].value))
            if hit:
                excel_row[notes_i].value = hit.get("Notes") or excel_row[notes_i].value
    wb.save(SHEET)
    sidecar = {
        "proof": "three-p-0922",
        "invent": False,
        "mail_send": False,
        "batch": proof.get("batch"),
        "rows": finished,
        "skipped": skipped,
        "wanted": proof["wanted"],
        "kimco_already": entered,
    }
    SHEET_JSON.write_text(json.dumps(sidecar, indent=2, default=str) + "\n")
    proof["rows"] = [
        {
            "invoice": r.get("Invoice #"),
            "kimco_id": r.get("KIMCO id"),
            "result": r.get("Result"),
            "batch": r.get("Batch"),
            "category": r.get(COL_EXCEPTION_CATEGORY),
            "owner": r.get(COL_EXCEPTION_OWNER),
            "invoice_pdf": r.get("invoice_pdf"),
            "packing_slip": r.get("packing_slip"),
            "comments_1": r.get("Comments_1_ids") or r.get("Comments_1"),
        }
        for r in finished
    ]
    proof["success"] = sum(1 for r in finished if r.get("Result") == "Success")
    proof["hold"] = sum(1 for r in finished if r.get("Result") == "HOLD")
    PROOF.write_text(json.dumps(proof, indent=2, default=str) + "\n")
    print(json.dumps(proof, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
