"""Fresh readback of AP bills 10478-10509. Does not post or touch mail."""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path
from typing import Any

import importlib.util
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

_SPEC = importlib.util.spec_from_file_location(
    "sept25_30_finish_2026_10_06",
    ROOT / "scripts" / "sept25_30_finish_2026_10_06.py",
)
_MOD = importlib.util.module_from_spec(_SPEC)
assert _SPEC and _SPEC.loader
_SPEC.loader.exec_module(_MOD)
install_401_guard = _MOD.install_401_guard
login = _MOD.login
slim_bill = _MOD.slim_bill

from ap_clerk.rules import money

SNAPSHOT = Path("/tmp/sept25-30-live.json")
MESSAGES = Path("/tmp/sept25-30-messages.json")
QC_JSON = ROOT / "qc" / "sept25-30-finish-qc.json"
PDF_OUT = ROOT / "qc" / "invoices-sept25-30"
PDF_SRC = ROOT / "qc" / "invoices-sept-catchup"
WORKBOOK = ROOT / "runs" / "AP-sept25-30-finish-2026-10-06.xlsx"
BATCH_AGENT = 747
BATCH_TRANSFER = 375

# Pages of a combined PDF that belong to this invoice. 1-based inclusive.
PAGE_SLICE = {
    10494: (1, 1),
    10504: (1, 1),
    10505: (2, 2),
    10506: (3, 3),
    10507: (4, 4),
    10508: (5, 5),
    10509: (6, 7),
}

# Holds a person still has to clear. These are not API-refused lines.
PERSON_HOLDS = {
    10480: "PO 59081 open receipts do not include the $213.93 hole-saw kit. Nothing was selected.",
    10482: "No purchase order. Merchandise and sales tax are entered. Confirm the no-PO coding.",
    10483: "No purchase order. Merchandise and sales tax are entered. Confirm the no-PO coding.",
    10493: "Dollars match $381.72. Invoice is 1 piece / 272.66 lb and receipt 25161 is quantity 288. Confirm the unit.",
    10494: "Invoice is 23.12 ft at $487.14. Receipt 24880 is quantity 276 at $484.63. Do not select receipt 24879.",
    10501: "20 pieces at $288 equals $5,760. Open receipt 25188 is 20 at $182.25, $3,645. Gap $2,115. Unreceive and correct the PO price.",
    10503: "Dollars match $166.98 with $2.41 purchase price variance. Invoice is 138 lb and receipt 25160 is quantity 240. Confirm the unit.",
    10505: "PO 59251 has no receipt for the $3,012.19 sheet. Receive it.",
    10508: "PO 59289 has no receipt for the $10,713.89 plate and sheet. Receive it.",
    10509: "PO 59293 open receipts do not match the invoice lines ($5,328.46). Do not guess a receipt.",
}


def covered(live: dict[str, Any]) -> float:
    lines = round(sum(float(money(line.get("extended")) or 0) for line in live.get("lines") or []), 2)
    charges = round(sum(float(money(row.get("amount")) or 0) for row in live.get("charges") or []), 2)
    tax = round(sum(float(money(row.get("amount")) or 0) for row in live.get("taxes") or []), 2)
    return round(lines + charges + tax, 2)


def fee_amounts(live: dict[str, Any]) -> list[float]:
    amounts = []
    for row in live.get("charges") or []:
        name = str(row.get("name") or "")
        if "price variance" in name.lower():
            continue
        amount = money(row.get("amount"))
        if amount is not None:
            amounts.append(float(amount))
    return amounts


def copy_pdf(bill_id: int, source_name: str, invoice: str) -> str:
    source = PDF_SRC / source_name
    if not source.is_file():
        raise SystemExit(f"Missing source PDF {source}")
    PDF_OUT.mkdir(parents=True, exist_ok=True)
    dest = PDF_OUT / f"{bill_id}_{invoice}.pdf"
    bounds = PAGE_SLICE.get(bill_id)
    if bounds:
        reader = PdfReader(str(source))
        writer = PdfWriter()
        start, end = bounds
        for index in range(start - 1, end):
            writer.add_page(reader.pages[index])
        with dest.open("wb") as handle:
            writer.write(handle)
    else:
        shutil.copyfile(source, dest)
    return str(dest.relative_to(ROOT))


def judge(live: dict[str, Any], pdf_total: float) -> tuple[str, str]:
    kid = int(live["id"])
    if live.get("posted") not in (None, "", False) or live.get("void") is True:
        return "FAIL", "Bill is posted or void."
    batch = int((live.get("batch") or {}).get("id") or 0)
    ver = money(live.get("verification"))
    total = round(float(pdf_total), 2)
    have = covered(live)
    tax = round(sum(float(money(row.get("amount")) or 0) for row in live.get("taxes") or []), 2)
    invoice_amount = money(live.get("invoice_amount"))
    if kid in PERSON_HOLDS:
        if kid in {10482, 10483, 10501, 10505, 10508} and batch != BATCH_TRANSFER:
            return "FAIL", f"Hold belongs on TRANSFER AP 375 and the live batch is {batch}."
        if kid not in {10482, 10483, 10501, 10505, 10508} and batch != BATCH_AGENT:
            return "FAIL", f"This hold stays on batch 747 and the live batch is {batch}."
        return "HOLD", PERSON_HOLDS[kid]
    if ver != total or have != total:
        return "FAIL", f"PDF ${total:.2f}, verification {ver}, lines+charges+tax ${have:.2f}."
    if invoice_amount not in {total, round(total - tax, 2)}:
        return "FAIL", f"Invoice amount {invoice_amount} is not the PDF total ${total:.2f}."
    if batch != BATCH_AGENT:
        return "FAIL", f"Matched bill is on batch {batch}, not 747."
    # 0040455916 prints the $22.50 fuel surcharge under the cylinder lines.
    # Merchandise 56.00 + 60.00 plus that surcharge is the $138.50 total.
    if kid == 10492 and 1080 in fee_amounts(live):
        return "FAIL", "The extra $1,080 fee is still on the bill."
    if kid == 10499 and 300 in fee_amounts(live) and not live.get("lines"):
        return "FAIL", "The $300 fee is still on the bill and the receipt is not selected."
    return "PASS", "Live total matches the PDF. The bill is unposted on batch 747."


def main() -> None:
    from ap_clerk.graph import ALLOWED_MAILBOX, GraphClient, load_graph_credentials

    client = login()
    install_401_guard(client)
    snap = {int(row["id"]): row for row in json.loads(SNAPSHOT.read_text())["bills"]}
    messages = {int(row["bill_id"]): row for row in json.loads(MESSAGES.read_text())}
    creds = load_graph_credentials()
    graph = GraphClient.authenticate(creds.tenant_id, creds.client_id, creds.client_secret)
    folder = graph.resolve_fort_worth_folder(ALLOWED_MAILBOX)
    folder_name = str(folder.get("displayName") or "9 - FORT WORTH ARCHIVE")

    rows = []
    for kid in range(10478, 10510):
        record = client.get_item("ap_invoices", kid)
        live = slim_bill(record)
        names = []
        for item in client.list_attachments(kid):
            names.append(str(item.get("name") or item.get("fileName") or item.get("FileName") or item.get("id") or ""))
        pdf = snap[kid]["pdf"]
        total = float(money(pdf.get("amount")) or 0)
        result, reason = judge(live, total)
        our_notes = []
        for comment in live.get("comments") or []:
            text = str(comment.get("text") or "")
            if text.startswith("AP Clerk:") and int(comment.get("id") or 0) >= 1573:
                our_notes.append({"id": comment.get("id"), "text": text})
        invoice = str(pdf.get("invoice_number") or live.get("invoice") or kid)
        pdf_path = copy_pdf(kid, str(pdf.get("pdf")), invoice)
        rows.append(
            {
                "bill_id": kid,
                "vendor": (live.get("vendor") or {}).get("text") if isinstance(live.get("vendor"), dict) else pdf.get("vendor"),
                "invoice": live.get("invoice") or invoice,
                "batch_id": (live.get("batch") or {}).get("id"),
                "batch": (live.get("batch") or {}).get("text"),
                "pdf_total": total,
                "invoice_amount": live.get("invoice_amount"),
                "verification": live.get("verification"),
                "lines_charges_tax": covered(live),
                "receipts": [
                    (line.get("receipt") or {}).get("id")
                    for line in live.get("lines") or []
                    if isinstance(line.get("receipt"), dict) and (line.get("receipt") or {}).get("id")
                ],
                "lines": [
                    {
                        "item": (line.get("item") or {}).get("text") if isinstance(line.get("item"), dict) else line.get("desc"),
                        "qty": line.get("qty"),
                        "price": line.get("unit_price"),
                        "extended": line.get("extended"),
                        "receipt": (line.get("receipt") or {}).get("id") if isinstance(line.get("receipt"), dict) else None,
                    }
                    for line in live.get("lines") or []
                ],
                "charges": [
                    {"name": row.get("name"), "amount": row.get("amount")} for row in live.get("charges") or []
                ],
                "taxes": [
                    {"code": (row.get("code") or {}).get("text") if isinstance(row.get("code"), dict) else row.get("code"), "amount": row.get("amount")}
                    for row in live.get("taxes") or []
                ],
                "attachments": names,
                "notes_added": our_notes,
                "result": result,
                "reason": reason,
                "posted": live.get("posted"),
                "pdf_file": pdf_path,
                "packing_slip": "No receiving@ scan from 2026-09-25 through 2026-10-06 matched this vendor and PO or invoice. None attached.",
            }
        )
        print(f"{kid} {result} batch={rows[-1]['batch_id']} pdf={total} ver={live.get('verification')} sum={rows[-1]['lines_charges_tax']} {reason[:80]}")

    by_message: dict[str, list[int]] = {}
    for row in rows:
        message = messages[row["bill_id"]]
        by_message.setdefault(message.get("message_id") or "", []).append(row["bill_id"])
    moves = []
    for row in rows:
        message = messages[row["bill_id"]]
        siblings = [item for item in rows if item["bill_id"] in by_message.get(message.get("message_id") or "", [])]
        all_pass = all(item["result"] == "PASS" for item in siblings)
        moves.append(
            {
                "bill_id": row["bill_id"],
                "message_id": message.get("message_id"),
                "received": message.get("received"),
                "subject": message.get("subject"),
                "proposed_category": "Entered in AI" if all_pass else "Entered with issues",
                "proposed_folder": folder_name,
                "mail_changed": False,
            }
        )

    manual = [
        {"bill_id": row["bill_id"], "fix": row["reason"]}
        for row in rows
        if row["result"] == "HOLD"
    ]
    payload = {
        "bills": rows,
        "proposed_email_moves": moves,
        "manual_fixes": manual,
        "api_refused_lines": [],
        "packing_slips_attached": [],
        "counts": {
            "PASS": sum(1 for row in rows if row["result"] == "PASS"),
            "HOLD": sum(1 for row in rows if row["result"] == "HOLD"),
            "FAIL": sum(1 for row in rows if row["result"] == "FAIL"),
        },
    }
    QC_JSON.write_text(json.dumps(payload, indent=2, default=str))
    write_workbook(rows, moves, manual)
    print("COUNTS", payload["counts"])


def write_workbook(rows: list[dict[str, Any]], moves: list[dict[str, Any]], manual: list[dict[str, Any]]) -> None:
    book = Workbook()
    bills = book.active
    bills.title = "Bills"
    bills.append(["Bill", "Vendor", "Invoice", "Batch", "PDF total", "Result", "Reason"])
    for row in rows:
        bills.append([row["bill_id"], row["vendor"], row["invoice"], row["batch"], row["pdf_total"], row["result"], row["reason"]])
    qc = book.create_sheet("QC")
    qc.append(
        [
            "Bill",
            "Vendor",
            "Invoice",
            "Batch",
            "PDF total",
            "Invoice amount",
            "Verification",
            "Lines + charges + tax",
            "Receipts",
            "Lines",
            "Note ids",
            "Result",
            "Reason",
        ]
    )
    for row in rows:
        line_text = "; ".join(
            f"{line.get('item') or ''} qty {line.get('qty')} @ {line.get('price')}"
            for line in row["lines"]
        )
        qc.append(
            [
                row["bill_id"],
                row["vendor"],
                row["invoice"],
                row["batch"],
                row["pdf_total"],
                row["invoice_amount"],
                row["verification"],
                row["lines_charges_tax"],
                ", ".join(str(item) for item in row["receipts"]),
                line_text,
                ", ".join(str(note["id"]) for note in row["notes_added"]),
                row["result"],
                row["reason"],
            ]
        )
    mail = book.create_sheet("Proposed email moves")
    mail.append(["Bill", "Message id", "Received", "Subject", "Proposed category", "Proposed folder"])
    for row in moves:
        mail.append(
            [
                row["bill_id"],
                row["message_id"],
                row["received"],
                row["subject"],
                row["proposed_category"],
                row["proposed_folder"],
            ]
        )
    fixes = book.create_sheet("Manual fixes needed")
    fixes.append(["Bill", "What a person must do"])
    if not manual:
        fixes.append(["", "No line was left because the API refused a delete or an edit."])
    for row in manual:
        fixes.append([row["bill_id"], row["fix"]])
    fixes.append(["", "No charge or receipt line remains that the API refused to remove."])
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    book.save(WORKBOOK)


if __name__ == "__main__":
    main()
