"""Enter the remaining 3P Industries invoice on live KIMCO. invent=false. Do not post.

Mailbox inventory (accountspayable@, all folders, received on or after 2026-08-01)
found one invoice that is not already a vendor-1 bill: 142280, the parts invoice
whose freight 142281 is KIMCO 10311. It is page 1 of the 2026-09-10 PDF whose
subject lists the packing-list numbers 76945–76947.

142173 stays excluded. The vendor PDF still shows qty 10 of 2975-1, not a
corrected invoice for 20. Existing bills, including HOLDs 10306, 10308, and
10312, are not modified.

142280 is parts. Header PO stays blank (PDF says SEE BELOW). Open receipts
whose part, PO, and qty match are selected at the receipt's PO extended cost.
A unit-price gap under $75 is selected and is not stacked as a unit-price PPV.
A unit-price gap of $75 or more is not selected. The bill gap is then $75 or
more, so the bill is HOLD price_variance: Comments_1 @Shawn (mention id 104),
then Transfer AP. No PPV Additional Charge is posted.

Packing-slip Success gate is suspended. No Mail.Send.
After header + this invoice's own page, the source message moves to
9 - FORT WORTH ARCHIVE.
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(ROOT / "scripts"))

from ap_clerk.auth import load_credentials
from ap_clerk.cli import _optional_graph_client
from ap_clerk.graph import ALLOWED_MAILBOX, FORT_WORTH_FOLDER_DISPLAY_NAME, GraphClient
from ap_clerk.kimco import KimcoClient, receipt_line_values_from_records, select_receipts_payload
from ap_clerk.rules import kimco_datetime, lookup_id, lookup_text, money
from ap_clerk.transfer_ap import apply_transfer_ap_batch_move
from gas_59081_retry import SHAWN_MENTION_ID, comments_1_shawn

BATCH_ID = 724
VENDOR_ID = 1
TERMS_ID = 4
CURRENCY_ID = 3
REMIT_ID = 120
PPV_HOLD = 75.0
DO_NOT_TOUCH_IDS = {10186, 10187, 10306, 10308, 10312}
PDF_PATH = Path("/tmp/3p-finish/pdfs/2026-09-10_09102026_007.pdf")
PAGES_JSON = Path("/tmp/3p-finish/pages.json")
OUT_JSON = ROOT / "artifacts" / "3p-finish-2026-09-25.json"
OUT_XLSX = ROOT / "artifacts" / "AP-3P-finish-2026-09-25.xlsx"
RUN_HEADERS = [
    "Invoice #",
    "Status",
    "Reason",
    "Type",
    "Amount",
    "KIMCO bill #",
    "Batch",
    "Transfer AP",
    "Receipts selected",
    "PPV",
    "Comments_1 id",
    "Email moved",
    "Note",
]

# Confirmed by OCR of the invoice page and by live open receipts on 2026-09-25.
SPEC: dict[str, Any] = {
    "invoice": "142280",
    "kind": "parts",
    "date": date(2026, 8, 28),
    "due": date(2026, 10, 27),
    "total": 1631.56,
    "pages": [0],
    "page_count": 20,
    "pos": ["58952", "58925", "58972", "59002", "59031"],
    "po_on_pdf": "SEE BELOW",
    "message_subject_needle": "INV # 76945, 76946",
    "receipts": [23918, 23919, 23913, 23914, 23922, 23915, 23916],
    "receipt_expect": {
        23918: {"qty": 5.0, "price": 12.30},
        23919: {"qty": 1.0, "price": 7.76},
        23913: {"qty": 25.0, "price": 19.82},
        23914: {"qty": 3.0, "price": 83.02},
        23922: {"qty": 2.0, "price": 83.02},
        23915: {"qty": 8.0, "price": 3.65},
        23916: {"qty": 20.0, "price": 3.65},
    },
    "small_gaps": [
        {
            "id": 23913,
            "part": "33213-1",
            "po": "58972",
            "qty": 25,
            "po_unit": 19.82,
            "po_ext": 495.50,
            "invoice_unit": 22.30,
            "invoice_ext": 557.50,
        }
    ],
    "skip_receipts": [
        {
            "id": 23917,
            "part": "1007044-1",
            "po": "59002",
            "qty": 2,
            "po_unit": 361.46,
            "po_ext": 722.92,
            "invoice_unit": 97.50,
            "invoice_ext": 195.00,
        },
        {
            "id": 23920,
            "part": "1007044-1",
            "po": "59031",
            "qty": 2,
            "po_unit": 361.46,
            "po_ext": 722.92,
            "invoice_unit": 97.50,
            "invoice_ext": 292.50,
            "combined_with": 23921,
        },
        {
            "id": 23921,
            "part": "1007044-1",
            "po": "59031",
            "qty": 1,
            "po_unit": 361.46,
            "po_ext": 361.46,
            "invoice_unit": 97.50,
            "invoice_ext": 292.50,
            "combined_with": 23920,
        },
    ],
}


def page_bytes() -> bytes:
    reader = PdfReader(str(PDF_PATH))
    if len(reader.pages) != SPEC["page_count"]:
        raise SystemExit(f"142280 PDF page count {len(reader.pages)} != {SPEC['page_count']}")
    writer = PdfWriter()
    for index in SPEC["pages"]:
        writer.add_page(reader.pages[index])
    buf = BytesIO()
    writer.write(buf)
    return buf.getvalue()


def existing_3p(client: KimcoClient) -> dict[str, int]:
    found: dict[str, int] = {}
    url = client._url("ap_invoices")
    offset = 0
    total = None
    while total is None or offset < total:
        response = client.request("GET", url, params={"pageSize": 2000, "offset": offset})
        if response.status_code != 200:
            raise SystemExit(f"AP list HTTP {response.status_code}")
        payload = response.json()
        items = payload.get("items") or []
        total = int(payload.get("totalCount") or 0)
        for item in items:
            values = item.get("values") or {}
            vendor = str(values.get("Vendor_$_Display_Name") or "")
            number = str(values.get("Invoice_Number") or "").strip()
            if "3P" not in vendor.upper():
                continue
            if number and item.get("id") not in (None, ""):
                found[number] = int(item["id"])
        if not items:
            break
        offset += len(items)
    return found


def snapshot(client: KimcoClient, kimco_id: int) -> dict[str, Any]:
    rec = client.get_item("ap_invoices", kimco_id)
    values = rec.get("values") or {}
    lists = rec.get("lists") or {}
    lines = []
    for line in lists.get("APInvoiceLine") or []:
        lv = line.get("values") or {}
        lines.append(
            {
                "id": line.get("id"),
                "qty": lv.get("Quantity"),
                "price": money(lv.get("Unit_Price")),
                "ext": money(lv.get("Extended_Amount")),
                "receipt": lookup_id(lv.get("Receipt")),
                "po": lookup_text(lv.get("Purchase_Order_Number")),
                "po_line": lookup_text(lv.get("Purchase_Order_Line")),
            }
        )
    charges = []
    for charge in lists.get("InvoiceAdditionalCharges") or []:
        cv = charge.get("values") or {}
        charges.append(
            {
                "id": charge.get("id"),
                "kind": cv.get("Additional_Charges"),
                "amount": money(cv.get("Amount")),
            }
        )
    comments = []
    for comment in lists.get("Comments_1") or []:
        comments.append(
            {
                "id": comment.get("id"),
                "html": ((comment.get("values") or {}).get("HtmlValue") or "")[:4000],
            }
        )
    try:
        atts = client.list_attachments(kimco_id)
    except Exception:
        atts = []
    return {
        "id": kimco_id,
        "invoice": values.get("Invoice_Number"),
        "type": values.get("Invoice_Type"),
        "date": values.get("Invoice_Date"),
        "due": values.get("Invoice_Due_Date"),
        "amount": money(values.get("Invoice_Amount")),
        "verification": money(values.get("Invoice_Verification_Amount")),
        "status": values.get("Status"),
        "posted": values.get("Posted"),
        "po": lookup_text(values.get("Purchase_Order")),
        "po_id": lookup_id(values.get("Purchase_Order")),
        "batch": lookup_text(values.get("AP_Invoice_Batch")),
        "batch_id": lookup_id(values.get("AP_Invoice_Batch")),
        "vendor": lookup_text(values.get("Vendor")),
        "modified": values.get("ModifiedOn"),
        "lines": lines,
        "charges": charges,
        "comments_1": comments,
        "attachments": [a.get("name") for a in atts],
        "comment_ids": [c.get("id") for c in comments],
        "line_ids": [line.get("id") for line in lines],
    }


def fingerprint(live: dict[str, Any]) -> dict[str, Any]:
    return {
        "amount": live.get("amount"),
        "verification": live.get("verification"),
        "batch_id": live.get("batch_id"),
        "posted": live.get("posted"),
        "status": live.get("status"),
        "po_id": live.get("po_id"),
        "comment_ids": live.get("comment_ids"),
        "line_ids": live.get("line_ids"),
        "modified": live.get("modified"),
        "attachments": live.get("attachments"),
    }


def prove_receipt(client: KimcoClient, row: dict[str, Any]) -> None:
    rec = client.get_item("receipts", int(row["id"]))
    values = rec.get("values") or {}
    ap = values.get("AP_Invoice_Number")
    ap_text = lookup_text(ap) if isinstance(ap, dict) else str(ap or "").strip()
    if values.get("Invoiced") is True or ap_text:
        raise SystemExit(f"receipt {row['id']} already billed on {ap_text or 'an invoice'}")
    qty = money(values.get("Quantity_Received"))
    price = money(values.get("PO_Item_Number_$_Unit_Price"))
    if qty != float(row["qty"]) or price != float(row["po_unit"] if "po_unit" in row else row["price"]):
        raise SystemExit(f"receipt {row['id']} live qty {qty} price {price} != expected {row}")


def prove_spec_receipts(client: KimcoClient) -> None:
    expect = SPEC["receipt_expect"]
    for rid in SPEC["receipts"]:
        wanted = expect[int(rid)]
        prove_receipt(client, {"id": rid, "qty": wanted["qty"], "price": wanted["price"]})
    for row in SPEC["skip_receipts"]:
        prove_receipt(client, row)


def header_payload() -> dict[str, Any]:
    return {
        "AP_Invoice_Batch": {"id": BATCH_ID},
        "Vendor": {"id": VENDOR_ID},
        "Invoice_Number": SPEC["invoice"],
        "Invoice_Type": 3,
        "Invoice_Date": kimco_datetime(SPEC["date"]),
        "Invoice_Verification_Amount": float(SPEC["total"]),
        "Invoice_Due_Date": kimco_datetime(SPEC["due"]),
        "Terms_Code": {"id": TERMS_ID},
        "Currency": {"id": CURRENCY_ID},
        "Remit_To_Address": {"id": REMIT_ID},
        "Transaction_Date": kimco_datetime(SPEC["date"]),
        "Comments": "API Agent",
    }


def select_open_receipts(client: KimcoClient, kimco_id: int, receipt_ids: list[int]) -> dict[str, Any]:
    invoice = client.get_item("ap_invoices", kimco_id)
    lines = []
    for rid in receipt_ids:
        receipt = client.get_item("receipts", int(rid))
        values = receipt.get("values") or {}
        ap = values.get("AP_Invoice_Number")
        ap_text = lookup_text(ap) if isinstance(ap, dict) else str(ap or "").strip()
        if values.get("Invoiced") is True or ap_text:
            raise SystemExit(f"receipt {rid} already billed on {ap_text or 'an invoice'}")
        lines.append(receipt_line_values_from_records(invoice, receipt))
    payload = select_receipts_payload(lines, invoice_id=kimco_id)
    put = client.request("PUT", client._record_url("ap_invoices", kimco_id), json=payload)
    return {
        "http": put.status_code,
        "error": "" if put.status_code < 400 else put.text[:600],
        "count": len(receipt_ids),
    }


def price_note() -> str:
    bits = []
    for row in SPEC["small_gaps"]:
        gap = round(float(row["invoice_ext"]) - float(row["po_ext"]), 2)
        bits.append(
            f"qty {row['qty']} of part {row['part']} on PO {row['po']} is "
            f"${row['invoice_unit']:.2f} on the invoice (${row['invoice_ext']:.2f}) "
            f"and receipt {row['id']} is at PO price ${row['po_unit']:.2f} "
            f"(${row['po_ext']:.2f}). |PPV| ${abs(gap):.2f} is under $75, so that "
            "receipt was selected at the PO extended cost and no unit-price PPV was stacked on it."
        )
    bits.append(
        "qty 2 of part 1007044-1 on PO 59002 is $97.50 on the invoice ($195.00) "
        "and receipt 23917 is already at PO price $361.46 ($722.92). |PPV| $527.92 "
        "is $75 or more, so that receipt was not selected and the unit price was not changed. "
        "Unreceive, set PO 59002 part 1007044-1 to $97.50, and re-receive qty 2."
    )
    bits.append(
        "qty 3 of part 1007044-1 on PO 59031 is $97.50 on the invoice ($292.50). "
        "Receipt 23920 qty 2 is already at PO price $361.46 ($722.92) and receipt 23921 "
        "qty 1 is already at PO price $361.46 ($361.46). Combined PO extended $1084.38. "
        "|PPV| $791.88 is $75 or more, so those receipts were not selected and the unit "
        "price was not changed. Unreceive, set PO 59031 part 1007044-1 to $97.50, "
        "and re-receive qty 3."
    )
    selected = ", ".join(str(rid) for rid in SPEC["receipts"])
    return (
        "AP Clerk: @Shawn McKibben HOLD price_variance on 3P 142280. "
        f"Matching receipts selected: {selected}. " + " ".join(bits)
    )


def price_html() -> str:
    text = price_note()
    mention = (
        f'<span data-mention-id="{SHAWN_MENTION_ID}" '
        'data-mention-name="Shawn McKibben" '
        'data-mention-email="Shawn.McKibben@kannonmfg.com" '
        'class="prosemirror-mention-node">@Shawn McKibben</span>'
    )
    return f"<p>{text.replace('@Shawn McKibben', mention, 1)}</p>"


def comment_id_matching(live: dict[str, Any], needle: str) -> int | None:
    for comment in live.get("comments_1") or []:
        if needle in str(comment.get("html") or ""):
            if comment.get("id") not in (None, ""):
                return int(comment["id"])
    return None


def archive_email(graph: GraphClient) -> dict[str, Any]:
    hits = graph.search_messages(ALLOWED_MAILBOX, "76945", top=10)
    chosen = None
    for msg in hits:
        subject = str(msg.get("subject") or "")
        if SPEC["message_subject_needle"] in subject and "3pindustries" in json.dumps(msg.get("from") or {}).lower():
            chosen = msg
            break
    if chosen is None:
        return {"status": "not-found", "folder": "", "moved_this_run": False, "email_moved": "no"}
    message_id = str(chosen.get("id") or "")
    full = graph.get_message(
        ALLOWED_MAILBOX,
        message_id,
        select="id,subject,parentFolderId,from,receivedDateTime",
    )
    subject = str(full.get("subject") or "")
    if SPEC["message_subject_needle"] not in subject:
        raise SystemExit(f"refusing to move unexpected subject {subject}")
    folder = graph.resolve_fort_worth_folder(ALLOWED_MAILBOX)
    fort_id = str(folder.get("id") or "")
    parent = str(full.get("parentFolderId") or "")
    display = folder.get("displayName") or FORT_WORTH_FOLDER_DISPLAY_NAME
    if fort_id and parent == fort_id:
        return {
            "status": "already",
            "folder": display,
            "moved_this_run": False,
            "email_moved": "yes",
            "subject": subject,
        }
    if not fort_id:
        return {"status": "no-folder", "folder": "", "moved_this_run": False, "email_moved": "no", "subject": subject}
    moved = graph.move_message(ALLOWED_MAILBOX, message_id, fort_id)
    new_id = str(moved.get("new_id") or message_id)
    after = graph.get_message(ALLOWED_MAILBOX, new_id, select="id,subject,parentFolderId")
    in_folder = str(after.get("parentFolderId") or "") == fort_id
    return {
        "status": "moved" if in_folder else moved.get("status"),
        "folder": display if in_folder else "",
        "moved_this_run": bool(in_folder),
        "email_moved": "yes" if in_folder else "no",
        "http": moved.get("http"),
        "subject": after.get("subject") or subject,
    }


def parse_pdf_date(value: str | None) -> str:
    if not value:
        return ""
    try:
        return datetime.strptime(value, "%m/%d/%Y").date().isoformat()
    except ValueError:
        return value


def inventoried_invoices() -> dict[str, dict[str, Any]]:
    if not PAGES_JSON.exists():
        return {}
    pages = json.loads(PAGES_JSON.read_text())
    found: dict[str, dict[str, Any]] = {}
    for rec in pages:
        number = str(rec.get("invoice") or "").strip()
        if not number:
            continue
        current = found.get(number)
        if current is None:
            found[number] = {
                "invoice": number,
                "pdf_date": parse_pdf_date(rec.get("date")),
                "pdf_total": (rec.get("totals") or [None])[-1],
                "subject": rec.get("subject"),
                "received": rec.get("received"),
                "freight": bool(rec.get("freight")),
                "corrected": bool(rec.get("corrected")),
                "file": rec.get("file"),
                "page": rec.get("page"),
            }
            continue
        if rec.get("corrected"):
            current["corrected"] = True
        if not current.get("pdf_date") and rec.get("date"):
            current["pdf_date"] = parse_pdf_date(rec.get("date"))
        if not current.get("pdf_total") and rec.get("totals"):
            current["pdf_total"] = rec["totals"][-1]
    return found


def skip_why(number: str, kimco_id: int | None, meta: dict[str, Any]) -> str:
    if number == "142173":
        return (
            f"excluded. Existing KIMCO bill {kimco_id}. The vendor PDF still shows "
            "qty 10 of part 2975-1 shipped on PO 58887, not a corrected invoice for 20. "
            "Existing bill was not modified."
        )
    if kimco_id is None:
        return "inventoried but not entered this run"
    extra = ""
    if meta.get("corrected"):
        extra = " A corrected PDF is also in the mailbox. Existing bill was not modified."
    elif meta.get("pdf_date") and meta["pdf_date"] < "2026-08-01":
        extra = " Invoice date is before 2026-08-01. The mailbox copy was received on or after Aug 1."
    return f"already in KIMCO bill {kimco_id}.{extra}"


def receipt_label(lines: list[dict[str, Any]]) -> str:
    bits = []
    for line in lines:
        bits.append(f"{line.get('receipt')} qty {line.get('qty')} @ {line.get('price')}")
    return "; ".join(bits) if bits else "none"


def write_outputs(
    entered: list[dict[str, Any]],
    skipped: list[dict[str, Any]],
    *,
    batch_name: str,
    untouched_ok: bool,
) -> None:
    payload = {
        "run": "3p-finish-2026-09-25",
        "host": "https://live.kimcoerp.com",
        "vendor": "999-3P INDUSTRIES",
        "vendor_id": VENDOR_ID,
        "batch_id": BATCH_ID,
        "batch_name": batch_name,
        "posted": False,
        "emails_sent": False,
        "packing_slip_gate": "suspended",
        "login": "API Agent key. Bills are not posted. Treyce Hodges (id 33) posts.",
        "untouched_verified": untouched_ok,
        "untouched": [
            {"invoice": "142179", "kimco_id": 10306},
            {"invoice": "142188", "kimco_id": 10308},
            {"invoice": "142216", "kimco_id": 10312},
            {"invoice": "142173", "kimco_id": 10187},
            {"invoice": "142231", "kimco_id": 10186},
        ],
        "scope": (
            "Every 3P invoice PDF in accountspayable@ from 2026-08-01 onward, "
            "including Inbox and 9 - FORT WORTH ARCHIVE. Only 142280 was not already "
            "a vendor-1 bill. 142173 was excluded because no corrected qty-20 invoice arrived."
        ),
        "invoices": entered,
        "skipped": skipped,
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(payload, indent=2, default=str) + "\n")

    wb = Workbook()
    first = wb.active
    first.title = f"This run ({len(entered)})"
    first.append(RUN_HEADERS)
    green = PatternFill("solid", fgColor="C6EFCE")
    yellow = PatternFill("solid", fgColor="FFEB9C")
    for row in entered:
        live = row.get("live") or {}
        batch_id = live.get("batch_id")
        batch = live.get("batch") or ""
        batch_cell = f"{batch} ({batch_id})" if batch_id else (row.get("batch") or "")
        status = row.get("status") or ""
        first.append(
            [
                row.get("invoice"),
                status,
                row.get("exception") or "",
                "parts" if row.get("kind") == "parts" else row.get("kind"),
                row.get("pdf_total"),
                row.get("kimco_id"),
                batch_cell,
                row.get("transfer_ap") or "no",
                receipt_label(row.get("receipts_selected") or []),
                row.get("ppv") if row.get("ppv") not in (None, "") else None,
                row.get("comments_1_id"),
                (row.get("email") or {}).get("email_moved") or "no",
                row.get("comments_1") or "",
            ]
        )
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in first[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for excel_row in first.iter_rows(min_row=2):
        status = str(excel_row[1].value or "")
        if status == "Success":
            excel_row[1].fill = green
        elif status == "HOLD":
            excel_row[1].fill = yellow
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [14, 12, 18, 12, 12, 16, 28, 14, 55, 12, 16, 14, 80]
    for index, width in enumerate(widths, 1):
        first.column_dimensions[get_column_letter(index)].width = width
    first.freeze_panes = "A2"
    first.auto_filter.ref = f"A1:M{first.max_row}"
    first.row_dimensions[1].height = 22

    skipped_ws = wb.create_sheet("Skipped")
    skipped_ws.append(
        ["Invoice #", "KIMCO bill #", "Why", "Invoice date", "PDF total", "Email subject", "Received"]
    )
    for row in skipped:
        skipped_ws.append(
            [
                row.get("invoice"),
                row.get("kimco_id"),
                row.get("why"),
                row.get("pdf_date"),
                row.get("pdf_total"),
                row.get("subject"),
                (row.get("received") or "")[:19],
            ]
        )
    for cell in skipped_ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for excel_row in skipped_ws.iter_rows(min_row=2):
        for cell in excel_row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, width in enumerate([14, 16, 88, 16, 14, 70, 22], 1):
        skipped_ws.column_dimensions[get_column_letter(index)].width = width
    skipped_ws.freeze_panes = "A2"
    skipped_ws.auto_filter.ref = f"A1:G{skipped_ws.max_row}"
    wb.save(OUT_XLSX)


def enter(client: KimcoClient, graph: GraphClient) -> dict[str, Any]:
    prove_spec_receipts(client)
    pdf = page_bytes()
    if PdfReader(BytesIO(pdf)).pages.__len__() != 1:
        raise SystemExit("refusing to attach more than the invoice page")
    created_id, _body, status, error = client.create("ap_invoices", header_payload())
    out: dict[str, Any] = {
        "invoice": SPEC["invoice"],
        "kind": "parts",
        "pdf_date": SPEC["date"].isoformat(),
        "pdf_due": SPEC["due"].isoformat(),
        "pdf_total": SPEC["total"],
        "po": None,
        "po_on_pdf": SPEC["po_on_pdf"],
        "pos": SPEC["pos"],
        "create_http": status,
        "create_error": error,
        "kimco_id": created_id,
        "ppv": None,
        "transfer_ap": "no",
        "comments_1": "",
        "comments_1_id": None,
        "exception": None,
    }
    if created_id is None:
        out["status"] = "Fail"
        out["why"] = f"header create HTTP {status} {error}"
        return out
    if int(created_id) in DO_NOT_TOUCH_IDS:
        raise SystemExit(f"refusing to edit KIMCO {created_id}")
    attach = client.try_official_attach(
        created_id,
        name="142280.pdf",
        content_type="application/pdf",
        size=len(pdf),
        content=pdf,
    )
    out["attach"] = attach
    if attach == "attached":
        out["email"] = archive_email(graph)
    else:
        out["email"] = {"status": "not-moved", "email_moved": "no", "moved_this_run": False, "folder": ""}
    selected = select_open_receipts(client, created_id, list(SPEC["receipts"]))
    out["select"] = selected
    live = snapshot(client, created_id)
    if selected["http"] >= 400:
        out["status"] = "Fail"
        out["why"] = f"select receipts HTTP {selected['http']} {selected['error']}"
        out["live"] = live
        return out
    selected_ids = {line["receipt"] for line in live["lines"]}
    expected_ids = set(SPEC["receipts"])
    out["receipts_selected"] = [
        {
            "id": line["receipt"],
            "qty": line["qty"],
            "price": line["price"],
            "ext": line["ext"],
            "po": line["po"],
            "po_line": line["po_line"],
        }
        for line in live["lines"]
    ]
    receipt_ext = round(sum(line["ext"] or 0 for line in live["lines"]), 2)
    gap = round(SPEC["total"] - receipt_ext, 2)
    out["receipt_extended"] = receipt_ext
    out["price_gap"] = gap
    note = price_note()
    if selected_ids != expected_ids:
        out["status"] = "Fail"
        out["exception"] = "price_variance"
        out["why"] = f"selected {sorted(selected_ids)} != expected {sorted(expected_ids)}"
        out["live"] = live
        out["comments_1"] = note
        return out
    # The $62 unit gap is inside this total. The unselected 1007044 lines make |gap| >= $75.
    if abs(gap) < PPV_HOLD:
        out["status"] = "Fail"
        out["why"] = f"expected |gap| >= 75 after leaving the 1007044 receipts off the bill, got {gap}"
        out["live"] = live
        return out
    comment = comments_1_shawn(client, created_id, price_html())
    live = snapshot(client, created_id)
    comment_id = comment_id_matching(live, "AP Clerk:")
    mention_ok = any('data-mention-id="104"' in str(c.get("html") or "") for c in live["comments_1"])
    moved: dict[str, Any] = {"status": "not-moved"}
    if comment.get("status") == "persisted" and comment_id and mention_ok:
        moved = apply_transfer_ap_batch_move(client, kimco_id=created_id)
        live = snapshot(client, created_id)
    out["live"] = live
    out["comments_1"] = note
    out["comments_1_id"] = comment_id
    out["comments_1_put"] = comment
    out["transfer_ap_detail"] = moved
    out["ppv"] = None
    on_transfer = live.get("batch_id") not in (None, BATCH_ID) and moved.get("status") in {
        "moved",
        "already-on-transfer-ap",
    }
    out["transfer_ap"] = "yes" if on_transfer else "no"
    unit_ok = all(
        line["price"] == SPEC["receipt_expect"][int(line["receipt"])]["price"]
        for line in live["lines"]
        if line.get("receipt") in SPEC["receipt_expect"] or int(line.get("receipt") or 0) in SPEC["receipt_expect"]
    )
    if on_transfer and mention_ok and attach == "attached" and live["posted"] in (None, "", False) and unit_ok:
        out["status"] = "HOLD"
        out["exception"] = "price_variance"
        out["why"] = note
    else:
        out["status"] = "HOLD"
        out["exception"] = "price_variance"
        out["why"] = (
            f"{note} transfer={moved.get('status')} comment={comment.get('status')} "
            f"mention={mention_ok} batch={live.get('batch_id')} gap={gap} attach={attach} unit_ok={unit_ok}"
        )
    return out


def main() -> int:
    if not PDF_PATH.exists():
        raise SystemExit(f"missing {PDF_PATH}")
    creds = load_credentials(target="live")
    if not creds.ready:
        raise SystemExit(creds.error)
    client = KimcoClient.authenticate(creds.instance_url, creds.key or "", creds.password or "", target="live")
    graph = _optional_graph_client()
    if graph is None:
        raise SystemExit("Graph client required")
    batch = client.get_item("ap_batches", BATCH_ID)
    batch_name = str((batch.get("values") or {}).get("AP_Invoice_Batch_ID") or "")
    if batch_name != "API Agent - 9/22/26 3P":
        raise SystemExit(f"batch 724 name is {batch_name!r}")
    print("BATCH", BATCH_ID, batch_name, flush=True)
    already = existing_3p(client)
    if already.get("142188") != 10308 or already.get("142179") != 10306 or already.get("142216") != 10312:
        raise SystemExit(
            f"HOLD ids changed 142179={already.get('142179')} 142188={already.get('142188')} 142216={already.get('142216')}"
        )
    if already.get("142173") != 10187:
        raise SystemExit(f"142173 id changed: {already.get('142173')}")
    if "142280" in already:
        raise SystemExit(f"142280 already in KIMCO {already['142280']}; refusing to modify it")
    before = {kid: fingerprint(snapshot(client, kid)) for kid in sorted(DO_NOT_TOUCH_IDS)}
    result = enter(client, graph)
    after = {kid: fingerprint(snapshot(client, kid)) for kid in sorted(DO_NOT_TOUCH_IDS)}
    untouched_ok = before == after
    if not untouched_ok:
        print("UNTOUCHED CHANGED", flush=True)
        for kid in sorted(DO_NOT_TOUCH_IDS):
            if before[kid] != after[kid]:
                print(kid, before[kid], after[kid], flush=True)
    inventoried = inventoried_invoices()
    skipped = []
    for number in sorted(inventoried):
        if number == "142280":
            continue
        meta = inventoried[number]
        kimco_id = already.get(number)
        skipped.append(
            {
                "invoice": number,
                "kimco_id": kimco_id,
                "why": skip_why(number, kimco_id, meta),
                "pdf_date": meta.get("pdf_date"),
                "pdf_total": meta.get("pdf_total"),
                "subject": meta.get("subject"),
                "received": meta.get("received"),
                "freight": meta.get("freight"),
                "corrected": meta.get("corrected"),
            }
        )
    live = result.get("live") or {}
    entered = []
    if result.get("kimco_id"):
        entered.append(
            {
                "invoice": result["invoice"],
                "pdf_date": result.get("pdf_date"),
                "pdf_total": result.get("pdf_total"),
                "type": result.get("kind"),
                "kind": result.get("kind"),
                "po": None,
                "po_on_pdf": result.get("po_on_pdf"),
                "pos": result.get("pos") or [],
                "kimco_id": result.get("kimco_id"),
                "status": result.get("status"),
                "exception": result.get("exception"),
                "why": result.get("why"),
                "receipts": result.get("receipts_selected") or [],
                "receipts_selected": result.get("receipts_selected") or [],
                "ppv": result.get("ppv"),
                "price_gap": result.get("price_gap"),
                "transfer_ap": result.get("transfer_ap"),
                "charges": live.get("charges") or [],
                "invoice_amount_after": live.get("amount"),
                "verification_amount": live.get("verification"),
                "comments_1": result.get("comments_1") or "",
                "comments_1_id": result.get("comments_1_id"),
                "email_folder": (result.get("email") or {}).get("folder") or "",
                "email_moved": (result.get("email") or {}).get("email_moved"),
                "email_moved_this_run": (result.get("email") or {}).get("moved_this_run"),
                "email_subject": (result.get("email") or {}).get("subject"),
                "batch_id": live.get("batch_id"),
                "batch": live.get("batch"),
                "posted": live.get("posted"),
                "invoice_type": live.get("type"),
                "attachments": live.get("attachments"),
                "email": result.get("email"),
                "live": live,
                "select": result.get("select"),
                "attach": result.get("attach"),
            }
        )
    write_outputs(entered, skipped, batch_name=batch_name, untouched_ok=untouched_ok)
    print(
        json.dumps(
            {
                "invoice": result.get("invoice"),
                "status": result.get("status"),
                "exception": result.get("exception"),
                "kimco_id": result.get("kimco_id"),
                "amount": result.get("pdf_total"),
                "live_amount": live.get("amount"),
                "verification": live.get("verification"),
                "batch": live.get("batch"),
                "batch_id": live.get("batch_id"),
                "posted": live.get("posted"),
                "gap": result.get("price_gap"),
                "transfer": result.get("transfer_ap"),
                "comment": result.get("comments_1_id"),
                "email": result.get("email"),
                "attach": result.get("attach"),
                "lines": len(live.get("lines") or []),
                "untouched_ok": untouched_ok,
                "skipped": len(skipped),
            },
            default=str,
        ),
        flush=True,
    )
    ok = (
        result.get("status") == "HOLD"
        and result.get("exception") == "price_variance"
        and result.get("transfer_ap") == "yes"
        and result.get("comments_1_id")
        and (result.get("email") or {}).get("email_moved") == "yes"
        and untouched_ok
        and live.get("posted") in (None, "", False)
    )
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
