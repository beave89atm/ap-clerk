"""QUALITY V1.1 gates. No live Graph or KIMCO I/O. Never prints secrets."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ap_clerk.cli import _process_invoice
from ap_clerk.gates import (
    RESULT_FAIL,
    RESULT_HOLD,
    RESULT_INCOMPLETE,
    RESULT_SUCCESS,
    RESULT_VALUES,
    attach_succeeded,
    drop_fee_disguised_as_ppv,
    find_live_po,
    finish_gate,
    merchandise_amount,
    preflight_parse_gate,
    success_is_legal,
)
from ap_clerk.graph import FLAG_AI_HOLD, FLAG_FLAGGED
from ap_clerk.inbox import PO_FILE_RE, skip_rows_for_report
from ap_clerk.pdf_invoice import is_purchase_order_document, parse_invoice_text
from ap_clerk.report import write_report
from ap_clerk.rules import classify_mail, match_receipts


def _sample_values():
    return {
        "Remit_To_Address": {"id": 1, "text": "remit"},
        "Terms_Code": {"id": 2, "text": "Net 30"},
        "Currency": {"id": 3, "text": "USD"},
    }


class RecordingKimco:
    def __init__(self, *, attach="no-pdf-on-vm", select="blocked-405", created_id=80):
        self.target = "live"
        self.attach = attach
        self.select = select
        self.created_id = created_id
        self.created: list[dict] = []
        self.selected: list[tuple] = []

    def create(self, service, values):
        self.created.append(values)
        return self.created_id, {"id": self.created_id, "values": values}, 200, ""

    def get_item(self, service, item_id):
        return {"id": item_id, "values": _sample_values()}

    def try_official_attach(self, *args, **kwargs):
        return self.attach

    def try_select_receipts(self, invoice_id, receipt_ids=None):
        self.selected.append((invoice_id, receipt_ids or []))
        return self.select

    def try_put_probe_rejected(self, service, item_id):
        return "blocked-405"


def _row(inv, *, kimco=None, po_index=None, receipts=None, samples=None, pdf_dir=None):
    return _process_invoice(
        kimco or RecordingKimco(),
        inv,
        batch={"id": 1},
        batch_label="API Agent - 9/8/26 (1)",
        invoice_by_number={},
        vendor_samples=samples
        or [{"vendor_id": 9, "vendor_text": inv.get("vendor") or "Vendor", "invoice_id": 100, "po_text": ""}],
        po_index=po_index or {},
        receipts=receipts,
        pdf_dir=pdf_dir,
        flag_outlook=False,
    )


def test_result_values_are_success_incomplete_hold_fail():
    assert RESULT_VALUES == (RESULT_SUCCESS, RESULT_INCOMPLETE, RESULT_HOLD, RESULT_FAIL)


def test_success_is_illegal_without_attach_and_receipts_when_po():
    assert success_is_legal(
        header_created=True,
        attach_status="blocked-405",
        po="58634",
        receipts_selected=False,
    ) is False
    result, why = finish_gate(
        header_created=True,
        attach_status="blocked-405",
        po="58634",
        receipts_selected=False,
        kimco_id=9711,
    )
    assert result == RESULT_INCOMPLETE
    assert "finish" in why
    assert "Not Entered in AI" in why
    assert attach_succeeded("blocked-405") is False
    assert attach_succeeded("attached") is True


def test_success_requires_header_attach_and_receipts_when_po(tmp_path: Path):
    pdf = tmp_path / "TXFT499356.pdf"
    pdf.write_bytes(b"%PDF-1.4 finished")
    row = _row(
        {
            "vendor": "Fastenal Company",
            "invoice_number": "TXFT499356",
            "date": "2026-08-26",
            "po": "58700",
            "amount": 40.0,
            "pdf_path": str(pdf),
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        kimco=RecordingKimco(attach="attached", select="selected"),
        po_index={"58700": {"id": 3, "text": "58700-FASTENAL", "vendor_id": 9, "lines": []}},
        receipts=[{"slip": "TXFT499356", "qty": 6, "part": "FAST-1", "po_line": 2, "po": "58700", "id": 44}],
    )
    assert row["Result"] == RESULT_SUCCESS
    assert row["Attach status"] == "attached"
    assert row["Notes"] == ""


def test_header_only_blocked_attach_is_incomplete_not_success(tmp_path: Path):
    pdf = tmp_path / "26167.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    row = _row(
        {
            "vendor": "Capital Machine Technologies, Inc",
            "invoice_number": "26167",
            "date": "2026-07-30",
            "po": "58634",
            "amount": 403.0,
            "pdf_path": str(pdf),
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        kimco=RecordingKimco(attach="blocked-405", select="blocked-405"),
        po_index={"58634": {"id": 3, "text": "58634-CAPITAL", "vendor_id": 45, "lines": []}},
        receipts=[{"slip": "26167", "qty": 1, "part": "BLADE", "po": "58634", "id": 9}],
        samples=[{"vendor_id": 45, "vendor_text": "Capital Machine", "invoice_id": 100, "po_text": ""}],
    )
    assert row["Result"] == RESULT_INCOMPLETE
    assert row["Result"] != RESULT_SUCCESS
    assert row["Flag status"] != "entered-in-ai"
    assert "finish" in row["Why"]
    assert row["Notes"] == ""


def test_incomplete_gets_ai_hold_not_entered_in_ai():
    class FakeGraph:
        def __init__(self):
            self.held = []
            self.matched = []

        def flag_hold(self, mailbox, message_id):
            self.held.append(message_id)
            return FLAG_AI_HOLD

        def flag_matched(self, mailbox, message_id):
            self.matched.append(message_id)
            return FLAG_FLAGGED

    graph = FakeGraph()
    row = _process_invoice(
        RecordingKimco(attach="blocked-405"),
        {
            "vendor": "Crosslink Powder Coating of TX, LLC",
            "invoice_number": "27756",
            "date": "2026-08-03",
            "po": None,
            "amount": 1252.18,
            "graph_message_id": "AAMk-incomplete",
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": ""},
        },
        batch={"id": 1},
        batch_label="API Agent - 9/8/26 (1)",
        invoice_by_number={},
        vendor_samples=[{"vendor_id": 9, "vendor_text": "Crosslink Powder Coating", "invoice_id": 100, "po_text": ""}],
        po_index={},
        pdf_dir=None,
        graph_client=graph,
        flag_outlook=True,
    )
    assert row["Result"] == RESULT_INCOMPLETE
    assert row["Flag status"] == FLAG_AI_HOLD
    assert row["Flag status"] != "entered-in-ai"
    assert graph.held == ["AAMk-incomplete"]
    assert graph.matched == []


def test_gas_0040323616_amount_from_pdf_not_subject():
    parsed = parse_invoice_text(
        "ORIGINAL INVOICE\nINVOICE DATE ACCOUNT NUMBER INVOICE NUMBER\n"
        "08/01/26   A3050      0040323616\nGas and Supply North Texas, LLC\n"
        "CUS P/O #\n58001\nAmount Due: 125.50\n",
        subject="Gas invoice 0040323616 Amount 9,999.00",
        from_address="billing@gasandsupply.com",
        filename="billing01_A3050_c.pdf",
    )
    assert parsed["invoice_number"] == "0040323616"
    assert parsed["amount"] == 125.50
    assert parsed["amount"] != 9999.00
    assert parsed["field_sources"]["amount"] == "pdf"
    assert parsed["field_sources"]["invoice_number"] == "pdf"


def test_preflight_holds_when_pdf_total_missing():
    parsed = parse_invoice_text(
        "Gas and Supply\nInvoice Number 0040323616\nInvoice Date 08/01/2026\n",
        subject="Gas 0040323616 Amount 418.93",
        filename="0040323616.pdf",
    )
    ok, why = preflight_parse_gate(parsed)
    assert ok is False
    assert "parse-error" in why
    assert "preflight-parse" in why
    row = _row({**parsed, "vendor": "Gas and Supply North Texas, LLC"})
    assert row["Result"] == RESULT_HOLD
    assert row["KIMCO id"] == ""


def test_msc_and_mcqueary_invoice_number_from_pdf_not_filename():
    msc = parse_invoice_text(
        "MSC INDUSTRIAL SUPPLY CO.\nCustomer Number Invoice Number\n02627782 64564711\nAmount Due $10.00\nInvoice Date 08/01/2026",
        filename="02627782_wrong.PDF",
        from_address="DoNotReply@invoices.mscdirect.com",
    )
    assert msc["invoice_number"] == "64564711"
    assert msc["invoice_number"] != "02627782"
    assert msc["field_sources"]["invoice_number"] == "pdf"

    mcq = parse_invoice_text(
        "McQueary Industries\nInvoice 98-76543\nInvoice Date 08/01/2026\nAmount Due 12.00",
        filename="Invoice_12-34567.pdf",
        from_name="McQueary",
    )
    assert mcq["invoice_number"] == "98-76543"
    assert mcq["invoice_number"] != "12-34567"

    filename_only = parse_invoice_text(
        "",
        filename="02627782.PDF",
        subject="MSC file",
        from_address="DoNotReply@invoices.mscdirect.com",
    )
    ok, why = preflight_parse_gate(filename_only)
    assert ok is False
    assert "invoice #" in why.lower() or "preflight-parse" in why


def test_legacy_po_pdf_is_rejected_as_parse_error():
    assert PO_FILE_RE.search("Purchase_Order_58861.pdf")
    assert is_purchase_order_document(filename="Purchase_Order_58861.pdf")
    text = "PURCHASE ORDER\nPO 58861\nShip To: Kannon\nQty 1 Steel"
    assert is_purchase_order_document(text=text, filename="58861.pdf")
    parsed = parse_invoice_text(text, filename="Purchase_Order_58861.pdf")
    assert parsed["is_purchase_order_doc"] is True
    ok, why = preflight_parse_gate(parsed)
    assert ok is False
    assert "purchase order" in why.lower()
    row = _row({**parsed, "vendor": "Legacy Wire Products", "invoice_number": "58861", "date": "2026-08-10", "amount": 1})
    assert row["Result"] == RESULT_HOLD
    assert row["KIMCO id"] == ""


def test_rmp_1470159_finds_po_by_vendor_and_part():
    parsed = parse_invoice_text(
        "RMP Industrial Supply Inc\nInvoice 1470159\nInvoice Date 07/30/2026\n"
        "625-5200-002    1    80.00\nAmount Due $80.00\n",
        from_name="RMP Industrial Supply Inc",
        filename="1470159.pdf",
    )
    assert parsed["invoice_number"] == "1470159"
    found = find_live_po(
        {
            "58690": {
                "id": 22,
                "text": "58690-RMP INDUSTRIAL",
                "vendor_id": 322,
                "vendor_text": "RMP Industrial Supply",
                "lines": [{"part": "625-5200-002", "po_line": 1, "qty": 1}],
            }
        },
        printed_po=parsed.get("po"),
        vendor="RMP Industrial Supply Inc",
        parts=["625-5200-002"],
    )
    assert found["info"] is not None
    assert found["number"] == "58690"
    assert found["how"] == "vendor+part"
    row = _row(
        {
            "vendor": "RMP Industrial Supply Inc",
            "invoice_number": "1470159",
            "date": "2026-07-30",
            "po": None,
            "pos": [],
            "amount": 80.0,
            "lines": [{"part": "625-5200-002", "qty": 1, "amount": 80.0}],
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": ""},
        },
        po_index={
            "58690": {
                "id": 22,
                "text": "58690-RMP INDUSTRIAL",
                "vendor_id": 322,
                "vendor_text": "RMP Industrial Supply",
                "lines": [{"part": "625-5200-002", "po_line": 1, "qty": 1, "amount": 80.0, "unit_price": 80.0}],
            }
        },
        samples=[{"vendor_id": 322, "vendor_text": "RMP Industrial Supply", "invoice_id": 44, "po_text": ""}],
        receipts=[{"po": "58690", "part": "625-5200-002", "qty": 1, "slip": "1470159"}],
    )
    assert row["PO"] == "58690"
    assert row["Result"] != RESULT_FAIL
    assert "Misc Type 4" not in row["Why"]


def test_willbanks_209663_209664_not_misc_has_po():
    for number, po in (("209663", "58510"), ("209664", "58511")):
        parsed = parse_invoice_text(
            f"Invoice {number} (USD)\nDate 7/30/2026\nWillbanks Metals\nYour PO  {po}  (6/24/2026)\nTotal Due   (USD) 100.00",
            from_name="Estevan Uribe",
        )
        assert parsed["invoice_number"] == number
        assert parsed["po"] == po
        row = _row(
            {
                **parsed,
                "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
            },
            po_index={
                po: {
                    "id": 8,
                    "text": f"{po}-WILLBANKS",
                    "vendor_id": 202,
                    "vendor_text": "Willbanks Metals",
                    "lines": [{"part": "STEEL", "amount": 100.0, "unit_price": 100.0}],
                }
            },
            samples=[{"vendor_id": 202, "vendor_text": "Willbanks Metals", "invoice_id": 9, "po_text": ""}],
            receipts=[{"po": po, "slip": number, "part": "STEEL", "qty": 1}],
        )
        assert row["PO"] == po
        assert row["Result"] != RESULT_FAIL
        assert "Misc Type 4" not in row["Why"]


def test_willbanks_printed_po_missing_on_live_is_hold_not_type_4():
    kimco = RecordingKimco()
    row = _row(
        {
            "vendor": "Willbanks Metals",
            "invoice_number": "209663",
            "date": "2026-07-30",
            "po": "58510",
            "amount": 100.0,
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        kimco=kimco,
        po_index={},
        samples=[{"vendor_id": 202, "vendor_text": "Willbanks Metals", "invoice_id": 9, "po_text": ""}],
    )
    assert row["Result"] == RESULT_HOLD
    assert "po" in row["Why"].lower()
    assert kimco.created == []


def test_capital_26167_second_pass_open_receipts_on_po():
    result = match_receipts(
        invoice_number="26167",
        invoice_lines=[{"part": "SAW-BLADE", "qty": 1}],
        receipts=[
            {"slip": "OTHER", "part": "DIFFERENT", "qty": 9, "po": "11111"},
            {"slip": "R-58634", "part": "SAW-BLADE", "qty": 1, "po": "58634", "po_line": 1},
        ],
        po_number="58634",
    )
    assert result["found"] is True
    assert result["hold_no_receipts"] is False
    assert result["matched"][0]["receipt"]["po"] == "58634"

    open_only = match_receipts(
        invoice_number="26167",
        invoice_lines=[{"part": "UNKNOWN-PART", "qty": 1}],
        receipts=[{"slip": "RECV-1", "part": "PO58634-01", "qty": None, "po": "58634", "name": "PO58634-CAPITAL"}],
        po_number="58634",
    )
    assert open_only["found"] is True
    assert open_only["hold_no_receipts"] is False
    assert open_only["matched"][0]["receipt"]["po"] == "58634"


def test_fastenal_txft499356_slip_equals_invoice():
    result = match_receipts(
        invoice_number="TXFT499356",
        invoice_lines=[],
        receipts=[{"slip": "TXFT499356", "qty": 6, "part": "FAST-42", "po": "58700"}],
        po_number="58700",
    )
    assert result["found"] is True
    assert result["hold_no_receipts"] is False


def test_aqpc_vendor_invoice_is_a_bill_not_noise():
    assert classify_mail(subject="American Quality Powder Coating") == "invoice"
    parsed = parse_invoice_text(
        "American Quality Powder Coating\nInvoice #: 4412\nInvoice Date 08/01/2026\nAmount Due $500.00",
        subject="AQPC job complete",
        from_name="American Quality Powder Coating",
        filename="AQPC-4412.pdf",
    )
    assert parsed["vendor"] == "American Quality Powder Coating"
    assert parsed["invoice_number"] == "4412"
    assert parsed["amount"] == 500.00
    assert parsed["parse_verified"] is True
    ok, _why = preflight_parse_gate(parsed)
    assert ok is True


def test_noise_skip_rows_name_bill_vs_noise_gate():
    rows = skip_rows_for_report(
        [
            {
                "vendor": "Bank",
                "class": "statement",
                "hold_reason": "statement",
                "receivedDateTime": "2026-08-01T12:00:00Z",
                "Flag status": "ai-hold",
            }
        ],
        "API Agent - 9/8/26",
    )
    assert rows[0]["Result"] == RESULT_HOLD
    assert "bill-vs-noise" in rows[0]["Why"]
    assert rows[0]["Notes"] == ""


def test_freight_is_not_double_counted_as_ppv():
    assert merchandise_amount(1128.67, [{"name": "Freight", "amount": 48.67}]) == 1080.00
    assert drop_fee_disguised_as_ppv(48.67, [{"name": "Freight", "amount": 48.67}]) == 0.0
    row = _row(
        {
            "vendor": "Xcaliber Industrial LLC",
            "invoice_number": "WB4337861572",
            "date": "2026-08-12",
            "po": "58896",
            "amount": 1128.67,
            "fees": [{"name": "Freight", "amount": 48.67}],
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        po_index={
            "58896": {
                "id": 4,
                "text": "58896-XCALIBER",
                "vendor_id": 339,
                "vendor_text": "Xcaliber",
                "lines": [{"part": "PART", "amount": 1080.00, "unit_price": 1080.00, "qty": 1}],
            }
        },
        samples=[{"vendor_id": 339, "vendor_text": "Xcaliber Industrial LLC", "invoice_id": 5, "po_text": ""}],
        receipts=[{"po": "58896", "slip": "WB4337861572", "part": "PART", "qty": 1}],
    )
    assert "Freight" in row["Fees and surcharges"]
    assert row["PPV"] == "none"


def test_nsa_and_coherent_aliases_still_resolve():
    nsa = _row(
        {
            "vendor": "National Specialty Alloys, Inc",
            "invoice_number": "453743",
            "date": "2026-08-26",
            "po": "59000",
            "amount": 50.0,
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        po_index={"59000": {"id": 12, "text": "59000-NATIONAL SPECIALTY ALLOYS", "vendor_id": 1386, "vendor_text": "National Specialty Alloys", "lines": []}},
        samples=[{"vendor_id": 1386, "vendor_text": "1386-NATIONAL SPECIALTY ALLOYS", "invoice_id": 44, "po_text": ""}],
        receipts=[{"po": "59000", "slip": "453743"}],
    )
    assert "vendor missing" not in nsa["Why"]
    coherent = _row(
        {
            "vendor": "Coherent Corp.",
            "invoice_number": "120953",
            "date": "2026-08-26",
            "po": "59001",
            "amount": 25.0,
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        po_index={"59001": {"id": 13, "text": "59001-COHERENT", "vendor_id": 1410, "vendor_text": "Coherent", "lines": []}},
        samples=[{"vendor_id": 1410, "vendor_text": "Coherent Corp.", "invoice_id": 55, "po_text": ""}],
        receipts=[{"po": "59001", "slip": "120953"}],
    )
    assert "vendor missing" not in coherent["Why"]


def test_excel_notes_column_stays_empty_for_treyce(tmp_path: Path):
    path = tmp_path / "AP-run-2026-09-08.xlsx"
    write_report(
        path,
        [
            {
                "Vendor": "Fastenal",
                "Invoice #": "TXFT1",
                "date": "2026-08-01",
                "PO": "58700",
                "Amount": 10,
                "Result": RESULT_INCOMPLETE,
                "Why": "Incomplete (finish): header created.",
                "KIMCO id": 1,
                "Batch": "API Agent - 9/8/26",
                "Fees and surcharges": "none",
                "PPV": "none",
                "Attach status": "blocked-405",
                "Flag in Outlook": "Yes",
                "Flag status": "ai-hold",
                "Notes": "agent must not write here",
            }
        ],
    )
    sheet = load_workbook(path).active
    headers = [sheet.cell(1, col).value for col in range(1, 16)]
    assert headers[5] == "Result"
    assert headers[-1] == "Notes"
    assert sheet.cell(2, 6).value == RESULT_INCOMPLETE
    assert not sheet.cell(2, 15).value


def test_process_invoice_cannot_return_success_without_finished_bill():
    row = _row(
        {
            "vendor": "Fastenal Company",
            "invoice_number": "TXFT499356",
            "date": "2026-08-26",
            "po": "58700",
            "amount": 40.0,
            "field_sources": {"invoice_number": "pdf", "date": "pdf", "amount": "pdf", "po": "pdf"},
        },
        po_index={"58700": {"id": 3, "text": "58700-FASTENAL", "vendor_id": 9, "lines": []}},
        receipts=[{"slip": "TXFT499356", "po": "58700", "qty": 6, "part": "FAST-1"}],
    )
    assert row["Result"] != RESULT_SUCCESS
    assert row["Result"] == RESULT_INCOMPLETE
    assert row["Flag status"] != "entered-in-ai"
